from __future__ import annotations

import json
import re
import shutil
from copy import copy
from dataclasses import asdict
from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

try:
    from . import my_pipeline_legacy as legacy
    from .pipeline_input_adapter import ensure_pipeline_xlsx, write_plan_json
    from .pipeline_agent_types import ControllerState, PipelinePaths, ReflectRegionPlan, ReflectStageResult
    from .reflect_plugin import ReflectContext, ReflectTarget, normalize_key, run_reflect_round, union_ranges
    from .runtime.engine import run_slot_vlm
except ImportError:
    import my_pipeline_legacy as legacy  # type: ignore
    from pipeline_input_adapter import ensure_pipeline_xlsx, write_plan_json  # type: ignore
    from pipeline_agent_types import ControllerState, PipelinePaths, ReflectRegionPlan, ReflectStageResult  # type: ignore
    from reflect_plugin import ReflectContext, ReflectTarget, normalize_key, run_reflect_round, union_ranges  # type: ignore
    from runtime.engine import run_slot_vlm  # type: ignore


PROJECT_ROOT = Path(__file__).resolve().parents[2]
AGENT_SKILLS_DIR = PROJECT_ROOT / "agent_skills"


def create_usage_summary() -> dict[str, Any]:
    return {
        "first_pass_agent": {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0},
        "reflect_checker": {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0},
        "reflect_model_assess": {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0},
        "patch_agents": [],
        "reflect_plugin_stats": {
            "rounds_started": 0,
            "stage1_need_fix_count": 0,
            "stage1_parse_failed_count": 0,
            "assessed_cells_total": 0,
            "edited_cells_total": 0,
            "scan_only_cells_total": 0,
            "high_risk_cells_total": 0,
            "targets_total": 0,
            "blank_targets_total": 0,
            "patch_jobs_total": 0,
            "patch_written_actual_edits_total": 0,
        },
    }


def load_skills_text(skills_inline: str, skills_file: str) -> str:
    parts: list[str] = []
    inline = (skills_inline or "").strip()
    if inline:
        items = [x.strip() for x in inline.split(",") if x.strip()]
        if items:
            parts.append("Inline skills:")
            parts.extend([f"- {x}" for x in items])
    fp = (skills_file or "").strip()
    if fp:
        p = Path(fp)
        if not p.is_absolute():
            preferred = AGENT_SKILLS_DIR / p
            p = preferred if preferred.exists() else (Path.cwd() / p)
        if not p.exists():
            raise legacy.PipelineError(f"agent_skills_file not found: {p}")
        raw = p.read_text(encoding="utf-8").strip()
        if raw:
            try:
                obj: Any = json.loads(raw)
                if isinstance(obj, dict):
                    parts.append("File skills (json object):")
                    parts.append(json.dumps(obj, ensure_ascii=False, indent=2))
                elif isinstance(obj, list):
                    parts.append("File skills (json list):")
                    parts.extend([f"- {it}" for it in obj])
                else:
                    parts.append("File skills:")
                    parts.append(str(obj))
            except Exception:
                parts.append("File skills:")
                parts.append(raw)
    return "\n".join(parts)


def _short_json(value: Any) -> str:
    try:
        return json.dumps(value, ensure_ascii=False)
    except Exception:
        return json.dumps(str(value), ensure_ascii=False)


def _looks_like_prompt_too_large_error(exc: Exception) -> bool:
    text = str(exc or "").lower()
    markers = (
        "maximum context",
        "context length",
        "prompt is too long",
        "prompt too long",
        "too many tokens",
        "token limit",
        "max prompt",
        "request too large",
        "context window",
    )
    return any(marker in text for marker in markers)


def _agent_fill_with_large_sheet_retry(
    *,
    log_fn: Callable[[str], None],
    stage_label: str,
    agent_fill_kwargs: dict[str, Any],
) -> dict[str, Any]:
    try:
        return legacy.agent_fill(**agent_fill_kwargs)
    except Exception as exc:
        if not _looks_like_prompt_too_large_error(exc):
            raise
        log_fn(
            f"[{stage_label}] prompt too large for model context; retrying with compact tool-first instruction "
            "that requires scan_structure before any local reads/writes."
        )
        compact_instruction = (
            f"{str(agent_fill_kwargs.get('instruction') or '').strip()}\n\n"
            "[Large workbook fallback]\n"
            "The workbook is too large to inline full table context in the prompt.\n"
            "Start by calling scan_structure() exactly once.\n"
            "Then use read_range(include_style=true) only on small local regions that are directly relevant.\n"
            "Do not request or assume whole-sheet dumps.\n"
            "Rely on workbook tools for table understanding before any write_cell() call.\n"
            "If target mapping stays ambiguous after local inspection, leave that field blank.\n"
        ).strip()
        retry_kwargs = dict(agent_fill_kwargs)
        retry_kwargs["instruction"] = compact_instruction
        retry_kwargs["rendered_html_path"] = None
        retry_kwargs["rendered_html_max_chars"] = 0
        # Use the leanest prompt path: no pre-inlined structure/html, force tool-first reasoning.
        retry_kwargs["form_context_mode"] = "html"
        retry_kwargs["context_summary"] = None
        retry_kwargs["compare_before_after_pngs"] = None
        return legacy.agent_fill(**retry_kwargs)


def _iter_sheet_data_validations(ws: Any) -> list[Any]:
    dvs = getattr(ws, "data_validations", None)
    items = getattr(dvs, "dataValidation", None)
    if not items:
        return []
    return list(items)


def _parse_validation_formula_options(wb: Any, ws: Any, formula: str, limit: int = 50) -> list[str]:
    raw = str(formula or "").strip()
    if not raw:
        return []
    if raw.startswith("="):
        raw = raw[1:]
    raw = raw.strip()
    if raw.startswith('"') and raw.endswith('"'):
        return [item.strip() for item in raw.strip('"').split(",") if item.strip()][:limit]

    ref_sheet = ws.title
    ref_expr = raw
    if "!" in raw:
        sheet_part, ref_part = raw.split("!", 1)
        ref_sheet = sheet_part.strip("'")
        ref_expr = ref_part

    ref_expr = ref_expr.replace("$", "").strip()
    values: list[str] = []
    try:
        if ":" in ref_expr:
            ref_ws = wb[ref_sheet]
            for row in ref_ws[ref_expr]:
                row_cells = row if isinstance(row, tuple) else (row,)
                for cell in row_cells:
                    text = str(cell.value or "").strip()
                    if text:
                        values.append(text)
                        if len(values) >= limit:
                            return values[:limit]
            return values[:limit]
    except Exception:
        pass

    try:
        defined_name = wb.defined_names.get(ref_expr)
    except Exception:
        defined_name = None
    if defined_name is not None:
        try:
            destinations = list(defined_name.destinations)
        except Exception:
            destinations = []
        for sheet_name, coord in destinations:
            try:
                ref_ws = wb[sheet_name]
                for row in ref_ws[coord]:
                    row_cells = row if isinstance(row, tuple) else (row,)
                    for cell in row_cells:
                        text = str(cell.value or "").strip()
                        if text:
                            values.append(text)
                            if len(values) >= limit:
                                return values[:limit]
            except Exception:
                continue
    return values[:limit]


def _extract_dropdown_constraints(xlsx_path: Path, sheet_name: str | None) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    items: list[dict[str, Any]] = []
    try:
        for dv in _iter_sheet_data_validations(ws):
            if str(getattr(dv, "type", "")).strip().lower() != "list":
                continue
            formula1 = str(getattr(dv, "formula1", "") or "").strip()
            options = _parse_validation_formula_options(wb, ws, formula1)
            if not options:
                continue
            for target_range in [seg.strip() for seg in str(getattr(dv, "sqref", "") or "").split() if seg.strip()]:
                items.append(
                    {
                        "target_range": target_range,
                        "sheet": ws.title,
                        "options": options[:20],
                        "formula1": formula1,
                    }
                )
                if len(items) >= 200:
                    return items
    finally:
        wb.close()
    return items


def _build_dropdown_context(dropdowns: list[dict[str, Any]], focus_range: str | None = None) -> str:
    filtered = []
    for item in dropdowns or []:
        target = str(item.get("target_range") or "").strip()
        if not target:
            continue
        if focus_range:
            try:
                if not legacy._range_intersects(target, focus_range):
                    continue
            except Exception:
                continue
        filtered.append(
            {
                "target_range": target,
                "options": list(item.get("options") or [])[:20],
            }
        )
        if len(filtered) >= 80:
            break
    if not filtered:
        return ""
    return (
        "Dropdown validation constraints from the Excel template:\n"
        f"{json.dumps(filtered, ensure_ascii=False, indent=2)}\n"
        "When filling cells inside these ranges, choose only from the listed options.\n"
    )


def _log_screenshot_check_details(
    log_fn: Callable[[str], None],
    stage_name: str,
    round_idx: int,
    screenshot_check: dict[str, Any],
) -> None:
    need_fix = bool(screenshot_check.get("need_fix"))
    parse_failed = bool(screenshot_check.get("bad_range_parse_failed"))
    confidence = screenshot_check.get("confidence")
    bad_ranges = [str(x) for x in (screenshot_check.get("bad_ranges", []) or []) if str(x).strip()]
    correct_ranges = [str(x) for x in (screenshot_check.get("correct_ranges", []) or []) if str(x).strip()]
    heuristic_findings = screenshot_check.get("heuristic_findings", []) or []
    reason = str(screenshot_check.get("reason") or "").strip()
    usage = screenshot_check.get("usage") if isinstance(screenshot_check.get("usage"), dict) else {}
    log_fn(
        f"[{stage_name.upper()}][SCREENSHOT] round={round_idx} "
        f"need_fix={need_fix} parse_failed={parse_failed} confidence={confidence} "
        f"bad_ranges={len(bad_ranges)} correct_ranges={len(correct_ranges)} "
        f"heuristic_findings={len(heuristic_findings)} "
        f"tokens={usage.get('total_tokens', 0)}"
    )
    if reason:
        log_fn(f"[{stage_name.upper()}][SCREENSHOT][REASON] round={round_idx} reason={_short_json(reason)}")
    if bad_ranges:
        log_fn(f"[{stage_name.upper()}][SCREENSHOT][BAD_RANGES] round={round_idx} ranges={_short_json(bad_ranges)}")
    if correct_ranges:
        preview = correct_ranges[:20]
        suffix = " ..." if len(correct_ranges) > len(preview) else ""
        log_fn(
            f"[{stage_name.upper()}][SCREENSHOT][CORRECT_RANGES] round={round_idx} "
            f"ranges={_short_json(preview)}{suffix}"
        )
    if heuristic_findings:
        log_fn(
            f"[{stage_name.upper()}][SCREENSHOT][HEURISTICS] round={round_idx} "
            f"findings={_short_json(heuristic_findings[:10])}"
        )


def _log_reflect_result_details(
    log_fn: Callable[[str], None],
    stage_name: str,
    round_idx: int,
    reflect_payload: dict[str, Any],
    screenshot_check: dict[str, Any],
) -> None:
    score_flag_order = [
        ("label_area", "fill_in_label_area"),
        ("duplicate_pattern", "duplicate_pattern"),
        ("template_inconsistency", "template_inconsistency"),
        ("high_risk_pattern", "high_risk_pattern"),
        ("special_char_anomaly", "special_char_anomaly"),
    ]
    assessed_cells = reflect_payload.get("edits_considered", []) if isinstance(reflect_payload.get("edits_considered"), list) else []
    risk_assessments = reflect_payload.get("risk_assessments", []) if isinstance(reflect_payload.get("risk_assessments"), list) else []
    reflect_targets = reflect_payload.get("reflect_targets", []) if isinstance(reflect_payload.get("reflect_targets"), list) else []
    patch_jobs = reflect_payload.get("patch_jobs", []) if isinstance(reflect_payload.get("patch_jobs"), list) else []
    high_risk_cells = [x for x in risk_assessments if isinstance(x, dict) and bool(x.get("is_high_risk"))]
    correct_cells = [x for x in risk_assessments if isinstance(x, dict) and bool(x.get("is_correct_region"))]
    safe_blank_cells = [x for x in risk_assessments if isinstance(x, dict) and bool(x.get("is_safe_blank"))]
    edited_cells = [x for x in assessed_cells if isinstance(x, dict) and bool(x.get("is_actual_edit"))]
    scan_only_cells = [x for x in assessed_cells if isinstance(x, dict) and not bool(x.get("is_actual_edit"))]
    missing_keys = [str(x) for x in (reflect_payload.get("missing_keys", []) or []) if str(x).strip()]
    correct_ranges = [str(x) for x in (reflect_payload.get("correct_ranges", []) or []) if str(x).strip()]
    log_fn(
        f"[{stage_name.upper()}][SUMMARY] round={round_idx} "
        f"stage1_need_fix={bool(screenshot_check.get('need_fix'))} "
        f"parse_failed={bool(screenshot_check.get('bad_range_parse_failed'))} "
        f"assessed={len(assessed_cells)} edited={len(edited_cells)} scan_only={len(scan_only_cells)} "
        f"correct={len(correct_cells)} safe_blank={len(safe_blank_cells)} high_risk={len(high_risk_cells)} "
        f"targets={len(reflect_targets)} patch_jobs={len(patch_jobs)} "
        f"missing_keys={len(missing_keys)} correct_ranges={len(correct_ranges)}"
    )
    if missing_keys:
        log_fn(f"[{stage_name.upper()}][MISSING_KEYS] round={round_idx} keys={_short_json(missing_keys)}")
    if patch_jobs:
        patch_regions = [str(job.get('region') or '').strip() for job in patch_jobs if str(job.get('region') or '').strip()]
        if patch_regions:
            log_fn(f"[{stage_name.upper()}][PATCH_REGIONS] round={round_idx} regions={_short_json(patch_regions)}")
    for assess in risk_assessments:
        if not isinstance(assess, dict):
            continue
        flags = {str(x) for x in (assess.get("risk_flags", []) or []) if str(x).strip()}
        flag_bits = {
            label: (flag in flags)
            for flag, label in score_flag_order
        }
        score = int(assess.get("score", 0) or 0)
        triggered = [label for _, label in score_flag_order if flag_bits[label]]
        log_fn(
            f"[{stage_name.upper()}][RISK_SCORE] round={round_idx} "
            f"cell={assess.get('target_cell')} "
            f"key={_short_json(assess.get('key', ''))} "
            f"score={score}/5 "
            f"fill_in_label_area={flag_bits['fill_in_label_area']} "
            f"duplicate_pattern={flag_bits['duplicate_pattern']} "
            f"template_inconsistency={flag_bits['template_inconsistency']} "
            f"high_risk_pattern={flag_bits['high_risk_pattern']} "
            f"special_char_anomaly={flag_bits['special_char_anomaly']} "
            f"triggered={_short_json(triggered)} "
            f"high_risk={assess.get('is_high_risk', False)}"
        )
        log_fn(
            f"[{stage_name.upper()}][ASSESS] round={round_idx} "
            f"cell={assess.get('target_cell')} "
            f"key={_short_json(assess.get('key', ''))} "
            f"score={score} "
            f"correct={assess.get('is_correct_region', False)} "
            f"safe_blank={assess.get('is_safe_blank', False)} "
            f"high_risk={assess.get('is_high_risk', False)} "
            f"flags={_short_json(sorted(flags))} "
            f"reason={_short_json(assess.get('reason', ''))}"
        )


def build_runtime_paths(args: Any, log_fn: Callable[[str], None]) -> PipelinePaths:
    job_dir = (
        Path(args.runtime_job_dir).resolve()
        if str(getattr(args, "runtime_job_dir", "") or "").strip()
        else (Path.cwd() / "runs" / "agent_pipeline")
    )
    job_dir.mkdir(parents=True, exist_ok=True)
    edit_log_jsonl = Path(args.reflect_edit_log_jsonl).resolve() if args.reflect_edit_log_jsonl else (job_dir / "edit_log.jsonl")
    model_io_dir = (
        Path(args.model_io_dump_dir).resolve()
        if str(args.model_io_dump_dir).strip()
        else (job_dir / "model_io" if args.save_model_io else None)
    )
    if model_io_dir is not None:
        model_io_dir.mkdir(parents=True, exist_ok=True)
    controller_debug_dir = job_dir / "controller"
    controller_debug_dir.mkdir(parents=True, exist_ok=True)
    source_input = Path(args.input).resolve()
    input_adapter_dir = job_dir / "input_adapter"
    xlsx_path, input_plan = ensure_pipeline_xlsx(
        src=source_input,
        dst=job_dir / "input.xlsx",
        work_dir=input_adapter_dir,
        log_fn=log_fn,
        excel_ensure_fn=legacy.ensure_xlsx,
    )
    input_plan_json = write_plan_json(input_plan, input_adapter_dir / "input_plan.json")
    before_fill_xlsx = job_dir / "before_fill.xlsx"
    shutil.copyfile(xlsx_path, before_fill_xlsx)
    return PipelinePaths(
        job_dir=job_dir,
        source_input=source_input,
        input_xlsx=xlsx_path,
        before_fill_xlsx=before_fill_xlsx,
        output_xlsx=Path(args.output),
        edit_log_jsonl=edit_log_jsonl,
        controller_debug_dir=controller_debug_dir,
        input_plan_json=input_plan_json,
        model_io_dir=model_io_dir,
    )


def load_initial_form(args: Any, paths: PipelinePaths, log_fn: Callable[[str], None]) -> dict[str, Any]:
    form: dict[str, Any] = {"pairs": [], "invalid_slots": []}
    try:
        dropdown_constraints = _extract_dropdown_constraints(paths.input_xlsx, args.sheet)
        form["dropdown_constraints"] = dropdown_constraints
        if dropdown_constraints:
            log_fn(f"[INFO] extracted dropdown constraints: {len(dropdown_constraints)}")
    except Exception as exc:
        log_fn(f"[WARN] failed to extract dropdown constraints: {exc}")
    if args.external_qwen_pairs_json:
        qwen_pairs_path = Path(args.external_qwen_pairs_json).resolve()
        if not qwen_pairs_path.exists():
            raise legacy.PipelineError(f"external_qwen_pairs_json not found: {qwen_pairs_path}")
        log_fn(f"[INFO] Using external qwen pairs: {qwen_pairs_path}")
        form = legacy.build_form_from_qwen(qwen_pairs_path, xlsx_path=paths.input_xlsx, sheet_name=args.sheet)
    else:
        log_fn("[INFO] initial global slot hints disabled; free_fill will run without precomputed field mappings.")
    if isinstance(form, dict) and isinstance(form.get("dropped_suspicious_pairs"), list) and form.get("dropped_suspicious_pairs"):
        log_fn(f"[INFO] dropped suspicious qwen pairs: {len(form.get('dropped_suspicious_pairs') or [])}")
    if isinstance(form, dict) and isinstance(form.get("invalid_slots"), list) and form.get("invalid_slots"):
        log_fn(f"[INFO] qwen invalid slots available for reflect/repair: {len(form.get('invalid_slots') or [])}")
    return form


def load_planner_policy(
    args: Any,
    paths: PipelinePaths,
    log_fn: Callable[[str], None],
) -> tuple[dict[str, Any], dict[str, Any], Path | None]:
    planner_policy: dict[str, Any] = {}
    planner_debug_root: Path | None = None
    if str(args.planner_json_path).strip():
        planner_policy = legacy._load_planner_policy(args.planner_json_path)
        planner_debug_root = Path(args.planner_json_path).resolve().parent
    elif args.agentic_planner_enable:
        planner_model = str(args.planner_model or args.agent_model).strip()
        planner_base_url = str(args.planner_base_url or args.agent_base_url).strip() or None
        planner_api_key = str(args.planner_api_key or args.agent_api_key).strip() or None
        planner_debug_dir = (
            Path(args.planner_debug_dir).resolve()
            if str(args.planner_debug_dir).strip()
            else (paths.job_dir / "planner")
        )
        planner_debug_root = planner_debug_dir
        planner_debug_dir.mkdir(parents=True, exist_ok=True)
        planner_struct_summary = legacy._scan_structure_for_planner(paths.input_xlsx, args.sheet)
        try:
            planner_policy = legacy._request_planner_policy(
                instruction=args.instruction,
                struct_summary=planner_struct_summary,
                sheet_name=args.sheet,
                model=planner_model,
                base_url=planner_base_url,
                api_key=planner_api_key,
                temperature=float(args.planner_temperature),
                model_io_dir=(paths.model_io_dir / "planner" if paths.model_io_dir is not None else None),
            )
            planner_failed = False
        except Exception as exc:
            planner_failed = True
            planner_policy = {
                "form_type": "generic_form",
                "global_policies": {
                    "leave_unspecified_blank": True,
                    "preserve_existing_content": True,
                    "avoid_guessing": True,
                },
                "field_groups": [],
                "blank_policies": [],
                "do_not_fill_sections": [],
                "checkbox_policies": [],
                "repeated_block_policy": {
                    "has_repeated_blocks": False,
                    "block_key": "",
                    "expected_count": None,
                },
                "high_confidence_fields": [],
                "uncertain_fields": [],
                "notes": [f"planner_failed: {exc}"],
            }
        planner_payload = {k: v for k, v in planner_policy.items() if not str(k).startswith("_")}
        (planner_debug_dir / "struct_summary.json").write_text(
            json.dumps(planner_struct_summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (planner_debug_dir / "plan.json").write_text(
            json.dumps(planner_policy, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (planner_debug_dir / "plan.normalized.json").write_text(
            json.dumps(planner_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (planner_debug_dir / "planner_policy_preview.txt").write_text(
            legacy._build_planner_policy_preview(planner_policy),
            encoding="utf-8",
        )
        (planner_debug_dir / "instruction_original.txt").write_text(
            str(getattr(args, "original_instruction", args.instruction) or ""),
            encoding="utf-8",
        )
        (planner_debug_dir / "instruction_effective.txt").write_text(
            str(args.instruction or ""),
            encoding="utf-8",
        )
        log_fn(
            f"[INFO] planner_status={'failed' if planner_failed else 'ok'} "
            f"plan={planner_debug_dir / 'plan.normalized.json'}"
        )
    workflow_policy = legacy._planner_workflow_policy(planner_policy)
    return planner_policy, workflow_policy, planner_debug_root


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def _coord_of(a1: str) -> tuple[int, int]:
    col_s, row = coordinate_from_string(a1)
    return int(row), int(column_index_from_string(col_s))


def _range_cells_set(range_list: list[str]) -> set[str]:
    cells: set[str] = set()
    for item in range_list:
        text = str(item or "").strip()
        if not text:
            continue
        try:
            cells.update(legacy._iter_range_cells(text))
        except Exception:
            continue
    return {str(cell).upper() for cell in cells}


def _border_tags(cell) -> list[str]:
    border = cell.border
    tags: list[str] = []
    if border is None:
        return tags
    if getattr(getattr(border, "top", None), "style", None):
        tags.append("T")
    if getattr(getattr(border, "bottom", None), "style", None):
        tags.append("B")
    if getattr(getattr(border, "left", None), "style", None):
        tags.append("L")
    if getattr(getattr(border, "right", None), "style", None):
        tags.append("R")
    return tags


def _scan_template_structure(xlsx_path: Path, sheet_name: str | None) -> dict[str, set[str]]:
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active
        min_row, min_col, max_row, max_col = legacy._get_used_bbox(ws)
        merged = legacy._merged_anchor_map(ws)
        nonempty_cells: set[str] = set()
        input_candidate_cells: set[str] = set()
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                a1 = f"{get_column_letter(c)}{r}".upper()
                if cell.value is not None and str(cell.value).strip() != "":
                    nonempty_cells.add(a1)
                tags = _border_tags(cell)
                is_merged = (r, c) in merged
                if not tags and not is_merged:
                    continue
                is_boxed = set(tags) >= {"T", "B", "L", "R"}
                has_bottom = "B" in tags
                if is_boxed or has_bottom or is_merged:
                    anchor = merged.get((r, c), (r, c, r, c))
                    input_candidate_cells.add(f"{get_column_letter(anchor[1])}{anchor[0]}".upper())
        static_template_cells = {cell for cell in nonempty_cells if cell not in input_candidate_cells}
        return {
            "nonempty_cells": nonempty_cells,
            "input_candidate_cells": input_candidate_cells,
            "static_template_cells": static_template_cells,
        }
    finally:
        wb.close()


def inspect_template_region_damage(
    template_xlsx: Path,
    current_xlsx: Path,
    sheet_name: str | None,
    region: str,
    preserve_cells: set[str] | None = None,
) -> dict[str, Any]:
    preserve = {str(x).upper() for x in (preserve_cells or set())}
    structure = _scan_template_structure(template_xlsx, sheet_name)
    static_template_cells = structure["static_template_cells"]
    src_wb = openpyxl.load_workbook(template_xlsx)
    dst_wb = openpyxl.load_workbook(current_xlsx)
    try:
        src_ws = src_wb[sheet_name] if (sheet_name and sheet_name in src_wb.sheetnames) else src_wb.active
        dst_ws = dst_wb[sheet_name] if (sheet_name and sheet_name in dst_wb.sheetnames) else dst_wb.active
        src_merged = legacy._merged_anchor_map(src_ws)
        dst_merged = legacy._merged_anchor_map(dst_ws)
        changed_cells: list[dict[str, Any]] = []
        seen: set[str] = set()
        for cell_ref in legacy._iter_range_cells(region):
            row_idx, col_idx = _coord_of(cell_ref)
            src_anchor = src_merged.get((row_idx, col_idx), (row_idx, col_idx, row_idx, col_idx))
            dst_anchor = dst_merged.get((row_idx, col_idx), (row_idx, col_idx, row_idx, col_idx))
            src_anchor_a1 = f"{get_column_letter(src_anchor[1])}{src_anchor[0]}".upper()
            dst_anchor_a1 = f"{get_column_letter(dst_anchor[1])}{dst_anchor[0]}".upper()
            if src_anchor_a1 in seen or src_anchor_a1 in preserve or src_anchor_a1 not in static_template_cells:
                continue
            seen.add(src_anchor_a1)
            src_cell = src_ws[src_anchor_a1]
            dst_cell = dst_ws[dst_anchor_a1]
            reasons: list[str] = []
            if src_cell.value != dst_cell.value:
                reasons.append("template_static_value_changed")
            if src_cell._style != dst_cell._style:
                reasons.append("template_static_style_changed")
            if tuple(src_anchor) != tuple(dst_anchor):
                reasons.append("template_merge_changed")
            if reasons:
                changed_cells.append(
                    {
                        "cell": src_anchor_a1,
                        "current_cell": dst_anchor_a1,
                        "template_value": src_cell.value,
                        "current_value": dst_cell.value,
                        "reasons": reasons,
                    }
                )
        restore_region = union_ranges([str(item["cell"]) for item in changed_cells]) if changed_cells else ""
        return {
            "need_restore": bool(changed_cells),
            "restore_region": restore_region,
            "changed_cells": changed_cells,
            "reasons": sorted({reason for item in changed_cells for reason in item.get("reasons", [])}),
        }
    finally:
        src_wb.close()
        dst_wb.close()


def restore_template_region(
    template_xlsx: Path,
    target_xlsx: Path,
    sheet_name: str | None,
    region: str,
    preserve_cells: set[str] | None = None,
) -> dict[str, Any]:
    preserve = {str(x).upper() for x in (preserve_cells or set())}
    src_wb = openpyxl.load_workbook(template_xlsx)
    dst_wb = openpyxl.load_workbook(target_xlsx)
    try:
        src_ws = src_wb[sheet_name] if (sheet_name and sheet_name in src_wb.sheetnames) else src_wb.active
        dst_ws = dst_wb[sheet_name] if (sheet_name and sheet_name in dst_wb.sheetnames) else dst_wb.active
        src_merged = legacy._merged_anchor_map(src_ws)
        dst_merged = legacy._merged_anchor_map(dst_ws)
        touched: set[str] = set()
        restored_cells: list[str] = []
        for cell_ref in legacy._iter_range_cells(region):
            row_idx, col_idx = _coord_of(cell_ref)
            src_anchor = src_merged.get((row_idx, col_idx), (row_idx, col_idx, row_idx, col_idx))
            dst_anchor = dst_merged.get((row_idx, col_idx), (row_idx, col_idx, row_idx, col_idx))
            src_anchor_a1 = f"{get_column_letter(src_anchor[1])}{src_anchor[0]}".upper()
            dst_anchor_a1 = f"{get_column_letter(dst_anchor[1])}{dst_anchor[0]}".upper()
            if dst_anchor_a1 in preserve or dst_anchor_a1 in touched:
                continue
            touched.add(dst_anchor_a1)
            src_cell = src_ws[src_anchor_a1]
            dst_cell = dst_ws[dst_anchor_a1]
            if src_cell.value != dst_cell.value or src_cell._style != dst_cell._style:
                restored_cells.append(dst_anchor_a1)
            dst_cell.value = src_cell.value
            dst_cell._style = copy(src_cell._style)
            dst_cell.number_format = src_cell.number_format
            dst_cell.alignment = copy(src_cell.alignment)
            dst_cell.font = copy(src_cell.font)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.border = copy(src_cell.border)
            dst_cell.protection = copy(src_cell.protection)
        if restored_cells:
            dst_wb.save(target_xlsx)
        return {"region": region, "restored_count": len(restored_cells), "restored_cells": restored_cells}
    finally:
        src_wb.close()
        dst_wb.close()


def _summarize_template_region(template_xlsx: Path, sheet_name: str | None, region: str, limit: int = 40) -> dict[str, Any]:
    wb = openpyxl.load_workbook(template_xlsx, read_only=True)
    try:
        ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active
        items: list[dict[str, Any]] = []
        seen: set[str] = set()
        for cell_ref in legacy._iter_range_cells(region):
            anchor = str(cell_ref).upper()
            if anchor in seen:
                continue
            seen.add(anchor)
            value = ws[anchor].value
            if value is None or str(value).strip() == "":
                continue
            items.append({"cell": anchor, "value": value})
            if len(items) >= max(1, int(limit)):
                break
        return {"region": region, "nonempty_template_cells": items}
    finally:
        wb.close()


def _is_noisy_reflect_key(key: str) -> bool:
    text = str(key or "").strip()
    if not text:
        return True
    compact = re.sub(r"\s+", " ", text)
    if re.fullmatch(r"[\d\W_]+", compact):
        return True
    alnum = re.sub(r"[^A-Za-z0-9]+", "", compact)
    if len(alnum) <= 1:
        return True
    return bool(re.fullmatch(r"\d+[A-Za-z]?", alnum))


def _keys_soft_match(a: str, b: str) -> bool:
    na = normalize_key(a)
    nb = normalize_key(b)
    if not na or not nb:
        return False
    if na == nb or na in nb or nb in na:
        return True
    ta = set(na.split())
    tb = set(nb.split())
    if not ta or not tb:
        return False
    overlap = len(ta & tb) / max(1, min(len(ta), len(tb)))
    return overlap >= 0.75


def _hint_matches_target(key: str, candidate_value_cell: str, region_targets: list[ReflectTarget]) -> bool:
    for target in region_targets:
        target_key = str(getattr(target, "key", "") or "")
        if _keys_soft_match(key, target_key):
            return True
        try:
            if candidate_value_cell and legacy._range_intersects(target.target_range, candidate_value_cell):
                return True
        except Exception:
            continue
    return False


def _select_patch_targets(patch_job: dict[str, Any]) -> list[dict[str, Any]]:
    all_targets = [x for x in (patch_job.get("targets", []) or []) if isinstance(x, dict)]
    prioritized = [
        x
        for x in all_targets
        if str(x.get("kind") or "") in {"region_blank_candidate", "screenshot_bad_range", "label_area_edit"}
    ]
    if not prioritized:
        return all_targets
    selected: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for item in prioritized + all_targets:
        sig = (
            str(item.get("kind") or ""),
            str(item.get("key") or ""),
            str(item.get("target_range") or ""),
        )
        if sig in seen:
            continue
        seen.add(sig)
        selected.append(item)
    return selected


def _select_patch_region_hints(region_hints: list[dict[str, Any]], patch_targets: list[dict[str, Any]]) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    target_keys = {
        normalize_key(str(x.get("key") or ""))
        for x in patch_targets
        if str(x.get("key") or "").strip()
    }
    for hint in region_hints:
        if not isinstance(hint, dict):
            continue
        key = str(hint.get("key") or "").strip()
        value_cell = str(hint.get("candidate_value_cell") or "").strip()
        norm_key = normalize_key(key)
        matched = bool(norm_key and norm_key in target_keys)
        if not matched:
            for target in patch_targets:
                try:
                    if value_cell and legacy._range_intersects(str(target.get("target_range") or ""), value_cell):
                        matched = True
                        break
                except Exception:
                    continue
        if matched:
            selected.append(hint)
    return selected


def _build_reflect_qwen_prompt_context(region: str, region_targets: list[ReflectTarget]) -> str:
    flag_order = [
        "label_area",
        "duplicate_pattern",
        "template_inconsistency",
        "high_risk_pattern",
        "special_char_anomaly",
    ]
    region_flags: list[str] = []
    for target in region_targets:
        for flag in getattr(target, "risk_flags", []) or []:
            if flag in flag_order and flag not in region_flags:
                region_flags.append(flag)
    if not region_flags:
        return ""
    flag_notes = {
        "label_area": "Do not treat title/header/label text as the value key for a slot.",
        "duplicate_pattern": "This region may belong to repeated template blocks. Stay inside the current local block.",
        "template_inconsistency": "Prefer the key/value pattern that matches neighboring fields in this same section.",
        "high_risk_pattern": "Be careful about vertical misalignment between a key row and the value row.",
        "special_char_anomaly": "Preserve accented or language-specific characters when matching the key text.",
    }
    parts = [
        f"Reflect risk context for region {region}: {', '.join(region_flags)}.",
        "Use these flags as constraints while deciding the correct key for each slot.",
    ]
    for flag in region_flags:
        parts.append(f"{flag}: {flag_notes[flag]}")
    return "\n".join(parts)


def _build_reflect_hint_builder(
    state: ControllerState,
    args: Any,
    log_fn: Callable[[str], None],
) -> Callable[[str, list[ReflectTarget], Path], list[dict[str, Any]]]:
    def build_reflect_region_hints(region: str, region_targets: list[ReflectTarget], out_dir: Path) -> list[dict[str, Any]]:
        out_dir.mkdir(parents=True, exist_ok=True)
        try:
            prep = legacy.preprocess_excel_html_css(
                xlsx_path=state.paths.before_fill_xlsx,
                out_dir=out_dir / "preprocess",
                sheet_name=args.sheet,
                dpi=args.dpi,
                canvas_bg=args.render_canvas_bg,
                png_scale=args.render_png_scale,
                max_render_rows=args.render_max_rows,
                max_render_cols=args.render_max_cols,
                focus_range=region,
                focus_pad_left_cols=3,
                focus_pad_top_rows=1,
                focus_pad_right_cols=3,
                focus_pad_bottom_rows=1,
            )
            qwen_out_dir = out_dir / "qwen_out"
            qwen_outputs = legacy.run_qwen_small(
                image_path=prep["sheet_png"],
                xlsx_path=state.paths.before_fill_xlsx,
                out_dir=qwen_out_dir,
                sheet_name=args.sheet,
                ckpt_path=args.ckpt or None,
                edges_json=prep["edges_json"],
                bounds_json=prep["bounds_json"],
                qwen_base_url=None,
                qwen_model=None,
                qwen_api_key=None,
                skip_qwen=True,
                log_fn=log_fn,
            )
            prompt_context_text = _build_reflect_qwen_prompt_context(region, region_targets)
            reflect_slots_json = qwen_out_dir / "slots_in_region.json"
            kept_slots = legacy._filter_slots_json_by_focus_range(
                slots_json_path=qwen_outputs["slots_json"],
                focus_range=region,
                out_json_path=reflect_slots_json,
            )
            if kept_slots <= 0:
                log_fn(f"[REFLECT_PLUGIN] region={region} kept_slots=0 after focus-range filtering; no local slot hints.")
                return []
            reflect_pairs_path = run_slot_vlm(
                image_path=prep["sheet_png"],
                slots_json=reflect_slots_json,
                out_dir=qwen_out_dir,
                model_name=args.agent_model,
                base_url=args.agent_base_url or None,
                api_key=args.agent_api_key or None,
                temperature=float(args.agent_temperature),
                prompt_context_text=prompt_context_text,
                batch_size=8,
                log_fn=log_fn,
                debug_dir=qwen_out_dir / "vlm_debug",
            )
            region_form_all = legacy.build_form_from_qwen(reflect_pairs_path, xlsx_path=state.paths.output_xlsx, sheet_name=args.sheet)
            region_form = legacy._filter_form_pairs_by_focus_range(region_form_all, region)
            patch_target_payload = [
                {
                    "key": str(getattr(target, "key", "") or ""),
                    "target_range": str(getattr(target, "target_range", "") or ""),
                    "source_cell": str(getattr(target, "source_cell", "") or ""),
                }
                for target in region_targets
            ]
            qwen_context_summary = _build_qwen_context_summary_from_region_form(
                region=region,
                patch_targets=patch_target_payload,
                region_form=region_form,
                current_xlsx=state.paths.output_xlsx,
                sheet_name=args.sheet,
            )
            hints: list[dict[str, Any]] = []
            for pair in region_form.get("pairs", []):
                if not isinstance(pair, dict):
                    continue
                key = str(pair.get("key") or "").strip()
                value_cell = str(pair.get("value_cell") or "").strip()
                if not key or not value_cell:
                    continue
                candidate_value = legacy._read_anchor_value_from_workbook(state.paths.output_xlsx, args.sheet, value_cell)
                is_blank_candidate = legacy._is_blankish_value(candidate_value)
                if is_blank_candidate and _is_noisy_reflect_key(key):
                    continue
                if not _hint_matches_target(key, value_cell, region_targets):
                    continue
                hints.append(
                    {
                        "key": key,
                        "target_range": region,
                        "candidate_value_cell": value_cell,
                        "candidate_value": candidate_value,
                        "is_blank_candidate": is_blank_candidate,
                        "hint_source": "reflect_blank_slot" if is_blank_candidate else "reflect_slot_kv",
                    }
                )
            for invalid in region_form.get("invalid_slots", []):
                if not isinstance(invalid, dict):
                    continue
                value_cell = str(invalid.get("value_cell") or invalid.get("candidate_value_cell") or "").strip()
                if not value_cell or not _hint_matches_target(str(invalid.get("key") or ""), value_cell, region_targets):
                    continue
                hint = dict(invalid)
                hint["candidate_value_cell"] = value_cell
                hint["avoid_cell"] = True
                hint["invalid_slot"] = True
                hint["target_range"] = region
                hint.setdefault("hint_source", "reflect_invalid_slot")
                hints.append(hint)
            if qwen_context_summary:
                _write_json(qwen_out_dir / "region_form_summary.json", qwen_context_summary)
                hints.append(
                    {
                        "hint_source": "qwen_context_summary",
                        "context_only": True,
                        "target_range": region,
                        "qwen_context_summary": qwen_context_summary,
                    }
                )
            return hints
        except Exception as exc:
            log_fn(f"[REFLECT_PLUGIN] failed to build region hints for region={region}: {exc}")
            return []

    return build_reflect_region_hints


def _build_protected_write_cells(
    region_plan: ReflectRegionPlan,
    reflect_stage: ReflectStageResult,
    before_nonempty_cells: dict[str, str],
    editable_cells: set[str] | None = None,
    current_region_nonempty_cells: dict[str, str] | None = None,
) -> dict[str, str]:
    protected: dict[str, str] = {}
    target_cells: set[str] = {str(x).upper() for x in (editable_cells or set()) if str(x).strip()}
    if not target_cells:
        for target in region_plan.patch_targets:
            cell = str(target.get("source_cell") or target.get("target_range") or "").strip().upper()
            if cell:
                try:
                    target_cells.update(legacy._iter_range_cells(cell))
                except Exception:
                    target_cells.add(cell)
    for cell, value in before_nonempty_cells.items():
        if cell not in target_cells:
            protected[cell] = f"template/original content: {str(value)[:80]}"
    for cell, value in (current_region_nonempty_cells or {}).items():
        if cell not in target_cells:
            protected.setdefault(cell, f"current filled content already present before this repair: {str(value)[:80]}")
    for rng in reflect_stage.correct_ranges:
        for cell in _range_cells_set([rng]):
            if cell not in target_cells:
                protected.setdefault(cell, "reflect marked this region as correct/safe_blank")
    for assess in reflect_stage.reflect_result.get("risk_assessments", []) or []:
        if not isinstance(assess, dict):
            continue
        if bool(assess.get("is_high_risk", False)):
            continue
        if (not bool(assess.get("is_correct_region", False))) and (not bool(assess.get("is_safe_blank", False))) and int(assess.get("score", 0) or 0) != 0:
            continue
        protect_ranges: list[str] = []
        target_range = str(assess.get("target_range") or "").strip()
        target_cell = str(assess.get("target_cell") or "").strip()
        candidate_value_cell = str(assess.get("candidate_value_cell") or "").strip()
        if target_range:
            protect_ranges.append(target_range)
        elif target_cell:
            protect_ranges.append(target_cell)
        if candidate_value_cell:
            protect_ranges.append(candidate_value_cell)
        for rng in protect_ranges:
            for cell in _range_cells_set([rng]):
                if cell not in target_cells:
                    protected.setdefault(cell, "previously checked and considered okay")
    return protected


def _latest_repair_review_stage(state: ControllerState) -> ReflectStageResult | None:
    current = state.current_reflect
    verification = state.verification
    if verification is None:
        return current
    if current is None:
        return verification
    return verification if int(verification.round_idx or 0) >= int(current.round_idx or 0) else current


def _build_mandatory_review_items(
    state: ControllerState,
    args: Any,
    region_plan: ReflectRegionPlan,
    before_reset_xlsx: Path,
) -> list[dict[str, Any]]:
    review_stage = _latest_repair_review_stage(state)
    ranges: list[str] = []
    for rng in review_stage.screenshot_check.get("bad_ranges", []) if review_stage else []:
        text = str(rng or "").strip()
        if text:
            try:
                if legacy._range_intersects(text, region_plan.region):
                    ranges.append(legacy._normalize_a1_range(text))
            except Exception:
                continue
    for target in region_plan.patch_targets:
        for key in ("target_range", "source_cell"):
            text = str(target.get(key) or "").strip()
            if text:
                try:
                    if legacy._range_intersects(text, region_plan.region):
                        ranges.append(legacy._normalize_a1_range(text))
                except Exception:
                    continue
    deduped = list(dict.fromkeys(ranges))
    items: list[dict[str, Any]] = []
    for rng in deduped:
        current_value = legacy._read_anchor_value_from_workbook(before_reset_xlsx, args.sheet, rng)
        template_value = legacy._read_anchor_value_from_workbook(state.paths.before_fill_xlsx, args.sheet, rng)
        item = {
            "range": rng,
            "previous_value": current_value,
            "template_value": template_value,
            "must_recheck": True,
            "hint_source": "mandatory_review_range",
        }
        items.append(item)
    return items


def _build_latest_repair_feedback(state: ControllerState, region: str) -> dict[str, Any]:
    review_stage = _latest_repair_review_stage(state)
    if review_stage is None:
        return {}
    bad_ranges: list[str] = []
    for rng in (review_stage.bad_ranges or []):
        text = str(rng or "").strip()
        if not text:
            continue
        try:
            if legacy._range_intersects(text, region):
                bad_ranges.append(legacy._normalize_a1_range(text))
        except Exception:
            continue
    correct_ranges: list[str] = []
    for rng in (review_stage.correct_ranges or []):
        text = str(rng or "").strip()
        if not text:
            continue
        try:
            if legacy._range_intersects(text, region):
                correct_ranges.append(legacy._normalize_a1_range(text))
        except Exception:
            continue
    screenshot = review_stage.screenshot_check if isinstance(review_stage.screenshot_check, dict) else {}
    reason = str(screenshot.get("reason") or "").strip()
    payload = {
        "stage_name": str(review_stage.stage_name or ""),
        "round_idx": int(review_stage.round_idx or 0),
        "need_fix": bool(screenshot.get("need_fix")),
        "bad_ranges": list(dict.fromkeys(bad_ranges)),
        "correct_ranges": list(dict.fromkeys(correct_ranges)),
        "reason": reason,
        "before_sheet_png": str(screenshot.get("before_sheet_png") or "").strip(),
        "after_sheet_png": str(screenshot.get("after_sheet_png") or "").strip(),
    }
    return payload if payload["bad_ranges"] or payload["reason"] else {}


def _collect_guarded_write_candidates(
    region_plan: ReflectRegionPlan,
    mandatory_review_items: list[dict[str, Any]],
    latest_repair_feedback: dict[str, Any] | None,
) -> list[str]:
    candidates: set[str] = set()
    avoid_cells: set[str] = set()

    def _add_cell_like(raw_value: Any, max_span: int = 12) -> None:
        text = str(raw_value or "").strip()
        if not text:
            return
        try:
            expanded = list(legacy._iter_range_cells(text))
        except Exception:
            expanded = [text]
        if len(expanded) > max_span:
            return
        for cell_ref in expanded:
            cell_text = str(cell_ref or "").strip().upper()
            if cell_text:
                candidates.add(cell_text)

    def _add_avoid_cell(raw_value: Any) -> None:
        text = str(raw_value or "").strip()
        if not text:
            return
        try:
            expanded = list(legacy._iter_range_cells(text))
        except Exception:
            expanded = [text]
        for cell_ref in expanded:
            cell_text = str(cell_ref or "").strip().upper()
            if cell_text:
                avoid_cells.add(cell_text)

    for hint in region_plan.selected_region_hints:
        if not isinstance(hint, dict):
            continue
        _add_cell_like(hint.get("candidate_value_cell"))
    for item in mandatory_review_items:
        if not isinstance(item, dict):
            continue
        _add_cell_like(item.get("range"))
    for target in region_plan.patch_targets:
        if not isinstance(target, dict):
            continue
        _add_cell_like(target.get("target_range"))
        _add_cell_like(target.get("source_cell"))
    for invalid in region_plan.invalid_slots:
        if not isinstance(invalid, dict):
            continue
        _add_avoid_cell(invalid.get("candidate_value_cell") or invalid.get("value_cell"))
    if latest_repair_feedback:
        for rng in latest_repair_feedback.get("bad_ranges", []) or []:
            _add_cell_like(rng)
    qwen_context_summary = region_plan.qwen_context_summary or {}
    for group in qwen_context_summary.get("candidate_key_groups", []) or []:
        if not isinstance(group, dict):
            continue
        for cell_ref in group.get("candidate_cells", []) or []:
            _add_cell_like(cell_ref, max_span=1)
        for cell_ref in group.get("blank_candidate_cells", []) or []:
            _add_cell_like(cell_ref, max_span=1)
    return sorted(cell for cell in candidates if cell and cell not in avoid_cells)


def _cell_sort_key(cell_ref: str) -> tuple[int, int, str]:
    text = str(cell_ref or "").strip().upper()
    if not text:
        return (10**9, 10**9, "")
    try:
        first = legacy._normalize_a1_range(text).split(",")[0].split(":")[0]
        col_text, row_idx = coordinate_from_string(first)
        return (int(row_idx), int(column_index_from_string(col_text)), first)
    except Exception:
        return (10**9, 10**9, text)


def _dedupe_cells(cells: list[str], limit: int = 12) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for cell in sorted((str(x or "").strip().upper() for x in cells if str(x or "").strip()), key=_cell_sort_key):
        if cell in seen:
            continue
        seen.add(cell)
        out.append(cell)
        if len(out) >= max(1, int(limit)):
            break
    return out


def _assemble_qwen_context_summary(
    region: str,
    patch_targets: list[dict[str, Any]],
    valid_items: list[dict[str, Any]],
    invalid_items: list[dict[str, Any]],
    raw_hint_count: int,
    source: str,
    blank_like_cells: list[str] | None = None,
) -> dict[str, Any]:
    if not valid_items and not invalid_items:
        return {}
    groups: dict[str, dict[str, Any]] = {}
    numeric_noise_cells: list[str] = []
    blank_candidates = {str(x or "").strip().upper() for x in (blank_like_cells or []) if str(x or "").strip()}

    def _add_item(key: str, cell: str, is_invalid: bool, reason: str = "", hint_source: str = "") -> None:
        nonlocal groups, numeric_noise_cells
        if not cell:
            return
        cell = cell.upper()
        if not cell:
            return
        if _is_noisy_reflect_key(key):
            numeric_noise_cells.append(cell)
            return
        norm_key = normalize_key(key) or key.lower()
        group = groups.setdefault(
            norm_key,
            {
                "key": key,
                "candidate_cells": [],
                "invalid_cells": [],
                "blank_candidate_cells": [],
                "sources": [],
                "reasons": [],
            },
        )
        if is_invalid:
            group["invalid_cells"].append(cell)
        else:
            group["candidate_cells"].append(cell)
        if cell in blank_candidates:
            group["blank_candidate_cells"].append(cell)
        if hint_source and hint_source not in group["sources"]:
            group["sources"].append(hint_source)
        if reason and reason not in group["reasons"]:
            group["reasons"].append(reason)
    for item in valid_items:
        _add_item(
            key=str(item.get("key") or "").strip(),
            cell=str(item.get("candidate_value_cell") or item.get("value_cell") or "").strip(),
            is_invalid=False,
            reason=str(item.get("reason") or "").strip(),
            hint_source=str(item.get("hint_source") or "qwen_region_form_valid").strip(),
        )
    for item in invalid_items:
        _add_item(
            key=str(item.get("key") or "").strip(),
            cell=str(item.get("candidate_value_cell") or item.get("value_cell") or "").strip(),
            is_invalid=True,
            reason=str(item.get("reason") or "").strip(),
            hint_source=str(item.get("hint_source") or "qwen_region_form_invalid").strip(),
        )
    candidate_key_groups: list[dict[str, Any]] = []
    negative_only_key_groups: list[dict[str, Any]] = []
    for group in groups.values():
        payload = {
            "key": str(group.get("key") or "").strip(),
            "candidate_cells": _dedupe_cells(list(group.get("candidate_cells") or [])),
            "invalid_cells": _dedupe_cells(list(group.get("invalid_cells") or [])),
            "blank_candidate_cells": _dedupe_cells(list(group.get("blank_candidate_cells") or [])),
            "sources": list(group.get("sources") or []),
            "reasons": list(group.get("reasons") or [])[:4],
        }
        if payload["candidate_cells"]:
            candidate_key_groups.append(payload)
        elif payload["invalid_cells"]:
            negative_only_key_groups.append(payload)
    candidate_key_groups.sort(
        key=lambda item: (
            -len(item.get("candidate_cells", [])),
            -len(item.get("invalid_cells", [])),
            str(item.get("key") or "").lower(),
        )
    )
    negative_only_key_groups.sort(
        key=lambda item: (
            -len(item.get("invalid_cells", [])),
            str(item.get("key") or "").lower(),
        )
    )
    patch_target_keys = sorted(
        {
            normalize_key(str(item.get("key") or ""))
            for item in (patch_targets or [])
            if str(item.get("key") or "").strip()
        }
    )
    summary_lines: list[str] = []
    for item in candidate_key_groups[:6]:
        line = (
            f"Qwen associates key '{item['key']}' with nearby candidate cells "
            f"{', '.join(item['candidate_cells'])}."
        )
        if item["invalid_cells"]:
            line += f" Nearby cells {', '.join(item['invalid_cells'])} were judged invalid/avoid."
        if item["blank_candidate_cells"]:
            line += f" Blank-looking candidates were seen at {', '.join(item['blank_candidate_cells'])}."
        summary_lines.append(line)
    for item in negative_only_key_groups[:4]:
        summary_lines.append(
            f"Qwen only found invalid/avoid cells for key '{item['key']}' at {', '.join(item['invalid_cells'])}; use this as evidence that the field exists nearby, but do not write back into those cells."
        )
    numeric_noise_cells = _dedupe_cells(numeric_noise_cells, limit=16)
    if numeric_noise_cells:
        summary_lines.append(
            f"Numeric-only or overlay-like pseudo-keys were detected at {', '.join(numeric_noise_cells)}; treat them as noise, not semantic field labels."
        )
    payload = {
        "region": region,
        "source": source,
        "raw_hint_count": int(raw_hint_count),
        "valid_item_count": len(valid_items),
        "invalid_item_count": len(invalid_items),
        "patch_target_keys": patch_target_keys,
        "candidate_key_groups": candidate_key_groups[:8],
        "negative_only_key_groups": negative_only_key_groups[:6],
        "numeric_noise_cells": numeric_noise_cells,
        "summary_lines": summary_lines[:12],
    }
    return payload if payload["summary_lines"] or payload["candidate_key_groups"] or payload["negative_only_key_groups"] else {}


def _build_qwen_context_summary_from_region_form(
    region: str,
    patch_targets: list[dict[str, Any]],
    region_form: dict[str, Any],
    current_xlsx: Path,
    sheet_name: str | None,
) -> dict[str, Any]:
    pairs = [x for x in ((region_form or {}).get("pairs", []) or []) if isinstance(x, dict)]
    invalid_slots = [x for x in ((region_form or {}).get("invalid_slots", []) or []) if isinstance(x, dict)]
    valid_items: list[dict[str, Any]] = []
    invalid_items: list[dict[str, Any]] = []
    blank_like_cells: list[str] = []
    for pair in pairs:
        cell = str(pair.get("value_cell") or "").strip()
        if not cell:
            continue
        candidate_value = legacy._read_anchor_value_from_workbook(current_xlsx, sheet_name, cell)
        if legacy._is_blankish_value(candidate_value):
            blank_like_cells.append(cell)
        valid_items.append(
            {
                "key": str(pair.get("key") or "").strip(),
                "value_cell": cell,
                "reason": str(pair.get("reason") or ""),
                "hint_source": "qwen_region_form_valid",
            }
        )
    for item in invalid_slots:
        cell = str(item.get("value_cell") or item.get("candidate_value_cell") or "").strip()
        if not cell:
            continue
        invalid_items.append(
            {
                "key": str(item.get("key") or "").strip(),
                "value_cell": cell,
                "reason": str(item.get("reason") or ""),
                "hint_source": "qwen_region_form_invalid",
            }
        )
    return _assemble_qwen_context_summary(
        region=region,
        patch_targets=patch_targets,
        valid_items=valid_items,
        invalid_items=invalid_items,
        raw_hint_count=len(pairs) + len(invalid_slots),
        source="region_form_full",
        blank_like_cells=blank_like_cells,
    )


def _build_qwen_context_summary(region: str, patch_targets: list[dict[str, Any]], raw_region_hints: list[dict[str, Any]]) -> dict[str, Any]:
    raw_hints = [x for x in (raw_region_hints or []) if isinstance(x, dict)]
    if not raw_hints:
        return {}
    embedded = [
        x.get("qwen_context_summary")
        for x in raw_hints
        if isinstance(x, dict)
        and bool(x.get("context_only"))
        and str(x.get("hint_source") or "") == "qwen_context_summary"
        and isinstance(x.get("qwen_context_summary"), dict)
    ]
    if embedded:
        summary = dict(embedded[0] or {})
        summary["matched_hint_count"] = len(
            [
                x
                for x in raw_hints
                if isinstance(x, dict)
                and not bool(x.get("context_only"))
            ]
        )
        return summary
    signal_hints = [x for x in raw_hints if not bool(x.get("context_only"))]
    if not signal_hints:
        return {}
    valid_items: list[dict[str, Any]] = []
    invalid_items: list[dict[str, Any]] = []
    blank_like_cells: list[str] = []
    for hint in signal_hints:
        payload = {
            "key": str(hint.get("key") or "").strip(),
            "candidate_value_cell": str(hint.get("candidate_value_cell") or hint.get("value_cell") or "").strip(),
            "reason": str(hint.get("reason") or "").strip(),
            "hint_source": str(hint.get("hint_source") or "").strip(),
        }
        if bool(hint.get("is_blank_candidate")):
            blank_like_cells.append(payload["candidate_value_cell"])
        if bool(hint.get("invalid_slot")) or bool(hint.get("avoid_cell")):
            invalid_items.append(payload)
        else:
            valid_items.append(payload)
    return _assemble_qwen_context_summary(
        region=region,
        patch_targets=patch_targets,
        valid_items=valid_items,
        invalid_items=invalid_items,
        raw_hint_count=len(signal_hints),
        source="matched_region_hints",
        blank_like_cells=blank_like_cells,
    )


def free_fill(state: ControllerState, args: Any, log_fn: Callable[[str], None]) -> dict[str, Any]:
    form_context_mode = str(getattr(args, "agent_form_context_mode", "struct") or "struct").strip().lower()
    first_pass_form_context_mode = "struct" if form_context_mode == "refill_html" else form_context_mode
    first_pass_html_path = None
    if first_pass_form_context_mode in {"html", "struct_html"}:
        try:
            prompt_render = legacy.preprocess_excel_html_css(
                xlsx_path=state.paths.input_xlsx,
                out_dir=state.paths.job_dir / "first_pass_prompt_render",
                sheet_name=args.sheet,
                dpi=args.dpi,
                canvas_bg=args.render_canvas_bg,
                png_scale=args.render_png_scale,
                max_render_rows=args.render_max_rows,
                max_render_cols=args.render_max_cols,
            )
            first_pass_html_path = prompt_render.get("html")
            if first_pass_html_path:
                log_fn(f"[INFO] first-pass agent prompt includes rendered HTML: {first_pass_html_path}")
        except Exception as exc:
            log_fn(f"[WARN] failed to prepare rendered HTML for first-pass prompt: {exc}")
    elif form_context_mode == "refill_html":
        log_fn("[INFO] first-pass uses structural context only; rendered HTML reserved for guided refill.")
    if state.paths.edit_log_jsonl.exists():
        state.paths.edit_log_jsonl.unlink()
    effective_instruction = str(args.instruction or "")
    dropdown_context = _build_dropdown_context(state.form.get("dropdown_constraints", []) if isinstance(state.form, dict) else [])
    if dropdown_context:
        effective_instruction = f"{effective_instruction}\n\n{dropdown_context}".strip()
    result = _agent_fill_with_large_sheet_retry(
        log_fn=log_fn,
        stage_label="FREE_FILL",
        agent_fill_kwargs=dict(
            template_xlsx=state.paths.input_xlsx,
            output_xlsx=state.paths.output_xlsx,
            instruction=effective_instruction,
            form={"pairs": [], "invalid_slots": []},
            model=args.agent_model,
            base_url=args.agent_base_url,
            api_key=args.agent_api_key or None,
            tool_mode="code",
            max_steps=args.max_steps,
            skills_text=state.skills_text,
            mapping_mode="off",
            sheet_name=args.sheet,
            edit_log_jsonl=state.paths.edit_log_jsonl,
            round_idx=0,
            history_dump_path=state.paths.job_dir / "first_pass_history.json",
            temperature=float(args.agent_temperature),
            planner_policy=None,
            planner_stage="first_pass",
            model_io_dir=(state.paths.model_io_dir / "first_pass_agent" if state.paths.model_io_dir is not None else None),
            agent_profile="free_fill",
            enforce_label_guard=False,
            enforce_bottom_border_guard=True,
            enforce_protected_cells=False,
            rendered_html_path=first_pass_html_path,
            rendered_html_max_chars=int(args.agent_rendered_html_max_chars),
            form_context_mode=first_pass_form_context_mode,
        ),
    )
    if bool(result.get("no_code_generated")):
        log_fn(
            "[FREE_FILL] agent produced no executable code; "
            f"continuing with unchanged workbook. reason={result.get('no_code_generated_reason', '')}"
        )
    state.usage_summary["first_pass_agent"] = result.get("usage", state.usage_summary["first_pass_agent"])
    try:
        render = legacy.preprocess_excel_html_css(
            xlsx_path=state.paths.output_xlsx,
            out_dir=state.paths.job_dir / "first_pass_visual",
            sheet_name=args.sheet,
            dpi=args.dpi,
            canvas_bg=args.render_canvas_bg,
            png_scale=args.render_png_scale,
            max_render_rows=args.render_max_rows,
            max_render_cols=args.render_max_cols,
        )
        state.first_pass_sheet_png = str(render.get("sheet_png") or "")
        if state.first_pass_sheet_png:
            result["first_pass_sheet_png"] = state.first_pass_sheet_png
            log_fn(f"[INFO] saved first-pass screenshot: {state.first_pass_sheet_png}")
    except Exception as exc:
        log_fn(f"[WARN] failed to render first-pass screenshot: {exc}")
    state.first_pass_result = result
    return result


def _run_assessment_stage(
    state: ControllerState,
    args: Any,
    log_fn: Callable[[str], None],
    round_idx: int,
    stage_name: str,
    build_hints: bool,
    preset_screenshot_check: dict[str, Any] | None = None,
) -> ReflectStageResult:
    stage_dir_root = "reflect_plugin" if build_hints else "verify"
    stage_dir = state.paths.job_dir / stage_dir_root / f"round_{round_idx}"
    stage_dir.mkdir(parents=True, exist_ok=True)
    checker_model = args.reflect_checker_model.strip() or args.agent_model
    checker_context_summary = {
        "struct_summary": state.first_pass_result.get("struct_summary", {}),
        "first_pass_write_log": (state.first_pass_result.get("write_log", []) or [])[:40],
    }
    screenshot_check: dict[str, Any] = {}
    if isinstance(preset_screenshot_check, dict) and preset_screenshot_check:
        screenshot_check = dict(preset_screenshot_check)
        log_fn(f"[{stage_name.upper()}] round={round_idx} reusing_prior_screenshot_check={json.dumps(screenshot_check, ensure_ascii=False)}")
        _log_screenshot_check_details(log_fn, stage_name, round_idx, screenshot_check)
        if build_hints:
            state.usage_summary["reflect_plugin_stats"]["rounds_started"] += 1
    else:
        try:
            screenshot_check = legacy.reflect_check_with_screenshot(
                before_xlsx_path=state.paths.before_fill_xlsx,
                after_xlsx_path=state.paths.output_xlsx,
                sheet_name=args.sheet,
                instruction=args.instruction,
                model=checker_model,
                base_url=args.agent_base_url or None,
                api_key=args.agent_api_key or None,
                out_dir=stage_dir / "check",
                dpi=args.dpi,
                checker_context_summary=checker_context_summary,
                planner_policy=state.planner_policy,
                model_io_dir=((state.paths.model_io_dir / f"{stage_name}_{round_idx}" / "stage1_checker") if state.paths.model_io_dir is not None else None),
            )
            log_fn(f"[{stage_name.upper()}] round={round_idx} screenshot_check={json.dumps(screenshot_check, ensure_ascii=False)}")
            _log_screenshot_check_details(log_fn, stage_name, round_idx, screenshot_check)
            if isinstance(screenshot_check.get("usage"), dict):
                state.usage_summary["reflect_checker"]["prompt_tokens"] += int(screenshot_check["usage"].get("prompt_tokens", 0))
                state.usage_summary["reflect_checker"]["completion_tokens"] += int(screenshot_check["usage"].get("completion_tokens", 0))
                state.usage_summary["reflect_checker"]["total_tokens"] += int(screenshot_check["usage"].get("total_tokens", 0))
            if build_hints:
                state.usage_summary["reflect_plugin_stats"]["rounds_started"] += 1
                if bool(screenshot_check.get("need_fix")):
                    state.usage_summary["reflect_plugin_stats"]["stage1_need_fix_count"] += 1
                if bool(screenshot_check.get("bad_range_parse_failed")):
                    state.usage_summary["reflect_plugin_stats"]["stage1_parse_failed_count"] += 1
        except Exception as exc:
            log_fn(f"[{stage_name.upper()}] round={round_idx} screenshot_check_failed: {exc}")

    use_region_qwen_hints = build_hints
    hint_builder = _build_reflect_hint_builder(state, args, log_fn) if use_region_qwen_hints else None
    reflect_result_obj = run_reflect_round(
        ReflectContext(
            workbook_path=state.paths.output_xlsx,
            instruction=args.instruction,
            form=state.form,
            edit_log_jsonl=state.paths.edit_log_jsonl,
            out_dir=stage_dir,
            model=checker_model,
            base_url=args.agent_base_url or None,
            api_key=args.agent_api_key or None,
            round_idx=round_idx,
            sheet_name=args.sheet,
            risk_threshold=args.reflect_risk_threshold,
            expand_left_cols=args.reflect_expand_left_cols,
            expand_right_cols=args.reflect_expand_right_cols,
            expand_top_rows=args.reflect_expand_top_rows,
            expand_bottom_rows=args.reflect_expand_bottom_rows,
            hint_builder=hint_builder,
            prior_context_summary={
                "write_log": state.first_pass_result.get("write_log", []),
                "history_dump_path": state.first_pass_result.get("history_dump_path", ""),
            },
            screenshot_check=screenshot_check,
            first_pass_sheet_png=state.first_pass_sheet_png,
            planner_policy_summary=legacy._planner_policy_prompt_block(state.planner_policy, "reflect"),
            planner_policy=state.planner_policy,
            model_io_dir=((state.paths.model_io_dir / f"{stage_name}_{round_idx}" / "stage2_assess") if state.paths.model_io_dir is not None else None),
            model_direct_ranges_only=False,
        )
    )
    reflect_payload = asdict(reflect_result_obj)
    _log_reflect_result_details(log_fn, stage_name, round_idx, reflect_payload, screenshot_check)
    if isinstance(reflect_payload.get("model_usage"), dict):
        state.usage_summary["reflect_model_assess"]["prompt_tokens"] += int(reflect_payload["model_usage"].get("prompt_tokens", 0))
        state.usage_summary["reflect_model_assess"]["completion_tokens"] += int(reflect_payload["model_usage"].get("completion_tokens", 0))
        state.usage_summary["reflect_model_assess"]["total_tokens"] += int(reflect_payload["model_usage"].get("total_tokens", 0))

    high_risk_cells = [x for x in (reflect_payload.get("risk_assessments", []) or []) if isinstance(x, dict) and bool(x.get("is_high_risk"))]
    correct_ranges = [str(x) for x in (reflect_payload.get("correct_ranges", []) or []) if str(x).strip()]
    patch_jobs = [x for x in (reflect_payload.get("patch_jobs", []) or []) if isinstance(x, dict)]
    bad_ranges = [str(x) for x in (screenshot_check.get("bad_ranges", []) or []) if str(x).strip()]
    for target in (reflect_payload.get("reflect_targets", []) or []):
        if not isinstance(target, dict):
            continue
        if str(target.get("kind") or "").strip() != "label_area_edit":
            continue
        text = str(target.get("target_range") or target.get("source_cell") or "").strip()
        if text:
            bad_ranges.append(text)
    bad_ranges.extend([str(job.get("region") or "").strip() for job in patch_jobs if str(job.get("region") or "").strip()])
    bad_ranges = list(dict.fromkeys(bad_ranges))

    region_plans: list[ReflectRegionPlan] = []
    correct_cells_all = _range_cells_set(correct_ranges)
    for patch_job in patch_jobs:
        region = str(patch_job.get("region") or "").strip()
        if not region:
            continue
        patch_targets = _select_patch_targets(patch_job)
        raw_region_hints = [x for x in (patch_job.get("region_hints", []) or []) if isinstance(x, dict)]
        matched_hints = _select_patch_region_hints(raw_region_hints, patch_targets)
        positive_hints = [x for x in matched_hints if not bool(x.get("invalid_slot")) and not bool(x.get("avoid_cell"))]
        invalid_slots = [x for x in matched_hints if bool(x.get("invalid_slot")) or bool(x.get("avoid_cell"))]
        qwen_context_summary = (
            _build_qwen_context_summary(region, patch_targets, raw_region_hints)
            if use_region_qwen_hints
            else {}
        )
        correct_cells_in_region = {cell for cell in correct_cells_all if legacy._range_intersects(cell, region)}
        damage = inspect_template_region_damage(
            template_xlsx=state.paths.before_fill_xlsx,
            current_xlsx=state.paths.output_xlsx,
            sheet_name=args.sheet,
            region=region,
            preserve_cells=correct_cells_in_region,
        )
        region_plans.append(
            ReflectRegionPlan(
                region=region,
                patch_targets=patch_targets,
                selected_region_hints=positive_hints,
                invalid_slots=invalid_slots,
                qwen_context_summary=qwen_context_summary,
                correct_ranges=[rng for rng in correct_ranges if legacy._range_intersects(rng, region)],
                need_restore_template=bool(damage.get("need_restore")),
                restore_template_region=str(damage.get("restore_region") or ""),
                restore_reasons=[str(x) for x in (damage.get("reasons") or []) if str(x).strip()],
                raw_region_hints=raw_region_hints,
                raw_patch_job=patch_job,
            )
        )

    if build_hints:
        assessed_cells = reflect_payload.get("edits_considered", []) if isinstance(reflect_payload.get("edits_considered"), list) else []
        edited_cells = [x for x in assessed_cells if isinstance(x, dict) and bool(x.get("is_actual_edit"))]
        scan_only_cells = [x for x in assessed_cells if isinstance(x, dict) and not bool(x.get("is_actual_edit"))]
        blank_targets = [x for x in (reflect_payload.get("reflect_targets", []) or []) if isinstance(x, dict) and str(x.get("kind") or "") == "region_blank_candidate"]
        state.usage_summary["reflect_plugin_stats"]["assessed_cells_total"] += len(assessed_cells)
        state.usage_summary["reflect_plugin_stats"]["edited_cells_total"] += len(edited_cells)
        state.usage_summary["reflect_plugin_stats"]["scan_only_cells_total"] += len(scan_only_cells)
        state.usage_summary["reflect_plugin_stats"]["high_risk_cells_total"] += len(high_risk_cells)
        state.usage_summary["reflect_plugin_stats"]["targets_total"] += len(reflect_payload.get("reflect_targets", []) or [])
        state.usage_summary["reflect_plugin_stats"]["blank_targets_total"] += len(blank_targets)
        state.usage_summary["reflect_plugin_stats"]["patch_jobs_total"] += len(patch_jobs)

    stage_result = ReflectStageResult(
        round_idx=round_idx,
        stage_name=stage_name,
        screenshot_check=screenshot_check,
        reflect_result=reflect_payload,
        bad_ranges=bad_ranges,
        correct_ranges=correct_ranges,
        regions=region_plans,
        has_high_risk=bool(high_risk_cells or bad_ranges),
        high_risk_cells=high_risk_cells,
        missing_keys=[str(x) for x in (reflect_payload.get("missing_keys", []) or []) if str(x).strip()],
        assess_summary={
            "regions": len(region_plans),
            "patch_jobs": len(patch_jobs),
            "high_risk_cells": len(high_risk_cells),
            "bad_ranges": len(bad_ranges),
        },
    )
    _write_json(stage_dir / "controller_region_plans.json", [asdict(x) for x in region_plans])
    _write_json(stage_dir / "controller_stage_result.json", asdict(stage_result))
    return stage_result


def reflect_and_hint(state: ControllerState, args: Any, log_fn: Callable[[str], None], round_idx: int) -> ReflectStageResult:
    result = _run_assessment_stage(
        state=state,
        args=args,
        log_fn=log_fn,
        round_idx=round_idx,
        stage_name="reflect_and_hint",
        build_hints=True,
    )
    state.current_reflect = result
    return result


def rebuild_from_verification(
    state: ControllerState,
    args: Any,
    log_fn: Callable[[str], None],
    round_idx: int,
) -> ReflectStageResult:
    if state.verification is None:
        raise legacy.PipelineError("rebuild_from_verification requires an existing verification result")
    result = _run_assessment_stage(
        state=state,
        args=args,
        log_fn=log_fn,
        round_idx=round_idx,
        stage_name="rebuild_from_verification",
        build_hints=True,
        preset_screenshot_check=state.verification.screenshot_check,
    )
    state.current_reflect = result
    return result


def guided_refill(
    state: ControllerState,
    args: Any,
    log_fn: Callable[[str], None],
    region_plan: ReflectRegionPlan,
    force_repair: bool = False,
    noop_count: int = 0,
) -> dict[str, Any]:
    requested_form_context_mode = str(getattr(args, "agent_form_context_mode", "struct") or "struct").strip().lower()
    form_context_mode = "struct"
    if state.current_reflect is None:
        raise legacy.PipelineError("guided_refill requires an active reflect stage result")
    region = region_plan.region
    round_idx = state.current_reflect.round_idx
    region_debug_dir = state.paths.job_dir / "reflect_plugin" / f"round_{round_idx}" / f"guided_refill_{region.replace(':', '_')}"
    region_debug_dir.mkdir(parents=True, exist_ok=True)
    refill_html_path = None
    if requested_form_context_mode == "refill_html":
        form_context_mode = "html"
        try:
            prompt_render = legacy.preprocess_excel_html_css(
                xlsx_path=state.paths.output_xlsx,
                out_dir=region_debug_dir / "prompt_render",
                sheet_name=args.sheet,
                dpi=args.dpi,
                canvas_bg=args.render_canvas_bg,
                png_scale=args.render_png_scale,
                max_render_rows=args.render_max_rows,
                max_render_cols=args.render_max_cols,
            )
            refill_html_path = prompt_render.get("html")
            if refill_html_path:
                log_fn(f"[GUIDED_REFILL] region={region} prompt includes rendered HTML: {refill_html_path}")
        except Exception as exc:
            log_fn(f"[WARN] failed to prepare rendered HTML for guided refill region={region}: {exc}")
    elif requested_form_context_mode in {"html", "struct_html"}:
        log_fn(
            f"[GUIDED_REFILL] region={region} rendered HTML disabled for local repair; "
            "using workbook tools and regional evidence only."
        )
    template_summary = _summarize_template_region(state.paths.before_fill_xlsx, args.sheet, region)
    positive_hints = list(region_plan.selected_region_hints)
    invalid_slots = list(region_plan.invalid_slots)
    qwen_context_summary = dict(region_plan.qwen_context_summary or {})
    latest_repair_feedback = _build_latest_repair_feedback(state, region)
    correct_cells_in_region: set[str] = set()
    for rng in region_plan.correct_ranges:
        try:
            correct_cells_in_region.update(legacy._iter_range_cells(rng))
        except Exception:
            continue
    patch_pairs: list[dict[str, str]] = []
    for hint in positive_hints:
        key = str(hint.get("key") or "").strip()
        candidate_value_cell = str(hint.get("candidate_value_cell") or "").strip()
        if key and candidate_value_cell:
            patch_pairs.append({"key": key, "value_cell": candidate_value_cell})
    mapping_mode = "hint" if patch_pairs else "off"
    if mapping_mode == "hint" and not state.workflow_policy.get("prefer_hints_when_available", True):
        mapping_mode = "off"
        log_fn(f"[GUIDED_REFILL] region={region} planner policy disabled hint mode; falling back to guarded local refill.")
    backup_xlsx = region_debug_dir / "before_guided_refill.xlsx"
    shutil.copyfile(state.paths.output_xlsx, backup_xlsx)
    mandatory_review_items = _build_mandatory_review_items(state, args, region_plan, backup_xlsx)
    guarded_write_candidate_set = set(_collect_guarded_write_candidates(region_plan, mandatory_review_items, latest_repair_feedback))
    guarded_write_candidates = sorted(guarded_write_candidate_set)
    protected_write_cells = _build_protected_write_cells(
        region_plan=region_plan,
        reflect_stage=state.current_reflect,
        before_nonempty_cells=legacy._collect_nonempty_cells(state.paths.before_fill_xlsx, args.sheet),
        editable_cells=guarded_write_candidate_set,
        current_region_nonempty_cells=legacy._collect_nonempty_cells_in_range(backup_xlsx, args.sheet, region),
    )
    reset_result = restore_template_region(
        template_xlsx=state.paths.before_fill_xlsx,
        target_xlsx=state.paths.output_xlsx,
        sheet_name=args.sheet,
        region=region,
        preserve_cells=correct_cells_in_region,
    )
    synthetic_review_hints = [
        {
            "key": str(item.get("range") or ""),
            "target_range": str(item.get("range") or ""),
            "candidate_value_cell": str(item.get("range") or ""),
            "candidate_value": item.get("previous_value"),
            "must_recheck": True,
            "hint_source": "mandatory_review_range",
        }
        for item in mandatory_review_items
        if str(item.get("range") or "").strip()
    ]
    refill_instruction = (
        f"{args.instruction}\n\n"
        f"[Guided refill region]\n"
        f"Focus region: {region}.\n"
        "This is a guarded repair pass, not a free fill pass.\n"
        "Use the instruction, current workbook state, original template summary, and selected hints to repair only this region.\n"
        "Do not edit outside the allowed write range.\n"
        "Do not overwrite label cells or protected cells.\n"
        "The reported bad range is where the error is visible, not always the literal final write target.\n"
        "If a flagged cell is a label/static cell, use nearby structure to locate the actual writable value cell for that field.\n"
        "Empty merged cells/ranges directly below or beside filled/dark labels are strong writable-value candidates; verify locally and prefer them over the label cell.\n"
        "For topic/section labels, if the source contains a clearly matching section, it is valid to write a concise evidence-grounded summary into the corresponding blank merged value range.\n"
        "If the evidence is ambiguous, prefer leaving a cell blank over guessing.\n"
    )
    refill_instruction += (
        f"The controller has already reset questioned cells in this region back to template state before this repair pass. "
        f"reset_restored_count={int(reset_result.get('restored_count', 0) or 0)}.\n"
    )
    if force_repair:
        refill_instruction += (
            f"This region already had {int(noop_count)} prior guided repair attempt(s) with zero actual edits.\n"
            "You must explicitly re-check the mandatory review cells/ranges below before deciding you are done.\n"
            "Do not immediately return done without reading nearby structure for those cells.\n"
        )
    if region_plan.patch_targets:
        refill_instruction += f"Repair targets:\n{json.dumps(region_plan.patch_targets, ensure_ascii=False)}\n"
    if mandatory_review_items:
        refill_instruction += (
            "Mandatory review cells/ranges confirmed by verifier:\n"
            f"{json.dumps(mandatory_review_items, ensure_ascii=False)}\n"
            "These cells were specifically flagged as wrong. Re-evaluate them from local structure and nearby labels.\n"
        )
    if guarded_write_candidates:
        refill_instruction += (
            "Controller-approved local writable candidates for this repair pass:\n"
            f"{json.dumps(guarded_write_candidates, ensure_ascii=False)}\n"
            "Prefer these cells first. If write_cell warns that another target is outside the approved set, move back to this list.\n"
        )
    if latest_repair_feedback:
        refill_instruction += (
            "Latest checker/verifier issue summary for this region (high-priority evidence):\n"
            f"{json.dumps(latest_repair_feedback, ensure_ascii=False)}\n"
            "If the checker says a value was written into the wrong sibling field or a required neighboring value cell is still blank/template-like, you must correct that placement rather than preserving the current layout.\n"
        )
    if positive_hints:
        refill_instruction += f"Selected region hints:\n{json.dumps(positive_hints, ensure_ascii=False)}\n"
    if invalid_slots:
        refill_instruction += f"Negative hints / invalid slots:\n{json.dumps(invalid_slots, ensure_ascii=False)}\n"
    dropdown_context = _build_dropdown_context(
        state.form.get("dropdown_constraints", []) if isinstance(state.form, dict) else [],
        focus_range=region,
    )
    if dropdown_context:
        refill_instruction += dropdown_context
    if qwen_context_summary:
        refill_instruction += (
            "Qwen local context summary (soft regional evidence; use it to understand nearby field structure even when no direct write-target hint survived selection):\n"
            f"{json.dumps(qwen_context_summary, ensure_ascii=False)}\n"
        )
        log_fn(
            f"[GUIDED_REFILL] region={region} qwen_context_candidates={len(qwen_context_summary.get('candidate_key_groups', []) or [])} "
            f"negative_only={len(qwen_context_summary.get('negative_only_key_groups', []) or [])} "
            f"raw_hints={int(qwen_context_summary.get('raw_hint_count', 0) or 0)}"
        )
    refill_instruction += f"Original template summary:\n{json.dumps(template_summary, ensure_ascii=False)}\n"
    _write_json(region_debug_dir / "qwen_context_summary.json", qwen_context_summary)
    patch_result = _agent_fill_with_large_sheet_retry(
        log_fn=log_fn,
        stage_label=f"GUIDED_REFILL region={region}",
        agent_fill_kwargs=dict(
            template_xlsx=state.paths.output_xlsx,
            output_xlsx=state.paths.output_xlsx,
            instruction=refill_instruction,
            form={
                "pairs": patch_pairs,
                "invalid_slots": invalid_slots,
                "guarded_write_candidates": guarded_write_candidates,
            },
            model=args.agent_model,
            base_url=args.agent_base_url,
            api_key=args.agent_api_key or None,
            tool_mode="code",
            max_steps=max(1, int(args.reflect_patch_steps)) if not force_repair else max(2, int(args.reflect_patch_steps) * 2),
            skills_text=state.skills_text,
            mapping_mode=mapping_mode,
            sheet_name=args.sheet,
            allowed_write_range=region,
            edit_log_jsonl=state.paths.edit_log_jsonl,
            round_idx=round_idx,
            reflect_hints=positive_hints + synthetic_review_hints + invalid_slots,
            context_summary={
                "first_pass_write_log": state.first_pass_result.get("write_log", []),
                "reflect_targets": region_plan.patch_targets,
                "template_summary": template_summary,
                "mandatory_review_items": mandatory_review_items,
                "reset_result": reset_result,
                "qwen_context_summary": qwen_context_summary,
                "latest_repair_feedback": latest_repair_feedback,
            },
            compare_before_after_pngs={
                "before": str(latest_repair_feedback.get("before_sheet_png") or state.current_reflect.screenshot_check.get("before_sheet_png") or ""),
                "after": str(latest_repair_feedback.get("after_sheet_png") or state.current_reflect.screenshot_check.get("after_sheet_png") or ""),
                "summary": json.dumps(
                    {
                        "bad_ranges": latest_repair_feedback.get("bad_ranges", state.current_reflect.bad_ranges),
                        "need_fix": latest_repair_feedback.get("need_fix", state.current_reflect.screenshot_check.get("need_fix")),
                        "reason": latest_repair_feedback.get("reason", state.current_reflect.screenshot_check.get("reason")),
                    },
                    ensure_ascii=False,
                ),
            },
            history_dump_path=region_debug_dir / "history.json",
            temperature=float(args.agent_temperature),
            protected_write_cells=protected_write_cells,
            planner_policy=state.planner_policy,
            planner_stage="refill",
            model_io_dir=((state.paths.model_io_dir / f"guided_refill_{round_idx}_{region.replace(':', '_')}") if state.paths.model_io_dir is not None else None),
            agent_profile="guided_refill",
            enforce_label_guard=True,
            enforce_bottom_border_guard=False,
            enforce_protected_cells=True,
            role_template_xlsx=state.paths.before_fill_xlsx,
            rendered_html_path=refill_html_path,
            rendered_html_max_chars=int(args.agent_rendered_html_max_chars),
            form_context_mode=form_context_mode,
        ),
    )
    actual_edits = [
        row for row in (patch_result.get("write_log", []) or [])
        if isinstance(row, dict) and bool(row.get("is_actual_edit"))
    ]
    if not actual_edits:
        shutil.copyfile(backup_xlsx, state.paths.output_xlsx)
        state.region_noop_counts[region] = int(state.region_noop_counts.get(region, 0) or 0) + 1
        log_fn(f"[GUIDED_REFILL] region={region} produced no actual edits; restored previous workbook snapshot.")
    else:
        state.region_noop_counts[region] = 0
        state.usage_summary["patch_agents"].append(
            {
                "round": round_idx,
                "region": region,
                "usage": patch_result.get("usage", {}),
            }
        )
        state.usage_summary["reflect_plugin_stats"]["patch_written_actual_edits_total"] += len(actual_edits)
        log_fn(f"[GUIDED_REFILL] region={region} actual_edits={len(actual_edits)}")
    patch_result["reset_result"] = reset_result
    patch_result["mandatory_review_items"] = mandatory_review_items
    patch_result["force_repair"] = bool(force_repair)
    patch_result["noop_count_before"] = int(noop_count)
    patch_result["noop_count_after"] = int(state.region_noop_counts.get(region, 0) or 0)
    patch_result["guarded_write_candidates"] = guarded_write_candidates
    patch_result["qwen_context_summary"] = qwen_context_summary
    patch_result["latest_repair_feedback"] = latest_repair_feedback
    _write_json(region_debug_dir / "patch_result.json", patch_result)
    state.regions_refilled.add(region)
    return patch_result


def verify_and_deliver(
    state: ControllerState,
    args: Any,
    log_fn: Callable[[str], None],
    round_idx: int,
) -> ReflectStageResult:
    result = _run_assessment_stage(
        state=state,
        args=args,
        log_fn=log_fn,
        round_idx=round_idx,
        stage_name="verify_and_deliver",
        build_hints=False,
    )
    state.verification = result
    if not result.has_high_risk and not result.bad_ranges:
        state.delivered = True
    return result


def finalize_usage_summary(state: ControllerState) -> dict[str, Any]:
    patch_prompt = sum(int((item.get("usage") or {}).get("prompt_tokens", 0)) for item in state.usage_summary["patch_agents"])
    patch_completion = sum(int((item.get("usage") or {}).get("completion_tokens", 0)) for item in state.usage_summary["patch_agents"])
    patch_total = sum(int((item.get("usage") or {}).get("total_tokens", 0)) for item in state.usage_summary["patch_agents"])
    state.usage_summary["patch_agents_total"] = {
        "prompt_tokens": patch_prompt,
        "completion_tokens": patch_completion,
        "total_tokens": patch_total,
    }
    state.usage_summary["grand_total"] = {
        "prompt_tokens": int(state.usage_summary["first_pass_agent"].get("prompt_tokens", 0))
        + int(state.usage_summary["reflect_checker"].get("prompt_tokens", 0))
        + int(state.usage_summary["reflect_model_assess"].get("prompt_tokens", 0))
        + patch_prompt,
        "completion_tokens": int(state.usage_summary["first_pass_agent"].get("completion_tokens", 0))
        + int(state.usage_summary["reflect_checker"].get("completion_tokens", 0))
        + int(state.usage_summary["reflect_model_assess"].get("completion_tokens", 0))
        + patch_completion,
        "total_tokens": int(state.usage_summary["first_pass_agent"].get("total_tokens", 0))
        + int(state.usage_summary["reflect_checker"].get("total_tokens", 0))
        + int(state.usage_summary["reflect_model_assess"].get("total_tokens", 0))
        + patch_total,
    }
    usage_summary_path = state.paths.job_dir / "run_usage_summary.json"
    _write_json(usage_summary_path, state.usage_summary)
    return state.usage_summary
