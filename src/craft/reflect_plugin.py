from __future__ import annotations

import json
import os
import re
import base64
import mimetypes
from difflib import SequenceMatcher
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Callable, Optional

import openpyxl
from openai import OpenAI
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries


def build_client(base_url: str | None, api_key: str | None) -> OpenAI:
    key = api_key or os.getenv("OPENAI_API_KEY") or os.getenv("PPCHAT_API_KEY")
    if not key:
        raise RuntimeError("Missing API key. Set OPENAI_API_KEY or PPCHAT_API_KEY.")
    return OpenAI(base_url=base_url or os.getenv("OPENAI_BASE_URL"), api_key=key)


def normalize_key(text: str) -> str:
    s = str(text or "").strip().lower()
    if s.endswith(":"):
        s = s[:-1].strip()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def normalize_a1_range(a1: str) -> str:
    s = str(a1 or "").strip().upper().replace(" ", "")
    if not s:
        raise ValueError("empty range")
    if ":" not in s:
        _ = range_boundaries(s)
        return f"{s}:{s}"
    c1, r1, c2, r2 = range_boundaries(s)
    left = f"{get_column_letter(min(c1, c2))}{min(r1, r2)}"
    right = f"{get_column_letter(max(c1, c2))}{max(r1, r2)}"
    return f"{left}:{right}"


def range_intersects(a: str, b: str) -> bool:
    a1, b1, a2, b2 = range_boundaries(normalize_a1_range(a))
    c1, d1, c2, d2 = range_boundaries(normalize_a1_range(b))
    rows_overlap = not (b2 < d1 or d2 < b1)
    cols_overlap = not (a2 < c1 or c2 < a1)
    return rows_overlap and cols_overlap


def union_ranges(ranges: list[str]) -> str:
    coords: list[tuple[int, int, int, int]] = []
    for rng in ranges:
        try:
            coords.append(range_boundaries(normalize_a1_range(rng)))
        except Exception:
            continue
    if not coords:
        raise ValueError("no valid ranges")
    min_col = min(c1 for c1, _, _, _ in coords)
    min_row = min(r1 for _, r1, _, _ in coords)
    max_col = max(c2 for _, _, c2, _ in coords)
    max_row = max(r2 for _, _, _, r2 in coords)
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def iter_range_cells(range_str: str) -> list[str]:
    out: list[str] = []
    for token in [x.strip() for x in str(range_str or "").split(",") if x.strip()]:
        c1, r1, c2, r2 = range_boundaries(normalize_a1_range(token))
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                out.append(f"{get_column_letter(cc)}{rr}")
    return out


def coord_of(a1: str) -> tuple[int, int]:
    col_s, row = coordinate_from_string(a1)
    return int(row), int(column_index_from_string(col_s))


def _encode_image_base64(image_path: str | Path) -> tuple[str, str]:
    path = Path(image_path)
    mime, _ = mimetypes.guess_type(path.name)
    mime_token = (mime or "image/png").split("/")[-1]
    try:
        import io
        from PIL import Image
    except Exception:
        raw = path.read_bytes()
        return base64.b64encode(raw).decode("ascii"), mime_token

    max_dim = 7900
    with Image.open(path) as img:
        width, height = img.size
        if max(width, height) > max_dim:
            scale = min(max_dim / float(max(width, height)), 1.0)
            new_size = (
                max(1, int(width * scale)),
                max(1, int(height * scale)),
            )
            resample = getattr(Image, "Resampling", Image).LANCZOS
            img = img.resize(new_size, resample)
            if mime_token == "jpeg" and str(img.mode or "").upper() in {"RGBA", "LA", "P"}:
                img = img.convert("RGB")
            buf = io.BytesIO()
            save_format = "JPEG" if mime_token == "jpeg" else "PNG"
            save_mime = "jpeg" if mime_token == "jpeg" else "png"
            img.save(buf, format=save_format, optimize=True)
            return base64.b64encode(buf.getvalue()).decode("ascii"), save_mime

    raw = path.read_bytes()
    return base64.b64encode(raw).decode("ascii"), mime_token


def extract_json_object(text: str) -> dict[str, Any]:
    raw = str(text or "").strip()
    if not raw:
        return {}
    if "```" in raw:
        m = re.search(r"```(?:json)?\s*(.*?)```", raw, flags=re.IGNORECASE | re.DOTALL)
        if m:
            raw = m.group(1).strip()
    try:
        obj = json.loads(raw)
        return obj if isinstance(obj, dict) else {}
    except Exception:
        pass
    m = re.search(r"\{.*\}", raw, flags=re.DOTALL)
    if not m:
        return {}
    try:
        obj = json.loads(m.group(0))
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}


def usage_to_dict(resp: Any) -> dict[str, int]:
    usage = getattr(resp, "usage", None)
    if not usage:
        return {}
    out: dict[str, int] = {}
    for k in ("prompt_tokens", "completion_tokens", "total_tokens"):
        v = getattr(usage, k, None)
        if isinstance(v, int):
            out[k] = v
    return out


def _sanitize_model_io_value(obj: Any) -> Any:
    if isinstance(obj, dict):
        out: dict[str, Any] = {}
        for key, value in obj.items():
            if key == "image_url" and isinstance(value, dict):
                inner = dict(value)
                url = inner.get("url")
                if isinstance(url, str) and url.startswith("data:image/"):
                    inner["url"] = f"[data_url_omitted length={len(url)}]"
                out[key] = _sanitize_model_io_value(inner)
            else:
                out[key] = _sanitize_model_io_value(value)
        return out
    if isinstance(obj, list):
        return [_sanitize_model_io_value(x) for x in obj]
    if isinstance(obj, tuple):
        return [_sanitize_model_io_value(x) for x in obj]
    if isinstance(obj, Path):
        return str(obj)
    return obj


def _dump_model_io(
    base_dir: Path | None,
    name: str,
    request_payload: Any,
    response_payload: Any = None,
    raw_text: str = "",
    usage: dict[str, Any] | None = None,
    extra: dict[str, Any] | None = None,
) -> None:
    if base_dir is None:
        return
    dump_dir = Path(base_dir) / str(name)
    dump_dir.mkdir(parents=True, exist_ok=True)
    request_obj = _sanitize_model_io_value(request_payload)
    response_obj = _sanitize_model_io_value(response_payload)
    summary = {
        "request": request_obj,
        "response": response_obj,
        "raw_text": str(raw_text or ""),
        "usage": usage or {},
        "extra": _sanitize_model_io_value(extra or {}),
    }
    (dump_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    (dump_dir / "request.json").write_text(
        json.dumps(request_obj, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    if response_obj is not None:
        (dump_dir / "response.json").write_text(
            json.dumps(response_obj, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
    if str(raw_text or "").strip():
        (dump_dir / "response.txt").write_text(str(raw_text), encoding="utf-8")


@dataclass
class EditRecord:
    sheet: str
    requested_cell: str
    target_cell: str
    old_value: Any
    new_value: Any
    round_idx: int = 0
    step_idx: int = 0
    source: str = "agent_write"
    is_actual_edit: bool = False
    key: str = ""


@dataclass
class EditRiskAssessment:
    target_cell: str
    key: str
    score: int
    risk_flags: list[str] = field(default_factory=list)
    is_high_risk: bool = False
    is_actual_edit: bool = False
    is_correct_region: bool = False
    is_safe_blank: bool = False
    candidate_value_cell: str = ""
    candidate_value: Any = None
    target_range: str = ""
    reason: str = ""


@dataclass
class ReflectTarget:
    kind: str
    sheet: str
    key: str
    target_range: str
    source_cell: str
    risk_flags: list[str] = field(default_factory=list)


@dataclass
class ReflectHint:
    key: str
    target_range: str
    candidate_value_cell: str
    candidate_value: Any
    risk_flags: list[str] = field(default_factory=list)
    hint_source: str = "reflect_slot_kv"


@dataclass
class PatchJob:
    region: str
    region_hints: list[dict[str, Any]] = field(default_factory=list)
    targets: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class ReflectRoundResult:
    edits_considered: list[dict[str, Any]] = field(default_factory=list)
    risk_assessments: list[dict[str, Any]] = field(default_factory=list)
    missing_keys: list[str] = field(default_factory=list)
    bad_ranges: list[str] = field(default_factory=list)
    correct_ranges: list[str] = field(default_factory=list)
    reflect_targets: list[dict[str, Any]] = field(default_factory=list)
    regions: list[str] = field(default_factory=list)
    region_hints: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    patch_jobs: list[dict[str, Any]] = field(default_factory=list)
    model_usage: dict[str, int] = field(default_factory=dict)


@dataclass
class ReflectCandidate:
    target_cell: str
    key: str
    candidate_value_cell: str = ""
    candidate_value: Any = None
    target_range: str = ""
    is_actual_edit: bool = False
    edited_label_cell: bool = False
    old_value: Any = None
    new_value: Any = None
    source: str = "key_scan"


@dataclass
class ReflectContext:
    workbook_path: Path
    instruction: str
    form: dict[str, Any]
    edit_log_jsonl: Path
    out_dir: Path
    model: str
    base_url: str | None
    api_key: str | None
    round_idx: int = 1
    sheet_name: str | None = None
    risk_threshold: int = 1
    expand_left_cols: int = 3
    expand_right_cols: int = 3
    expand_top_rows: int = 1
    expand_bottom_rows: int = 1
    hint_builder: Optional[Callable[[str, list[ReflectTarget], Path], list[dict[str, Any]]]] = None
    prior_context_summary: Optional[dict[str, Any]] = None
    screenshot_check: Optional[dict[str, Any]] = None
    first_pass_sheet_png: str = ""
    planner_policy_summary: str = ""
    planner_policy: Optional[dict[str, Any]] = None
    model_io_dir: Optional[Path] = None
    model_direct_ranges_only: bool = False


def _cell_has_border(cell) -> bool:
    border = cell.border
    if border is None:
        return False
    return any(getattr(getattr(border, side, None), "style", None) for side in ("top", "bottom", "left", "right"))


def _cell_has_fill(cell) -> bool:
    fill = cell.fill
    if fill is None:
        return False
    return bool(getattr(fill, "fill_type", None))


def _get_used_bbox(ws) -> tuple[int, int, int, int]:
    max_r = max(1, int(getattr(ws, "max_row", 1) or 1))
    max_c = max(1, int(getattr(ws, "max_column", 1) or 1))
    merged_cells: set[tuple[int, int]] = set()
    for rng in ws.merged_cells.ranges:
        for rr in range(rng.min_row, rng.max_row + 1):
            for cc in range(rng.min_col, rng.max_col + 1):
                merged_cells.add((rr, cc))

    min_row = min_col = None
    max_row = max_col = None
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            cell = ws.cell(row=r, column=c)
            has_content = cell.value is not None and str(cell.value).strip() != ""
            structural = _cell_has_border(cell) or _cell_has_fill(cell) or ((r, c) in merged_cells)
            if has_content or structural:
                if min_row is None:
                    min_row = max_row = r
                    min_col = max_col = c
                else:
                    min_row = min(min_row, r)
                    min_col = min(min_col, c)
                    max_row = max(max_row, r)
                    max_col = max(max_col, c)
    if min_row is None:
        return 1, 1, 1, 1
    return int(min_row), int(min_col), int(max_row), int(max_col)


def _get_form_bbox(
    ws,
    form: dict[str, Any],
    key_left_pad_cols: int = 8,
    right_pad_cols: int = 2,
    top_pad_rows: int = 2,
    bottom_pad_rows: int = 2,
) -> tuple[int, int, int, int] | None:
    used_min_row, used_min_col, used_max_row, used_max_col = _get_used_bbox(ws)
    rects: list[tuple[int, int, int, int]] = []
    for pair in form.get("pairs", []) or []:
        if not isinstance(pair, dict):
            continue
        value_cell = str(pair.get("value_cell") or "").strip()
        if not value_cell:
            continue
        for token in [x.strip() for x in value_cell.split(",") if x.strip()]:
            try:
                c1, r1, c2, r2 = range_boundaries(normalize_a1_range(token))
            except Exception:
                continue
            rects.append((r1, c1, r2, c2))
    if not rects:
        return None
    min_row = min(r1 for r1, _, _, _ in rects)
    min_col = min(c1 for _, c1, _, _ in rects)
    max_row = max(r2 for _, _, r2, _ in rects)
    max_col = max(c2 for _, _, _, c2 in rects)
    min_row = max(used_min_row, min_row - max(0, int(top_pad_rows)))
    min_col = max(used_min_col, min_col - max(0, int(key_left_pad_cols)))
    max_row = min(used_max_row, max_row + max(0, int(bottom_pad_rows)))
    max_col = min(used_max_col, max_col + max(0, int(right_pad_cols)))
    return int(min_row), int(min_col), int(max_row), int(max_col)


def _get_reflect_scan_bbox(ws, form: dict[str, Any]) -> tuple[int, int, int, int]:
    return _get_form_bbox(ws, form) or _get_used_bbox(ws)


def _merged_anchor_map(ws) -> dict[tuple[int, int], tuple[int, int, int, int]]:
    m: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for rng in ws.merged_cells.ranges:
        r0, c0, r1, c1 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        for rr in range(r0, r1 + 1):
            for cc in range(c0, c1 + 1):
                m[(rr, cc)] = (r0, c0, r1, c1)
    return m


def _build_sheet_structure(ws, form: dict[str, Any]) -> dict[str, Any]:
    min_row, min_col, max_row, max_col = _get_reflect_scan_bbox(ws, form)
    merged = _merged_anchor_map(ws)
    labels: dict[str, str] = {}
    input_cells: set[str] = set()
    input_candidates: set[str] = set()
    row_labels: dict[int, list[str]] = defaultdict(list)
    merged_ranges = {str(rng): str(rng.start_cell.coordinate) for rng in ws.merged_cells.ranges}

    for pair in form.get("pairs", []) or []:
        if not isinstance(pair, dict):
            continue
        value_cell = pair.get("value_cell")
        if not isinstance(value_cell, str) or not value_cell.strip():
            continue
        for cell_ref in iter_range_cells(value_cell):
            try:
                r, c = coord_of(cell_ref)
            except Exception:
                continue
            anchor = merged.get((r, c))
            if anchor:
                input_cells.add(f"{get_column_letter(anchor[1])}{anchor[0]}")
            else:
                input_cells.add(cell_ref.upper())

    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            a1 = f"{get_column_letter(c)}{r}"
            if isinstance(cell.value, str) and cell.value.strip():
                labels[a1] = cell.value.strip()
                row_labels[r].append(a1)
            tags = []
            border = cell.border
            if border is not None:
                for side_name, tag in (("top", "T"), ("bottom", "B"), ("left", "L"), ("right", "R")):
                    if getattr(getattr(border, side_name, None), "style", None):
                        tags.append(tag)
            is_merged = (r, c) in merged
            if set(tags) >= {"T", "B", "L", "R"} or "B" in tags or is_merged:
                anchor = merged.get((r, c))
                if anchor:
                    input_candidates.add(f"{get_column_letter(anchor[1])}{anchor[0]}")
                else:
                    input_candidates.add(a1)

    return {
        "labels": labels,
        "input_cells": input_cells,
        "input_candidates": input_candidates,
        "row_labels": row_labels,
        "merged_ranges": merged_ranges,
        "scan_bbox": (min_row, min_col, max_row, max_col),
    }


def _build_duplicate_risk_regions(
    ws,
    form: dict[str, Any],
    pad_left_cols: int,
    pad_right_cols: int,
    pad_top_rows: int,
    pad_bottom_rows: int,
) -> list[str]:
    merged = _merged_anchor_map(ws)
    value_to_cells: dict[str, list[str]] = defaultdict(list)
    seen_anchors: set[str] = set()
    min_row, min_col, max_row, max_col = _get_reflect_scan_bbox(ws, form)

    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            anchor = merged.get((r, c), (r, c, r, c))
            anchor_a1 = f"{get_column_letter(anchor[1])}{anchor[0]}"
            if anchor_a1 in seen_anchors:
                continue
            seen_anchors.add(anchor_a1)
            cell = ws.cell(row=anchor[0], column=anchor[1])
            value = cell.value
            if value is None:
                continue
            norm = normalize_key(value)
            if not norm:
                continue
            value_to_cells[norm].append(anchor_a1)

    risk_regions: list[str] = []
    for cells in value_to_cells.values():
        if len(cells) < 2:
            continue
        for cell_ref in cells:
            try:
                risk_regions.append(
                    expand_reflect_target(
                        ws,
                        cell_ref,
                        pad_left_cols=pad_left_cols,
                        pad_right_cols=pad_right_cols,
                        pad_top_rows=pad_top_rows,
                        pad_bottom_rows=pad_bottom_rows,
                    )
                )
            except Exception:
                continue
    return risk_regions


def _read_current_form_values(ws, form: dict[str, Any]) -> list[dict[str, Any]]:
    merged = _merged_anchor_map(ws)
    rows: list[dict[str, Any]] = []
    for pair in form.get("pairs", []) or []:
        if not isinstance(pair, dict):
            continue
        key = str(pair.get("key") or "").strip()
        value_cell = str(pair.get("value_cell") or "").strip()
        if not key or not value_cell:
            continue
        first_token = value_cell.split(",")[0].strip()
        try:
            c1, r1, _, _ = range_boundaries(normalize_a1_range(first_token))
        except Exception:
            continue
        anchor = merged.get((r1, c1), (r1, c1, r1, c1))
        anchor_cell = ws.cell(row=anchor[0], column=anchor[1]).coordinate
        rows.append(
            {
                "key": key,
                "value_cell": value_cell,
                "current_value": ws[anchor_cell].value,
            }
        )
    return rows


def _is_probable_key_text(text: str) -> bool:
    value = str(text or "").strip()
    if not value:
        return False
    compact = re.sub(r"\s+", " ", value)
    if len(compact) > 64:
        return False
    alnum = re.sub(r"[^A-Za-z0-9]+", "", compact)
    if len(alnum) <= 1:
        return False
    if re.fullmatch(r"[\d\W_]+", compact):
        return False
    if re.fullmatch(r"[$€¥]?\d[\d,./:-]*%?", compact):
        return False
    if compact.lower() in {"yes", "no", "n/a", "na"}:
        return False
    return True


def _anchor_a1(merged: dict[tuple[int, int], tuple[int, int, int, int]], row_idx: int, col_idx: int) -> str:
    anchor = merged.get((row_idx, col_idx), (row_idx, col_idx, row_idx, col_idx))
    return f"{get_column_letter(anchor[1])}{anchor[0]}"


def _read_anchor_value(ws, merged: dict[tuple[int, int], tuple[int, int, int, int]], cell_ref: str) -> Any:
    try:
        row_idx, col_idx = coord_of(cell_ref)
    except Exception:
        return None
    anchor_a1 = _anchor_a1(merged, row_idx, col_idx)
    return ws[anchor_a1].value


def _guess_value_cell_for_key(
    ws,
    structure: dict[str, Any],
    key_cell: str,
) -> str:
    try:
        row_idx, col_idx = coord_of(key_cell)
    except Exception:
        return ""
    merged = _merged_anchor_map(ws)
    labels = {str(x).upper() for x in (structure.get("labels") or {}).keys()}
    preferred = {str(x).upper() for x in (structure.get("input_cells") or set())}
    candidates = preferred | {str(x).upper() for x in (structure.get("input_candidates") or set())}

    search_positions: list[tuple[int, int]] = []
    for dc in range(1, 7):
        search_positions.append((row_idx, col_idx + dc))
    for dr in (1, 2):
        for dc in range(0, 5):
            search_positions.append((row_idx + dr, col_idx + dc))

    for rr, cc in search_positions:
        if rr < 1 or cc < 1:
            continue
        a1 = _anchor_a1(merged, rr, cc).upper()
        if a1 == key_cell.upper():
            continue
        if a1 in candidates:
            return a1

    for rr, cc in search_positions:
        if rr < 1 or cc < 1:
            continue
        a1 = _anchor_a1(merged, rr, cc).upper()
        if a1 == key_cell.upper() or a1 in labels:
            continue
        return a1
    return ""


def _build_reflect_candidates(
    ws,
    form: dict[str, Any],
    edits: list[EditRecord],
    structure: dict[str, Any],
) -> list[ReflectCandidate]:
    merged = _merged_anchor_map(ws)
    edit_by_cell = {str(edit.target_cell).upper(): edit for edit in edits}
    explicit_pairs_by_key: dict[str, str] = {}
    for pair in form.get("pairs", []) or []:
        if not isinstance(pair, dict):
            continue
        key = normalize_key(str(pair.get("key") or ""))
        value_cell = str(pair.get("value_cell") or "").strip()
        if key and value_cell and key not in explicit_pairs_by_key:
            explicit_pairs_by_key[key] = value_cell.split(",")[0].strip().upper()

    candidates: list[ReflectCandidate] = []
    seen: set[tuple[str, str]] = set()
    for key_cell, raw_key in (structure.get("labels") or {}).items():
        key_text = str(raw_key or "").strip()
        if not _is_probable_key_text(key_text):
            continue
        norm_key = normalize_key(key_text)
        candidate_value_cell = explicit_pairs_by_key.get(norm_key) or _guess_value_cell_for_key(ws, structure, key_cell)
        candidate_value = _read_anchor_value(ws, merged, candidate_value_cell) if candidate_value_cell else None
        is_actual_edit = False
        edited_label_cell = False
        old_value = None
        new_value = None
        if candidate_value_cell and candidate_value_cell.upper() in edit_by_cell:
            rec = edit_by_cell[candidate_value_cell.upper()]
            is_actual_edit = True
            old_value = rec.old_value
            new_value = rec.new_value
        elif key_cell.upper() in edit_by_cell:
            rec = edit_by_cell[key_cell.upper()]
            is_actual_edit = True
            edited_label_cell = True
            old_value = rec.old_value
            new_value = rec.new_value
        target_range = union_ranges([f"{key_cell}:{key_cell}", f"{candidate_value_cell}:{candidate_value_cell}"]) if candidate_value_cell else normalize_a1_range(key_cell)
        sig = (key_cell.upper(), candidate_value_cell.upper() if candidate_value_cell else "")
        if sig in seen:
            continue
        seen.add(sig)
        candidates.append(
            ReflectCandidate(
                target_cell=key_cell.upper(),
                key=key_text,
                candidate_value_cell=candidate_value_cell.upper() if candidate_value_cell else "",
                candidate_value=candidate_value,
                target_range=target_range,
                is_actual_edit=is_actual_edit,
                edited_label_cell=edited_label_cell,
                old_value=old_value,
                new_value=new_value,
                source="key_scan",
            )
        )
    return candidates


def _reason_implies_correct_state(reason: str) -> tuple[bool, bool]:
    text = str(reason or "").strip().lower()
    if not text:
        return False, False
    safe_blank = any(
        phrase in text
        for phrase in (
            "intentionally blank",
            "should remain blank",
            "correctly left blank",
        )
    )
    correct = safe_blank or any(
        phrase in text
        for phrase in (
            "correctly filled",
            "correctly placed",
            "already correct",
            "form is correctly filled",
            "value correctly filled",
        )
    )
    return correct, safe_blank


def _is_blankish_value(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return str(value).strip() == ""
    return False


def _soft_key_match_score(a: str, b: str) -> float:
    aa = normalize_key(a)
    bb = normalize_key(b)
    if not aa or not bb:
        return 0.0
    if aa == bb:
        return 1.0
    ratio = SequenceMatcher(None, aa, bb).ratio()
    ta = set(aa.split())
    tb = set(bb.split())
    overlap = (len(ta & tb) / max(1, len(ta | tb))) if (ta or tb) else 0.0
    contains = 1.0 if (aa in bb or bb in aa) and min(len(aa), len(bb)) >= 4 else 0.0
    return max(ratio, overlap, contains)


def _planner_blank_like_targets(planner_policy: dict[str, Any] | None) -> list[dict[str, str]]:
    plan = planner_policy if isinstance(planner_policy, dict) else {}
    targets: list[dict[str, str]] = []
    for item in plan.get("blank_policies", []) or []:
        if not isinstance(item, dict):
            continue
        target = str(item.get("target") or "").strip()
        reason = str(item.get("reason") or "").strip()
        if target:
            targets.append({"kind": "blank_policy", "target": target, "reason": reason})
    for item in plan.get("uncertain_fields", []) or []:
        if not isinstance(item, dict):
            continue
        target = str(item.get("target") or "").strip()
        action = str(item.get("action") or "").strip().lower()
        reason = str(item.get("reason") or "").strip()
        if target and action in {"leave_blank", "leave blank", "blank", "do_not_fill", "do not fill", "skip"}:
            targets.append({"kind": "uncertain_blank", "target": target, "reason": reason or action})
    return targets


def _planner_blank_policy_for_key(key_text: str, planner_policy: dict[str, Any] | None) -> dict[str, str] | None:
    best: dict[str, str] | None = None
    best_score = 0.0
    for item in _planner_blank_like_targets(planner_policy):
        score = _soft_key_match_score(key_text, str(item.get("target") or ""))
        if score > best_score:
            best_score = score
            best = item
    if best is None or best_score < 0.74:
        return None
    out = dict(best)
    out["match_score"] = f"{best_score:.4f}"
    return out


def _load_effective_edits(edit_log_jsonl: Path, sheet_name: str | None) -> list[EditRecord]:
    if not edit_log_jsonl.exists():
        return []
    latest_by_cell: dict[tuple[str, str], EditRecord] = {}
    with edit_log_jsonl.open("r", encoding="utf-8") as f:
        for line in f:
            raw = line.strip()
            if not raw:
                continue
            try:
                obj = json.loads(raw)
            except Exception:
                continue
            sheet = str(obj.get("sheet") or "")
            target_cell = str(obj.get("target_cell") or obj.get("cell") or "").upper()
            if not sheet or not target_cell:
                continue
            if sheet_name and sheet != sheet_name:
                continue
            rec = EditRecord(
                sheet=sheet,
                requested_cell=str(obj.get("requested_cell") or target_cell).upper(),
                target_cell=target_cell,
                old_value=obj.get("old_value"),
                new_value=obj.get("new_value"),
                round_idx=int(obj.get("round_idx", 0) or 0),
                step_idx=int(obj.get("step_idx", 0) or 0),
                source=str(obj.get("source") or "agent_write"),
                is_actual_edit=bool(obj.get("is_actual_edit", obj.get("old_value") != obj.get("new_value"))),
                key=str(obj.get("key") or ""),
            )
            if not rec.is_actual_edit:
                continue
            latest_by_cell[(sheet, target_cell)] = rec
    return list(latest_by_cell.values())


def _match_edit_key(edit: EditRecord, form: dict[str, Any]) -> str:
    if edit.key:
        return edit.key
    for pair in form.get("pairs", []) or []:
        if not isinstance(pair, dict):
            continue
        value_cell = pair.get("value_cell")
        key = pair.get("key")
        if not isinstance(value_cell, str) or not isinstance(key, str):
            continue
        for tok in [x.strip() for x in value_cell.split(",") if x.strip()]:
            try:
                if range_intersects(tok, edit.target_cell):
                    return key
            except Exception:
                continue
    return ""


def _programmatic_flags(
    candidate: ReflectCandidate,
    structure: dict[str, Any],
    screenshot_bad_ranges: list[str] | None = None,
) -> list[str]:
    flags: set[str] = set()
    target = candidate.target_cell.upper()
    if bool(candidate.edited_label_cell):
        flags.add("label_area")
    duplicate_risk_regions: list[str] = structure.get("duplicate_risk_regions", [])
    for region in duplicate_risk_regions:
        try:
            if range_intersects(candidate.target_range or target, region):
                flags.add("duplicate_pattern")
                break
        except Exception:
            continue
    for screenshot_bad_range in screenshot_bad_ranges or []:
        try:
            if range_intersects(candidate.target_range or target, screenshot_bad_range):
                flags.add("screenshot_bad_range_hit")
                break
        except Exception:
            continue
    return sorted(flags)


def _model_assess(
    ctx: ReflectContext,
    candidates: list[ReflectCandidate],
    filled_pairs: list[dict[str, Any]],
) -> tuple[dict[str, dict[str, bool]], list[str], list[str], list[str], str, dict[str, int]]:
    client = build_client(ctx.base_url, ctx.api_key)
    direct_range_clause = (
        "Also return bad_ranges: A1 ranges that should be repaired, and correct_ranges: A1 ranges that are already correct and should be protected.\n"
        if ctx.model_direct_ranges_only
        else ""
    )
    range_schema_clause = (
        ","
        "\"bad_ranges\":[\"A1:B2\"],"
        "\"correct_ranges\":[\"C3:D4\"]"
        if ctx.model_direct_ranges_only
        else ""
    )
    prompt = (
        "You are a reflect risk checker for Excel form filling.\n"
        "Evaluate the key candidates listed below using the filled form screenshot.\n"
        "This file is a form. Treat a form area as correct only when the correct key/label and the correct corresponding value both appear together in the right place; a region is not correct if only the key is correct or only the value is correct.\n"
        "If a value is written into a key/label/title/static-template cell, that candidate is not correct even if the value text itself looks plausible.\n"
        "If one sibling field contains a combined value while a neighboring sibling field still shows template label text or remains blank, mark the candidate as not correct; the correct information must be split across the proper form fields.\n"
        "If the supposed value cell still visibly shows template label text like '(First)' or other original form text, that is a missing/wrong value, not a correct region.\n"
        "If a field is intentionally blank and that blank state is correct, mark is_safe_blank=true and is_correct_region=true.\n"
        "For each candidate, return booleans for:\n"
        "- is_correct_region: whether this key and its corresponding value are already correct in the filled form\n"
        "- is_safe_blank: whether this field should intentionally remain blank and should be protected\n"
        "- template_inconsistency: whether this edit is inconsistent with similar filled fields in the same sheet, "
        "including style/pattern shifts such as font-like formatting conventions or background-color style changes inside repeated template sections\n"
        "- high_risk_pattern: whether this key-value fill appears to belong to a risky vertical/top-bottom form pattern where the value may be placed on the wrong row/cell\n"
        "- special_char_anomaly: whether this edit likely replaced language-specific accented characters with plain English letters\n"
        "Also return missing_keys: keys that appear to be required by the instruction but are still missing in the filled workbook.\n"
        f"{direct_range_clause}"
        "Return JSON only with schema:\n"
        "{"
        "\"assessments\":[{\"target_cell\":\"A1\",\"key\":\"...\",\"candidate_value_cell\":\"B1\",\"is_correct_region\":false,\"is_safe_blank\":false,\"template_inconsistency\":true,\"high_risk_pattern\":false,\"special_char_anomaly\":false,\"reason\":\"...\"}],"
        f"{range_schema_clause}"
        "\"missing_keys\":[\"...\"]"
        "}\n\n"
        "Few-shot guidance:\n"
        "- Wrong: candidate key is a first-name field, but the screenshot shows the neighboring last-name field contains 'Doe, Jane' while the first-name value cell still shows '(First)' or stays blank. Then is_correct_region=false and the reason should say the first name is missing / misplaced.\n"
        "- Wrong: a number or other value appears in a title/key/static-template cell from the blank form, while the actual value field is blank. Then is_correct_region=false even if the value text itself is valid.\n"
        "- Wrong: a static label cell such as 'CONTACT NAME:' becomes 'CONTACT NAME: Alex Rivera' in the filled form, while the neighboring value area remains empty. Then is_correct_region=false because the value was appended into the label cell instead of the separate value field.\n"
        "- Correct only when the key area and the value area line up in the correct role; nearby but misplaced content is still incorrect.\n\n"
        f"Instruction:\n{ctx.instruction}\n\n"
        f"Key candidates:\n{json.dumps([asdict(x) for x in candidates], ensure_ascii=False)}\n\n"
        f"Current filled pairs:\n{json.dumps(filled_pairs, ensure_ascii=False)}\n"
    )
    if ctx.prior_context_summary:
        prompt += f"\nPrior first-pass context summary:\n{json.dumps(ctx.prior_context_summary, ensure_ascii=False)}\n"
    if ctx.screenshot_check:
        prompt += f"\nBefore/after screenshot check summary:\n{json.dumps(ctx.screenshot_check, ensure_ascii=False)}\n"
    if str(ctx.planner_policy_summary or "").strip():
        prompt += f"\nPlanner execution policies:\n{ctx.planner_policy_summary}\n"
    message_content: list[dict[str, Any]] = [{"type": "text", "text": prompt}]
    after_png = str((ctx.screenshot_check or {}).get("after_sheet_png") or "").strip()
    before_png = str((ctx.screenshot_check or {}).get("before_sheet_png") or "").strip()
    try:
        if before_png and Path(before_png).exists():
            b64_before, mime_before = _encode_image_base64(before_png)
            message_content.append({"type": "text", "text": "Template / BEFORE screenshot:"})
            message_content.append({"type": "image_url", "image_url": {"url": f"data:image/{mime_before};base64,{b64_before}"}})
        if after_png and Path(after_png).exists():
            b64_after, mime_after = _encode_image_base64(after_png)
            message_content.append({"type": "text", "text": "Filled / AFTER screenshot:"})
            message_content.append({"type": "image_url", "image_url": {"url": f"data:image/{mime_after};base64,{b64_after}"}})
    except Exception:
        pass
    messages = [{"role": "user", "content": message_content}]
    resp = client.chat.completions.create(
        model=ctx.model,
        temperature=1.0,
        messages=messages,
    )
    raw = str(resp.choices[0].message.content or "").strip()
    usage = usage_to_dict(resp)
    _dump_model_io(
        ctx.model_io_dir,
        f"round_{int(ctx.round_idx):03d}_stage2_assess",
        request_payload={"model": ctx.model, "temperature": 1.0, "messages": messages},
        response_payload={"content": raw},
        raw_text=raw,
        usage=usage,
        extra={
            "candidate_count": len(candidates),
            "filled_pairs_count": len(filled_pairs),
            "sheet_name": ctx.sheet_name,
        },
    )
    payload = extract_json_object(raw)
    assessments_by_cell: dict[str, dict[str, bool]] = {}
    for item in payload.get("assessments", []) if isinstance(payload, dict) else []:
        if not isinstance(item, dict):
            continue
        cell = str(item.get("target_cell") or "").upper()
        if not cell:
            continue
        assessments_by_cell[cell] = {
            "is_correct_region": bool(item.get("is_correct_region")),
            "is_safe_blank": bool(item.get("is_safe_blank")),
            "template_inconsistency": bool(item.get("template_inconsistency")),
            "high_risk_pattern": bool(item.get("high_risk_pattern")),
            "special_char_anomaly": bool(item.get("special_char_anomaly")),
            "reason": str(item.get("reason") or ""),
        }
    missing_keys = []
    for item in payload.get("missing_keys", []) if isinstance(payload, dict) else []:
        if isinstance(item, str) and item.strip():
            missing_keys.append(item.strip())
    bad_ranges: list[str] = []
    correct_ranges: list[str] = []
    for item in payload.get("bad_ranges", []) if isinstance(payload, dict) else []:
        try:
            bad_ranges.append(normalize_a1_range(str(item or "").strip()))
        except Exception:
            continue
    for item in payload.get("correct_ranges", []) if isinstance(payload, dict) else []:
        try:
            correct_ranges.append(normalize_a1_range(str(item or "").strip()))
        except Exception:
            continue
    return assessments_by_cell, missing_keys, bad_ranges, correct_ranges, raw, usage


def expand_reflect_target(
    ws,
    target_range: str,
    pad_left_cols: int,
    pad_right_cols: int,
    pad_top_rows: int,
    pad_bottom_rows: int,
) -> str:
    merged = _merged_anchor_map(ws)
    min_row, min_col, max_row, max_col = _get_used_bbox(ws)
    base_ranges: list[str] = []
    for cell_ref in iter_range_cells(target_range):
        try:
            row_idx, col_idx = coord_of(cell_ref)
        except Exception:
            continue
        anchor = merged.get((row_idx, col_idx))
        if anchor:
            base_ranges.append(
                f"{get_column_letter(anchor[1])}{anchor[0]}:{get_column_letter(anchor[3])}{anchor[2]}"
            )
        else:
            base_ranges.append(f"{cell_ref}:{cell_ref}")
    if not base_ranges:
        base_ranges = [normalize_a1_range(target_range)]
    c1, r1, c2, r2 = range_boundaries(union_ranges(base_ranges))
    c1 = max(min_col, c1 - max(0, int(pad_left_cols)))
    c2 = min(max_col, c2 + max(0, int(pad_right_cols)))
    r1 = max(min_row, r1 - max(0, int(pad_top_rows)))
    r2 = min(max_row, r2 + max(0, int(pad_bottom_rows)))
    return f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"


def merge_touching_regions(regions: list[str]) -> list[str]:
    pending = [normalize_a1_range(x) for x in regions if str(x or "").strip()]
    changed = True
    while changed:
        changed = False
        out: list[str] = []
        while pending:
            cur = pending.pop(0)
            c1, r1, c2, r2 = range_boundaries(cur)
            merged_any = False
            for idx, other in enumerate(pending):
                oc1, or1, oc2, or2 = range_boundaries(other)
                touch = not (c2 + 1 < oc1 or oc2 + 1 < c1 or r2 + 1 < or1 or or2 + 1 < r1)
                if touch:
                    cur = (
                        f"{get_column_letter(min(c1, oc1))}{min(r1, or1)}:"
                        f"{get_column_letter(max(c2, oc2))}{max(r2, or2)}"
                    )
                    pending.pop(idx)
                    changed = True
                    merged_any = True
                    break
            if merged_any:
                pending.insert(0, cur)
            else:
                out.append(cur)
        pending = out
    seen: set[str] = set()
    deduped: list[str] = []
    for item in pending:
        if item not in seen:
            seen.add(item)
            deduped.append(item)
    return deduped


def run_reflect_round(context: ReflectContext) -> ReflectRoundResult:
    wb = openpyxl.load_workbook(context.workbook_path)
    try:
        ws = wb[context.sheet_name] if context.sheet_name else wb.active
        structure = _build_sheet_structure(ws, context.form)
        structure["duplicate_risk_regions"] = _build_duplicate_risk_regions(
            ws,
            form=context.form,
            pad_left_cols=context.expand_left_cols,
            pad_right_cols=context.expand_right_cols,
            pad_top_rows=context.expand_top_rows,
            pad_bottom_rows=context.expand_bottom_rows,
        )
        filled_pairs = _read_current_form_values(ws, context.form)
        edits = _load_effective_edits(context.edit_log_jsonl, ws.title)
        for edit in edits:
            if not edit.key:
                edit.key = _match_edit_key(edit, context.form)

        candidates = _build_reflect_candidates(ws, context.form, edits, structure)
        model_flags, missing_keys, model_bad_ranges, model_correct_ranges, model_raw, model_usage = _model_assess(context, candidates, filled_pairs)
        screenshot_bad_ranges: list[str] = []
        screenshot_need_fix = False
        screenshot_correct_ranges: list[str] = []
        if isinstance(context.screenshot_check, dict):
            screenshot_need_fix = bool(context.screenshot_check.get("need_fix"))
            if screenshot_need_fix:
                for item in context.screenshot_check.get("bad_ranges", []) or []:
                    try:
                        screenshot_bad_ranges.append(normalize_a1_range(str(item or "").strip()))
                    except Exception:
                        continue
                if not screenshot_bad_ranges:
                    legacy_bad_range = str(context.screenshot_check.get("bad_range") or "").strip()
                    if legacy_bad_range:
                        try:
                            screenshot_bad_ranges.append(normalize_a1_range(legacy_bad_range))
                        except Exception:
                            pass
            for item in context.screenshot_check.get("correct_ranges", []) or []:
                try:
                    screenshot_correct_ranges.append(normalize_a1_range(str(item or "").strip()))
                except Exception:
                    continue

        assessments: list[EditRiskAssessment] = []
        if context.model_direct_ranges_only:
            screenshot_bad_ranges = []
            screenshot_correct_ranges = []
            screenshot_need_fix = False

        bad_ranges: list[str] = list(model_bad_ranges if context.model_direct_ranges_only else screenshot_bad_ranges)
        correct_ranges: list[str] = list(model_correct_ranges if context.model_direct_ranges_only else screenshot_correct_ranges)
        for candidate in candidates:
            flags = (
                []
                if context.model_direct_ranges_only
                else _programmatic_flags(
                    candidate,
                    structure,
                    screenshot_bad_ranges=screenshot_bad_ranges,
                )
            )
            cell_flags = model_flags.get(candidate.target_cell.upper(), {})
            is_correct_region = bool(cell_flags.get("is_correct_region"))
            is_safe_blank = bool(cell_flags.get("is_safe_blank"))
            candidate_range = candidate.target_range or candidate.target_cell
            planner_blank_rule = _planner_blank_policy_for_key(candidate.key, context.planner_policy)
            if any(range_intersects(candidate_range, rng) for rng in correct_ranges):
                is_correct_region = True
            if (not context.model_direct_ranges_only) and cell_flags.get("template_inconsistency"):
                flags.append("template_inconsistency")
            if (not context.model_direct_ranges_only) and cell_flags.get("high_risk_pattern"):
                flags.append("high_risk_pattern")
            if (not context.model_direct_ranges_only) and cell_flags.get("special_char_anomaly"):
                flags.append("special_char_anomaly")
            planner_reason = ""
            if planner_blank_rule:
                target_name = str(planner_blank_rule.get("target") or candidate.key or "").strip()
                rule_reason = str(planner_blank_rule.get("reason") or "").strip()
                if _is_blankish_value(candidate.candidate_value):
                    is_correct_region = True
                    is_safe_blank = True
                    planner_reason = f"Planner policy says '{target_name}' should remain blank."
                else:
                    flags.append("planner_blank_violation")
                    planner_reason = f"Planner policy says '{target_name}' should remain blank, but the field is filled."
                if rule_reason:
                    planner_reason += f" Reason: {rule_reason}."
            flags = sorted(set(flags))
            hard_programmatic_flags = set() if context.model_direct_ranges_only else {flag for flag in flags if flag == "label_area"}
            reason_text = str(cell_flags.get("reason") or "")
            if planner_reason:
                reason_text = (planner_reason + (" " + reason_text if reason_text else "")).strip()
            if "label_area" in hard_programmatic_flags:
                is_correct_region = False
                is_safe_blank = False
                if not reason_text:
                    reason_text = (
                        "A value was written into a detected key/label cell. "
                        "Treat this as a label-area overwrite that must be repaired."
                    )
            reason_correct, reason_safe_blank = _reason_implies_correct_state(reason_text)
            if reason_correct and not hard_programmatic_flags:
                is_correct_region = True
            if reason_safe_blank and not hard_programmatic_flags:
                is_safe_blank = True
            if (is_correct_region or is_safe_blank) and not hard_programmatic_flags:
                flags = []
            score = len(flags)
            if is_correct_region or is_safe_blank:
                try:
                    correct_ranges.append(normalize_a1_range(candidate_range))
                except Exception:
                    pass
            direct_bad_hit = any(range_intersects(candidate_range, rng) for rng in bad_ranges) if context.model_direct_ranges_only else False
            assessments.append(
                EditRiskAssessment(
                    target_cell=candidate.target_cell,
                    key=candidate.key,
                    score=score,
                    risk_flags=flags,
                    is_high_risk=(
                        (
                            direct_bad_hit
                            if context.model_direct_ranges_only
                            else (
                                not is_correct_region
                                and not is_safe_blank
                                and ("label_area" in hard_programmatic_flags or score >= max(1, int(context.risk_threshold)))
                            )
                        )
                    ),
                    is_actual_edit=bool(candidate.is_actual_edit),
                    is_correct_region=is_correct_region,
                    is_safe_blank=is_safe_blank,
                    candidate_value_cell=candidate.candidate_value_cell,
                    candidate_value=candidate.candidate_value,
                    target_range=candidate_range,
                    reason=reason_text,
                )
            )

        targets: list[ReflectTarget] = []
        if context.model_direct_ranges_only:
            for model_bad_range in bad_ranges:
                targets.append(
                    ReflectTarget(
                        kind="model_bad_range",
                        sheet=ws.title,
                        key="",
                        target_range=model_bad_range,
                        source_cell=model_bad_range,
                        risk_flags=["model_bad_range"],
                    )
                )
        else:
            for item in assessments:
                candidate_range = item.target_range or item.target_cell
                if item.is_correct_region or item.is_safe_blank:
                    continue
                if "label_area" in (item.risk_flags or []):
                    targets.append(
                        ReflectTarget(
                            kind="label_area_edit",
                            sheet=ws.title,
                            key=item.key,
                            target_range=candidate_range,
                            source_cell=item.target_cell,
                            risk_flags=item.risk_flags,
                        )
                    )
                if item.is_high_risk:
                    targets.append(
                        ReflectTarget(
                            kind="high_risk_edit",
                            sheet=ws.title,
                            key=item.key,
                            target_range=candidate_range,
                            source_cell=item.target_cell,
                            risk_flags=item.risk_flags,
                        )
                    )

        if (not context.model_direct_ranges_only) and screenshot_need_fix:
            for screenshot_bad_range in screenshot_bad_ranges:
                targets.append(
                    ReflectTarget(
                        kind="screenshot_bad_range",
                        sheet=ws.title,
                        key="",
                        target_range=screenshot_bad_range,
                        source_cell=screenshot_bad_range,
                        risk_flags=["screenshot_need_fix"],
                    )
                )

        region_to_targets: dict[str, list[ReflectTarget]] = defaultdict(list)
        initial_regions: list[str] = []
        for target in targets:
            expanded = expand_reflect_target(
                ws=ws,
                target_range=target.target_range,
                pad_left_cols=context.expand_left_cols,
                pad_right_cols=context.expand_right_cols,
                pad_top_rows=context.expand_top_rows,
                pad_bottom_rows=context.expand_bottom_rows,
            )
            initial_regions.append(expanded)
            region_to_targets[expanded].append(target)

        merged_regions = merge_touching_regions(initial_regions)
        final_region_targets: dict[str, list[ReflectTarget]] = defaultdict(list)
        for region in merged_regions:
            for original_region, original_targets in region_to_targets.items():
                try:
                    if range_intersects(region, original_region):
                        final_region_targets[region].extend(original_targets)
                except Exception:
                    continue

        region_hints: dict[str, list[dict[str, Any]]] = {}
        patch_jobs: list[dict[str, Any]] = []
        for idx, region in enumerate(merged_regions, start=1):
            dedup_targets: list[ReflectTarget] = []
            seen_target_keys: set[tuple[str, str, str]] = set()
            for target in final_region_targets.get(region, []):
                sig = (target.kind, target.key, target.target_range)
                if sig in seen_target_keys:
                    continue
                seen_target_keys.add(sig)
                dedup_targets.append(target)
            hints: list[dict[str, Any]] = []
            if context.hint_builder is not None:
                hints = context.hint_builder(region, dedup_targets, context.out_dir / f"region_{idx}") or []
            region_hints[region] = hints
            patch_jobs.append(
                asdict(
                    PatchJob(
                        region=region,
                        region_hints=hints,
                        targets=[asdict(x) for x in dedup_targets],
                    )
                )
            )

        result = ReflectRoundResult(
            edits_considered=[asdict(x) for x in candidates],
            risk_assessments=[asdict(x) for x in assessments],
            missing_keys=missing_keys,
            bad_ranges=sorted(set(bad_ranges)),
            correct_ranges=sorted(set(correct_ranges)),
            reflect_targets=[asdict(x) for x in targets],
            regions=merged_regions,
            region_hints=region_hints,
            patch_jobs=patch_jobs,
            model_usage=model_usage,
        )
        context.out_dir.mkdir(parents=True, exist_ok=True)
        payload = asdict(result)
        payload["model_raw"] = model_raw
        payload["sheet"] = ws.title
        (context.out_dir / f"reflect_round_{context.round_idx}_model_raw.txt").write_text(
            str(model_raw or ""),
            encoding="utf-8",
        )
        (context.out_dir / f"reflect_round_{context.round_idx}_parsed.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
        (context.out_dir / f"reflect_round_{context.round_idx}.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
        return result
    finally:
        wb.close()
