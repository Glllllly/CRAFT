#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
import re
import statistics
import unicodedata
from bisect import bisect_left
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


THIN_GRAY = Side(style="thin", color="BFBFBF")
THIN_TABLE = Side(style="thin", color="7A7A7A")
TABLE_BORDER = Border(left=THIN_TABLE, right=THIN_TABLE, top=THIN_TABLE, bottom=THIN_TABLE)
BLOCK_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
TITLE_FILL = PatternFill(fill_type="solid", fgColor="EAF2FF")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="F3F6F9")


def _ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def _safe_text(value: Any) -> str:
    text = re.sub(r"\s+\n", "\n", str(value or ""))
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _median(values: Iterable[float], default: float) -> float:
    clean = [float(v) for v in values if v is not None and not math.isnan(float(v))]
    if not clean:
        return default
    return float(statistics.median(clean))


def _cluster_positions(values: Iterable[float], tol: float) -> list[float]:
    raw = sorted(float(v) for v in values if v is not None and not math.isnan(float(v)))
    if not raw:
        return []
    groups: list[list[float]] = [[raw[0]]]
    for value in raw[1:]:
        if abs(value - groups[-1][-1]) <= tol:
            groups[-1].append(value)
        else:
            groups.append([value])
    return [sum(group) / len(group) for group in groups]


def _collect_edge_values(elements: list[dict[str, Any]]) -> tuple[list[float], list[float]]:
    x_values: list[float] = []
    y_values: list[float] = []
    for el in elements:
        x0, y0, x1, y1 = el["bbox"]
        x_values.extend([x0, x1])
        y_values.extend([y0, y1])
    return x_values, y_values


class _HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.in_table = False
        self.rows: list[list[dict[str, Any]]] = []
        self.current_row: list[dict[str, Any]] | None = None
        self.current_cell: dict[str, Any] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        attr_map = {k.lower(): v for k, v in attrs}
        if tag == "table" and not self.in_table:
            self.in_table = True
            return
        if not self.in_table:
            return
        if tag == "tr":
            self.current_row = []
            self.rows.append(self.current_row)
            return
        if tag in {"td", "th"} and self.current_row is not None:
            rowspan = int(attr_map.get("rowspan") or 1)
            colspan = int(attr_map.get("colspan") or 1)
            self.current_cell = {"text_parts": [], "rowspan": rowspan, "colspan": colspan}
            self.current_row.append(self.current_cell)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "table" and self.in_table:
            self.in_table = False
            return
        if not self.in_table:
            return
        if tag in {"td", "th"} and self.current_cell is not None:
            self.current_cell["text"] = _safe_text("".join(self.current_cell["text_parts"]))
            self.current_cell = None
            return
        if tag == "tr":
            self.current_row = None

    def handle_data(self, data: str) -> None:
        if self.in_table and self.current_cell is not None:
            self.current_cell["text_parts"].append(data)


def _parse_html_table(html: str) -> tuple[dict[tuple[int, int], str], list[tuple[int, int, int, int]], int, int]:
    parser = _HTMLTableParser()
    parser.feed(html or "")
    rows = parser.rows
    if not rows:
        return {}, [], 0, 0
    values: dict[tuple[int, int], str] = {}
    merges: list[tuple[int, int, int, int]] = []
    occupied: set[tuple[int, int]] = set()
    max_row = 0
    max_col = 0

    for r_idx, row in enumerate(rows, start=1):
        c_idx = 1
        for cell in row:
            while (r_idx, c_idx) in occupied:
                c_idx += 1
            rowspan = int(cell.get("rowspan", 1) or 1)
            colspan = int(cell.get("colspan", 1) or 1)
            text = _safe_text(cell.get("text"))
            values[(r_idx, c_idx)] = text
            if rowspan > 1 or colspan > 1:
                merges.append((r_idx, c_idx, r_idx + rowspan - 1, c_idx + colspan - 1))
            for rr in range(r_idx, r_idx + rowspan):
                for cc in range(c_idx, c_idx + colspan):
                    occupied.add((rr, cc))
                    max_row = max(max_row, rr)
                    max_col = max(max_col, cc)
            c_idx += colspan
    return values, merges, max_row, max_col


def _infer_table_col_widths(cell_box_list: list[Any], expected_cols: int) -> list[float]:
    boxes = [_normalize_bbox(box) for box in cell_box_list]
    boxes = [box for box in boxes if box]
    if not boxes or expected_cols <= 0:
        return [16.0] * max(1, expected_cols)
    x_values: list[float] = []
    for box in boxes:
        x_values.extend([box[0], box[2]])
    x_edges = _cluster_positions(x_values, tol=3.0)
    if len(x_edges) >= expected_cols + 1:
        spans = [max(12.0, x_edges[i + 1] - x_edges[i]) for i in range(expected_cols)]
        return [min(40.0, max(6.0, span / 7.0)) for span in spans[:expected_cols]]
    return [16.0] * expected_cols


def _text_display_width(text: str) -> float:
    width = 0.0
    for ch in str(text or ""):
        if ch == "\n":
            continue
        if unicodedata.east_asian_width(ch) in {"F", "W"}:
            width += 2.0
        else:
            width += 1.0
    return width


def _infer_content_col_widths(
    values: dict[tuple[int, int], str],
    merges: list[tuple[int, int, int, int]],
    expected_cols: int,
) -> list[float]:
    if expected_cols <= 0:
        return [16.0]

    merge_map = {(r0, c0): (r1, c1) for r0, c0, r1, c1 in merges}
    col_widths = [8.0] * expected_cols

    for (row, col), raw_text in values.items():
        text = _safe_text(raw_text)
        if not text:
            continue
        end_row, end_col = merge_map.get((row, col), (row, col))
        span_cols = max(1, end_col - col + 1)
        line_width = max((_text_display_width(line) for line in text.splitlines()), default=0.0)
        # Excel width is roughly char count plus a small padding.
        est_width = min(40.0, max(6.0, line_width / span_cols + 2.0))
        for col_idx in range(col - 1, min(expected_cols, end_col)):
            col_widths[col_idx] = max(col_widths[col_idx], est_width)
    return col_widths


def _resolve_table_col_widths(
    cell_box_list: list[Any],
    values: dict[tuple[int, int], str],
    merges: list[tuple[int, int, int, int]],
    expected_cols: int,
    mode: str,
) -> list[float]:
    geometry_widths = _infer_table_col_widths(cell_box_list, expected_cols=expected_cols)
    if mode == "geometry":
        return geometry_widths
    content_widths = _infer_content_col_widths(values, merges, expected_cols=expected_cols)
    return [
        min(45.0, max(geometry_widths[idx], content_widths[idx]))
        for idx in range(min(len(geometry_widths), len(content_widths)))
    ]


def _collect_page_flow_blocks(page_payload: dict[str, Any], column_width_mode: str) -> tuple[list[dict[str, Any]], int]:
    parsing_res_list = page_payload.get("parsing_res_list") or []
    table_res_list = page_payload.get("table_res_list") or []
    blocks: list[dict[str, Any]] = []
    table_idx = 0
    max_table_cols = 1

    for block in sorted(
        parsing_res_list,
        key=lambda b: (
            _normalize_bbox(b.get("block_bbox"))[1] if _normalize_bbox(b.get("block_bbox")) else 0.0,
            _normalize_bbox(b.get("block_bbox"))[0] if _normalize_bbox(b.get("block_bbox")) else 0.0,
        ),
    ):
        label = str(block.get("block_label") or "").strip()
        bbox = _normalize_bbox(block.get("block_bbox"))
        if not bbox:
            continue
        if "table" in label.lower():
            table_res = table_res_list[table_idx] if table_idx < len(table_res_list) else {}
            html = str(table_res.get("pred_html") or block.get("block_content") or "")
            values, merges, max_row, max_col = _parse_html_table(html)
            col_widths = _resolve_table_col_widths(
                table_res.get("cell_box_list") or [],
                values=values,
                merges=merges,
                expected_cols=max_col,
                mode=column_width_mode,
            )
            max_table_cols = max(max_table_cols, max_col)
            blocks.append(
                {
                    "kind": "table_block",
                    "label": label or "table",
                    "bbox": bbox,
                    "values": values,
                    "merges": merges,
                    "rows": max_row,
                    "cols": max_col,
                    "col_widths": col_widths,
                    "text": "",
                }
            )
            table_idx += 1
        else:
            text = _safe_text(block.get("block_content")) or _label_placeholder(label)
            blocks.append(
                {
                    "kind": "text_block",
                    "label": label or "text",
                    "bbox": bbox,
                    "text": text,
                }
            )
    return blocks, max_table_cols


def _normalize_bbox(box: Any) -> list[float] | None:
    if box is None:
        return None
    if isinstance(box, list) and len(box) == 4 and not isinstance(box[0], list):
        x0, y0, x1, y1 = (float(v) for v in box)
        return [min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1)]
    if isinstance(box, (list, tuple)) and box and isinstance(box[0], (list, tuple)):
        xs: list[float] = []
        ys: list[float] = []
        for point in box:
            if len(point) >= 2:
                xs.append(float(point[0]))
                ys.append(float(point[1]))
        if xs and ys:
            return [min(xs), min(ys), max(xs), max(ys)]
    return None


def _center_in_bbox(cx: float, cy: float, bbox: list[float]) -> bool:
    return bbox[0] <= cx <= bbox[2] and bbox[1] <= cy <= bbox[3]


def _bbox_overlap_ratio(a: list[float], b: list[float]) -> float:
    x0 = max(a[0], b[0])
    y0 = max(a[1], b[1])
    x1 = min(a[2], b[2])
    y1 = min(a[3], b[3])
    if x1 <= x0 or y1 <= y0:
        return 0.0
    inter = (x1 - x0) * (y1 - y0)
    area_a = max(1.0, (a[2] - a[0]) * (a[3] - a[1]))
    return inter / area_a


def _join_texts_by_layout(items: list[dict[str, Any]]) -> str:
    if not items:
        return ""
    items = sorted(items, key=lambda it: (it["bbox"][1], it["bbox"][0]))
    heights = [max(1.0, it["bbox"][3] - it["bbox"][1]) for it in items]
    line_tol = max(8.0, _median(heights, default=20.0) * 0.6)
    lines: list[list[dict[str, Any]]] = []
    for item in items:
        if not lines:
            lines.append([item])
            continue
        prev_line = lines[-1]
        prev_cy = sum((it["bbox"][1] + it["bbox"][3]) / 2.0 for it in prev_line) / len(prev_line)
        cy = (item["bbox"][1] + item["bbox"][3]) / 2.0
        if abs(cy - prev_cy) <= line_tol:
            prev_line.append(item)
        else:
            lines.append([item])
    line_texts: list[str] = []
    for line in lines:
        line = sorted(line, key=lambda it: it["bbox"][0])
        parts = [line[0]["text"]]
        for prev, curr in zip(line, line[1:]):
            gap = curr["bbox"][0] - prev["bbox"][2]
            if gap > max(4.0, min(prev["bbox"][3] - prev["bbox"][1], curr["bbox"][3] - curr["bbox"][1]) * 0.35):
                parts.append(" ")
            parts.append(curr["text"])
        line_texts.append("".join(parts).strip())
    return "\n".join(text for text in line_texts if text)


def _doc_id_from_summary(summary_path: Path) -> str:
    return summary_path.parent.name


def _safe_filename_stem(text: str) -> str:
    stem = re.sub(r"[^\w.-]+", "_", str(text or "").strip(), flags=re.UNICODE)
    stem = stem.strip("._")
    return stem or "document"


def _resolve_ref(path_text: str, base_dir: Path) -> Path | None:
    raw = Path(path_text)
    candidates = []
    if raw.is_absolute():
        candidates.append(raw)
    else:
        candidates.append(raw)
        candidates.append(base_dir / raw)
        candidates.append(Path.cwd() / raw)
        candidates.append(base_dir / raw.name)
        candidates.append(base_dir / "json" / raw.name)
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def discover_summaries(input_path: Path, recursive: bool) -> list[Path]:
    if input_path.is_file():
        if input_path.name.lower() == "summary.json":
            return [input_path]
        raise ValueError(f"Expected a summary.json file or a directory, got: {input_path}")
    if not input_path.is_dir():
        raise FileNotFoundError(f"Input path not found: {input_path}")
    iterator = input_path.rglob("summary.json") if recursive else input_path.glob("*/summary.json")
    return sorted(p for p in iterator if p.is_file())


def _load_page_payload(page_json_path: Path) -> dict[str, Any]:
    payload = json.loads(page_json_path.read_text(encoding="utf-8"))
    if isinstance(payload, dict) and "res" in payload and isinstance(payload["res"], dict):
        return payload["res"]
    if not isinstance(payload, dict):
        raise ValueError(f"Unexpected page json payload type in {page_json_path}: {type(payload)!r}")
    return payload


def _label_placeholder(label: str) -> str:
    normalized = str(label or "").strip().lower()
    if "image" in normalized or "figure" in normalized:
        return "[image]"
    if "chart" in normalized:
        return "[chart]"
    if "formula" in normalized:
        return "[formula]"
    if "seal" in normalized:
        return "[seal]"
    return f"[{normalized or 'block'}]"


def _style_descriptor(label: str, text: str, kind: str) -> dict[str, Any]:
    normalized = str(label or "").strip().lower()
    style = {
        "font": Font(name="Calibri", size=11),
        "alignment": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "border": BLOCK_BORDER if kind != "table_cell" else TABLE_BORDER,
        "fill": None,
    }
    if kind == "table_cell":
        if len(text) <= 20 and "\n" not in text:
            style["alignment"] = Alignment(horizontal="center", vertical="center", wrap_text=True)
        return style
    if "doc_title" in normalized:
        style["font"] = Font(name="Calibri", size=16, bold=True)
        style["alignment"] = Alignment(horizontal="center", vertical="center", wrap_text=True)
        style["fill"] = TITLE_FILL
    elif "title" in normalized or "header" in normalized:
        style["font"] = Font(name="Calibri", size=13, bold=True)
        style["alignment"] = Alignment(horizontal="center", vertical="center", wrap_text=True)
        style["fill"] = HEADER_FILL
    elif "caption" in normalized or "footer" in normalized or "reference" in normalized:
        style["font"] = Font(name="Calibri", size=10, italic=True)
    elif normalized in {"text", "paragraph", "content"}:
        style["alignment"] = Alignment(horizontal="left", vertical="top", wrap_text=True)
    return style


def _set_sheet_base_widths(ws: openpyxl.worksheet.worksheet.Worksheet, total_cols: int, widths: list[float] | None = None) -> None:
    widths = widths or []
    for col_idx in range(1, total_cols + 1):
        width = widths[col_idx - 1] if col_idx - 1 < len(widths) else 16.0
        ws.column_dimensions[get_column_letter(col_idx)].width = round(float(width), 2)


def _apply_text_block(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    start_row: int,
    total_cols: int,
    block: dict[str, Any],
) -> int:
    text = _safe_text(block["text"])
    if not text:
        return start_row
    style = _style_descriptor(block["label"], text, kind="layout_block")
    cell = ws.cell(row=start_row, column=1, value=text)
    cell.font = style["font"]
    cell.alignment = style["alignment"]
    cell.border = style["border"]
    if style["fill"] is not None:
        cell.fill = style["fill"]
    for col in range(1, total_cols + 1):
        ws.cell(row=start_row, column=col).border = style["border"]
        if style["fill"] is not None:
            ws.cell(row=start_row, column=col).fill = style["fill"]
    if total_cols > 1:
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=total_cols)
    bbox = block.get("bbox") or [0, 0, 0, 0]
    row_height = min(60.0, max(20.0, (float(bbox[3]) - float(bbox[1])) * 0.85))
    ws.row_dimensions[start_row].height = round(row_height, 2)
    return start_row + 2


def _apply_table_block(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    start_row: int,
    total_cols: int,
    block: dict[str, Any],
) -> int:
    cols = max(1, int(block.get("cols") or 1))
    values = block.get("values") or {}
    merges = block.get("merges") or []
    for r in range(1, int(block.get("rows") or 0) + 1):
        ws.row_dimensions[start_row + r - 1].height = 24
    for r in range(1, int(block.get("rows") or 0) + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=start_row + r - 1, column=c)
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if r == 1:
                cell.font = Font(name="Calibri", size=11, bold=True)
                cell.fill = HEADER_FILL
            else:
                cell.font = Font(name="Calibri", size=11)
    for (r, c), text in values.items():
        cell = ws.cell(row=start_row + r - 1, column=c, value=text)
        cell.border = TABLE_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if r == 1:
            cell.font = Font(name="Calibri", size=11, bold=True)
            cell.fill = HEADER_FILL
        else:
            cell.font = Font(name="Calibri", size=11)
    for r0, c0, r1, c1 in merges:
        ws.merge_cells(
            start_row=start_row + r0 - 1,
            start_column=c0,
            end_row=start_row + r1 - 1,
            end_column=c1,
        )
    if total_cols > cols:
        for r in range(start_row, start_row + int(block.get("rows") or 0)):
            for c in range(cols + 1, total_cols + 1):
                ws.cell(row=r, column=c).border = BLOCK_BORDER
    return start_row + int(block.get("rows") or 0) + 1


def _collect_table_elements(page_payload: dict[str, Any], table_blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    table_res_list = page_payload.get("table_res_list") or []
    out: list[dict[str, Any]] = []
    for table_idx, table_res in enumerate(table_res_list):
        table_block = table_blocks[table_idx] if table_idx < len(table_blocks) else {}
        cell_boxes = table_res.get("cell_box_list") or []
        ocr_pred = table_res.get("table_ocr_pred") or {}
        rec_boxes_raw = ocr_pred.get("rec_boxes") or []
        rec_texts = ocr_pred.get("rec_texts") or []
        rec_scores = ocr_pred.get("rec_scores") or []

        rec_items: list[dict[str, Any]] = []
        for i, raw_box in enumerate(rec_boxes_raw):
            bbox = _normalize_bbox(raw_box)
            text = _safe_text(rec_texts[i] if i < len(rec_texts) else "")
            score = float(rec_scores[i] if i < len(rec_scores) else 0.0)
            if not bbox or not text:
                continue
            rec_items.append({"bbox": bbox, "text": text, "score": score, "used": False})

        for cell_idx, raw_cell_box in enumerate(cell_boxes):
            cell_bbox = _normalize_bbox(raw_cell_box)
            if not cell_bbox:
                continue
            matched: list[dict[str, Any]] = []
            for item in rec_items:
                cx = (item["bbox"][0] + item["bbox"][2]) / 2.0
                cy = (item["bbox"][1] + item["bbox"][3]) / 2.0
                if _center_in_bbox(cx, cy, cell_bbox) or _bbox_overlap_ratio(item["bbox"], cell_bbox) >= 0.5:
                    matched.append(item)
                    item["used"] = True
            cell_text = _join_texts_by_layout(matched)
            out.append(
                {
                    "kind": "table_cell",
                    "label": "table_cell",
                    "bbox": cell_bbox,
                    "text": cell_text,
                    "score": max([float(it["score"]) for it in matched], default=1.0),
                    "table_index": table_idx,
                    "block_id": table_block.get("block_id"),
                    "cell_index": cell_idx,
                }
            )

        unmatched = [item for item in rec_items if not item["used"]]
        if unmatched:
            for extra_idx, item in enumerate(unmatched):
                out.append(
                    {
                        "kind": "table_text_fallback",
                        "label": "table_text_fallback",
                        "bbox": item["bbox"],
                        "text": item["text"],
                        "score": float(item["score"]),
                        "table_index": table_idx,
                        "block_id": table_block.get("block_id"),
                        "cell_index": len(cell_boxes) + extra_idx,
                    }
                )
    return out


def _collect_page_elements(page_payload: dict[str, Any]) -> list[dict[str, Any]]:
    parsing_res_list = page_payload.get("parsing_res_list") or []
    table_blocks = [
        block for block in parsing_res_list if "table" in str(block.get("block_label") or "").lower()
    ]
    elements = _collect_table_elements(page_payload, table_blocks)

    for block in parsing_res_list:
        label = str(block.get("block_label") or "").strip()
        bbox = _normalize_bbox(block.get("block_bbox"))
        if not bbox:
            continue
        if "table" in label.lower():
            continue
        text = _safe_text(block.get("block_content"))
        if not text:
            text = _label_placeholder(label)
        elements.append(
            {
                "kind": "layout_block",
                "label": label or "text",
                "bbox": bbox,
                "text": text,
                "score": 1.0,
                "block_id": block.get("block_id"),
                "block_order": block.get("block_order"),
            }
        )
    return elements


def _build_edges(elements: list[dict[str, Any]], page_payload: dict[str, Any]) -> tuple[list[float], list[float]]:
    page_w = float(page_payload.get("width") or 0.0)
    page_h = float(page_payload.get("height") or 0.0)
    x_values: list[float] = [0.0]
    y_values: list[float] = [0.0]
    if page_w > 0:
        x_values.append(page_w)
    if page_h > 0:
        y_values.append(page_h)

    table_elements = [el for el in elements if el["kind"] in {"table_cell", "table_text_fallback"}]
    layout_elements = [el for el in elements if el["kind"] == "layout_block"]

    if table_elements:
        table_widths = [max(1.0, el["bbox"][2] - el["bbox"][0]) for el in table_elements]
        table_heights = [max(1.0, el["bbox"][3] - el["bbox"][1]) for el in table_elements]
        table_median_w = _median(table_widths, default=60.0)
        table_median_h = _median(table_heights, default=24.0)
        table_x_values, table_y_values = _collect_edge_values(table_elements)
        x_values.extend(_cluster_positions(table_x_values, tol=max(1.5, min(6.0, table_median_w * 0.04))))
        y_values.extend(_cluster_positions(table_y_values, tol=max(1.5, min(6.0, table_median_h * 0.08))))

    if layout_elements:
        layout_widths = [max(1.0, el["bbox"][2] - el["bbox"][0]) for el in layout_elements]
        layout_heights = [max(1.0, el["bbox"][3] - el["bbox"][1]) for el in layout_elements]
        layout_median_w = _median(layout_widths, default=160.0)
        layout_median_h = _median(layout_heights, default=36.0)
        layout_x_values, layout_y_values = _collect_edge_values(layout_elements)
        x_values.extend(_cluster_positions(layout_x_values, tol=max(4.0, min(24.0, layout_median_w * 0.1))))
        y_values.extend(_cluster_positions(layout_y_values, tol=max(4.0, min(24.0, layout_median_h * 0.3))))

    # Final merge keeps table boundaries much tighter than layout boundaries.
    x_edges = _cluster_positions(x_values, tol=2.0 if table_elements else 6.0)
    y_edges = _cluster_positions(y_values, tol=2.5 if table_elements else 6.0)
    if len(x_edges) < 2:
        x_edges = [0.0, max(page_w, 100.0)]
    if len(y_edges) < 2:
        y_edges = [0.0, max(page_h, 100.0)]
    return x_edges, y_edges


def _nearest_edge_index(edges: list[float], value: float) -> int:
    if not edges:
        return 0
    idx = bisect_left(edges, value)
    if idx <= 0:
        return 0
    if idx >= len(edges):
        return len(edges) - 1
    before = edges[idx - 1]
    after = edges[idx]
    return idx - 1 if abs(before - value) <= abs(after - value) else idx


def _resolve_range(edges_x: list[float], edges_y: list[float], bbox: list[float]) -> tuple[int, int, int, int]:
    x0, y0, x1, y1 = bbox
    c0 = _nearest_edge_index(edges_x, x0) + 1
    c1 = _nearest_edge_index(edges_x, x1) + 1
    r0 = _nearest_edge_index(edges_y, y0) + 1
    r1 = _nearest_edge_index(edges_y, y1) + 1
    if c1 < c0:
        c0, c1 = c1, c0
    if r1 < r0:
        r0, r1 = r1, r0
    return r0, c0, max(r0, r1), max(c0, c1)


def _range_cells(r0: int, c0: int, r1: int, c1: int) -> list[tuple[int, int]]:
    cells: list[tuple[int, int]] = []
    for row in range(r0, r1 + 1):
        for col in range(c0, c1 + 1):
            cells.append((row, col))
    return cells


def _overlaps(occupied: set[tuple[int, int]], r0: int, c0: int, r1: int, c1: int) -> bool:
    return any(cell in occupied for cell in _range_cells(r0, c0, r1, c1))


def _find_free_anchor(occupied: set[tuple[int, int]], row: int, col: int, max_col: int) -> tuple[int, int]:
    cur_row = row
    cur_col = col
    while (cur_row, cur_col) in occupied:
        cur_col += 1
        if cur_col > max_col:
            cur_row += 1
            cur_col = col
    return cur_row, cur_col


def _apply_dimensions(ws: openpyxl.worksheet.worksheet.Worksheet, x_edges: list[float], y_edges: list[float]) -> None:
    for idx in range(len(x_edges) - 1):
        span_px = max(8.0, x_edges[idx + 1] - x_edges[idx])
        width_chars = min(45.0, max(2.0, span_px / 7.0))
        ws.column_dimensions[get_column_letter(idx + 1)].width = round(width_chars, 2)
    for idx in range(len(y_edges) - 1):
        span_px = max(12.0, y_edges[idx + 1] - y_edges[idx])
        height_pt = min(120.0, max(14.0, span_px * 0.75))
        ws.row_dimensions[idx + 1].height = round(height_pt, 2)


def _write_layout_sheet(
    wb: openpyxl.Workbook,
    page_payload: dict[str, Any],
    elements: list[dict[str, Any]],
    page_index: int,
) -> None:
    ws = wb.create_sheet(title=f"page_{page_index + 1:03d}")
    x_edges, y_edges = _build_edges(elements, page_payload)
    _apply_dimensions(ws, x_edges, y_edges)

    sort_order = {"table_cell": 0, "table_text_fallback": 1, "layout_block": 2}
    occupied: set[tuple[int, int]] = set()
    max_col = max(1, len(x_edges))

    for element in sorted(
        elements,
        key=lambda el: (
            sort_order.get(el["kind"], 9),
            el["bbox"][1],
            el["bbox"][0],
            el["bbox"][3] - el["bbox"][1],
        ),
    ):
        r0, c0, r1, c1 = _resolve_range(x_edges, y_edges, element["bbox"])
        if _overlaps(occupied, r0, c0, r1, c1):
            r0, c0 = _find_free_anchor(occupied, r0, c0, max_col=max_col)
            r1, c1 = r0, c0

        text = _safe_text(element["text"])
        if not text:
            continue
        cell = ws.cell(row=r0, column=c0, value=text)
        style = _style_descriptor(element["label"], text, kind=element["kind"])
        cell.font = style["font"]
        cell.alignment = style["alignment"]
        cell.border = style["border"]
        if style["fill"] is not None:
            cell.fill = style["fill"]

        for row, col in _range_cells(r0, c0, r1, c1):
            occupied.add((row, col))
            ws.cell(row=row, column=col).border = style["border"]
            if style["fill"] is not None:
                ws.cell(row=row, column=col).fill = style["fill"]
        if r1 > r0 or c1 > c0:
            ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)

    ws.freeze_panes = "A1"


def _write_blocks_sheet(wb: openpyxl.Workbook, page_payload: dict[str, Any], page_index: int) -> None:
    ws = wb.create_sheet(title=f"blocks_{page_index + 1:03d}")
    ws.append(["order", "label", "text", "bbox"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.border = BLOCK_BORDER
    for block in page_payload.get("parsing_res_list", []) or []:
        ws.append(
            [
                block.get("block_order"),
                block.get("block_label"),
                _safe_text(block.get("block_content")),
                json.dumps(block.get("block_bbox"), ensure_ascii=False),
            ]
        )
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 28
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BLOCK_BORDER


def _write_meta_sheet(wb: openpyxl.Workbook, summary: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "meta"
    rows = [
        ("input_path", summary.get("input_path")),
        ("page_count", summary.get("page_count")),
        ("doc_dir", summary.get("doc_dir")),
        ("combined_markdown_path", summary.get("combined_markdown_path")),
    ]
    for key, value in rows:
        ws.append([key, value])
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80


def _write_flow_sheet(
    wb: openpyxl.Workbook,
    page_payload: dict[str, Any],
    page_index: int,
    column_width_mode: str,
) -> None:
    ws = wb.create_sheet(title=f"page_{page_index + 1:03d}")
    blocks, max_table_cols = _collect_page_flow_blocks(page_payload, column_width_mode=column_width_mode)
    total_cols = max(4, max_table_cols)

    table_widths: list[float] = []
    for block in blocks:
        if block["kind"] == "table_block":
            widths = block.get("col_widths") or []
            if len(widths) > len(table_widths):
                table_widths = list(widths)
    if len(table_widths) < total_cols:
        table_widths.extend([16.0] * (total_cols - len(table_widths)))
    _set_sheet_base_widths(ws, total_cols=total_cols, widths=table_widths)

    current_row = 1
    for block in blocks:
        if block["kind"] == "text_block":
            current_row = _apply_text_block(ws, start_row=current_row, total_cols=total_cols, block=block)
        elif block["kind"] == "table_block":
            current_row = _apply_table_block(ws, start_row=current_row, total_cols=total_cols, block=block)
    ws.freeze_panes = "A1"


def convert_summary(summary_path: Path, out_dir: Path, column_width_mode: str) -> Path:
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    summary_dir = summary_path.parent
    page_jsons: list[Path] = []
    for page_info in summary.get("pages", []) or []:
        json_ref = page_info.get("json_path")
        if json_ref:
            resolved = _resolve_ref(str(json_ref), base_dir=summary_dir)
            if resolved:
                page_jsons.append(resolved)
    if not page_jsons:
        page_jsons = sorted((summary_dir / "json").glob("*.json"))
    if not page_jsons:
        raise FileNotFoundError(f"No page json files found for summary: {summary_path}")

    wb = openpyxl.Workbook()
    _write_meta_sheet(wb, summary)
    for idx, page_json in enumerate(page_jsons):
        page_payload = _load_page_payload(page_json)
        _write_flow_sheet(
            wb,
            page_payload=page_payload,
            page_index=idx,
            column_width_mode=column_width_mode,
        )
        _write_blocks_sheet(wb, page_payload=page_payload, page_index=idx)

    input_name = Path(str(summary.get("input_path") or "")).stem
    doc_name = _safe_filename_stem(input_name) if input_name else _doc_id_from_summary(summary_path)
    out_path = out_dir / f"{doc_name}_full_layout.xlsx"
    wb.save(out_path)
    return out_path


def build_arg_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(
        description="Reconstruct a full-page Excel layout from PP-StructureV3 outputs, including non-table text blocks."
    )
    ap.add_argument(
        "--input_path",
        default="out/pp_parse",
        help="A summary.json file or a directory containing PP-StructureV3 output folders.",
    )
    ap.add_argument(
        "--out_dir",
        default="out/pp_rebuild",
        help="Directory for reconstructed xlsx files.",
    )
    ap.add_argument("--recursive", action="store_true", help="Recursively discover summary.json files.")
    ap.add_argument(
        "--column_width_mode",
        choices=["geometry", "content"],
        default="geometry",
        help="`geometry`: preserve PDF-like widths; `content`: widen columns using text length on top of geometry.",
    )
    return ap


def main() -> int:
    parser = build_arg_parser()
    args = parser.parse_args()

    input_path = Path(args.input_path)
    out_dir = _ensure_dir(Path(args.out_dir))
    summaries = discover_summaries(input_path=input_path, recursive=bool(args.recursive))
    if not summaries:
        raise FileNotFoundError(f"No summary.json found under: {input_path}")

    ok = 0
    for idx, summary_path in enumerate(summaries, start=1):
        out_path = convert_summary(
            summary_path=summary_path,
            out_dir=out_dir,
            column_width_mode=args.column_width_mode,
        )
        ok += 1
        print(f"[OK] {idx}/{len(summaries)} {summary_path} -> {out_path}")
    print(f"[DONE] total={len(summaries)} ok={ok} out_dir={out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
