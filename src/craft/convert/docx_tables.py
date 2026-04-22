#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
    Alignment = Border = Font = PatternFill = Side = None
    get_column_letter = None


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


@dataclass
class CellStyle:
    fill_color: str | None = None
    border: dict[str, dict[str, str]] = field(default_factory=dict)
    horizontal: str | None = None
    vertical: str | None = None
    bold: bool = False
    italic: bool = False
    underline: bool = False
    font_color: str | None = None
    font_size: float | None = None


@dataclass
class CellData:
    text: str
    row: int
    col: int
    row_span: int = 1
    col_span: int = 1
    width_twips: int | None = None
    style: CellStyle = field(default_factory=CellStyle)


@dataclass
class TableData:
    n_rows: int
    n_cols: int
    cells: list[CellData]
    grid_widths: list[int | None]


def _qname(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


def _normalize_text(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _iter_text_parts(cell: ET.Element) -> Iterable[str]:
    for child in cell.iter():
        if child.tag == _qname("t"):
            yield child.text or ""
        elif child.tag == _qname("tab"):
            yield "\t"
        elif child.tag in {_qname("br"), _qname("cr")}:
            yield "\n"


def _cell_text(cell: ET.Element) -> str:
    return _normalize_text("".join(_iter_text_parts(cell)))


def _get_w_val(element: ET.Element | None) -> str | None:
    if element is None:
        return None
    return element.get(_qname("val")) or element.get("val")


def _parse_int(value: str | None, default: int = 0) -> int:
    try:
        return int(value or "")
    except (TypeError, ValueError):
        return default


def _grid_span(cell: ET.Element) -> int:
    tc_pr = cell.find("w:tcPr", NS)
    if tc_pr is None:
        return 1
    return max(1, _parse_int(_get_w_val(tc_pr.find("w:gridSpan", NS)), default=1))


def _v_merge_state(cell: ET.Element) -> str | None:
    tc_pr = cell.find("w:tcPr", NS)
    if tc_pr is None:
        return None
    value = _get_w_val(tc_pr.find("w:vMerge", NS))
    if value is None:
        if tc_pr.find("w:vMerge", NS) is not None:
            return "continue"
        return None
    return value or "continue"


def _cell_width_twips(cell: ET.Element) -> int | None:
    tc_pr = cell.find("w:tcPr", NS)
    if tc_pr is None:
        return None
    tc_w = tc_pr.find("w:tcW", NS)
    if tc_w is None:
        return None
    width_type = tc_w.get(_qname("type")) or tc_w.get("type")
    width_raw = tc_w.get(_qname("w")) or tc_w.get("w")
    if width_type != "dxa":
        return None
    width = _parse_int(width_raw, default=0)
    return width or None


def _normalize_color(value: str | None) -> str | None:
    if not value:
        return None
    value = value.strip().lstrip("#").upper()
    if value in {"AUTO", "NONE"}:
        return None
    if len(value) == 6 and all(ch in "0123456789ABCDEF" for ch in value):
        return value
    return None


def _parse_shading(tc_pr: ET.Element | None) -> str | None:
    if tc_pr is None:
        return None
    shd = tc_pr.find("w:shd", NS)
    if shd is None:
        return None
    fill = shd.get(_qname("fill")) or shd.get("fill")
    return _normalize_color(fill)


def _parse_borders(tc_pr: ET.Element | None) -> dict[str, dict[str, str]]:
    if tc_pr is None:
        return {}
    tc_borders = tc_pr.find("w:tcBorders", NS)
    if tc_borders is None:
        return {}

    borders: dict[str, dict[str, str]] = {}
    side_map = {
        "top": "top",
        "bottom": "bottom",
        "left": "left",
        "right": "right",
    }
    for word_side, excel_side in side_map.items():
        border_el = tc_borders.find(f"w:{word_side}", NS)
        if border_el is None:
            continue
        borders[excel_side] = {
            "val": _get_w_val(border_el) or "",
            "sz": border_el.get(_qname("sz")) or border_el.get("sz") or "",
            "color": border_el.get(_qname("color")) or border_el.get("color") or "",
        }
    return borders


def _parse_alignment(cell: ET.Element) -> tuple[str | None, str | None]:
    paragraphs = cell.findall("./w:p", NS)
    horizontal: str | None = None
    vertical: str | None = None

    tc_pr = cell.find("w:tcPr", NS)
    if tc_pr is not None:
        v_align = tc_pr.find("w:vAlign", NS)
        vertical = _get_w_val(v_align)

    for paragraph in paragraphs:
        p_pr = paragraph.find("w:pPr", NS)
        if p_pr is None:
            continue
        jc = p_pr.find("w:jc", NS)
        value = _get_w_val(jc)
        if value:
            horizontal = value
            break
    return horizontal, vertical


def _first_run_props(cell: ET.Element) -> dict[str, str | bool | float | None]:
    for run in cell.findall(".//w:r", NS):
        r_pr = run.find("w:rPr", NS)
        if r_pr is None:
            continue
        color = _normalize_color(_get_w_val(r_pr.find("w:color", NS)))
        size_raw = _get_w_val(r_pr.find("w:sz", NS))
        size = None
        if size_raw:
            try:
                size = float(size_raw) / 2.0
            except ValueError:
                size = None
        return {
            "bold": r_pr.find("w:b", NS) is not None,
            "italic": r_pr.find("w:i", NS) is not None,
            "underline": r_pr.find("w:u", NS) is not None,
            "font_color": color,
            "font_size": size,
        }
    return {
        "bold": False,
        "italic": False,
        "underline": False,
        "font_color": None,
        "font_size": None,
    }


def _parse_cell_style(cell: ET.Element) -> CellStyle:
    tc_pr = cell.find("w:tcPr", NS)
    horizontal, vertical = _parse_alignment(cell)
    run_props = _first_run_props(cell)
    return CellStyle(
        fill_color=_parse_shading(tc_pr),
        border=_parse_borders(tc_pr),
        horizontal=horizontal,
        vertical=vertical,
        bold=bool(run_props["bold"]),
        italic=bool(run_props["italic"]),
        underline=bool(run_props["underline"]),
        font_color=run_props["font_color"] if isinstance(run_props["font_color"], str) else None,
        font_size=run_props["font_size"] if isinstance(run_props["font_size"], float) else None,
    )


def _parse_tbl_grid(tbl: ET.Element) -> list[int | None]:
    tbl_grid = tbl.find("./w:tblGrid", NS)
    if tbl_grid is None:
        return []
    widths: list[int | None] = []
    for grid_col in tbl_grid.findall("./w:gridCol", NS):
        width_raw = grid_col.get(_qname("w")) or grid_col.get("w")
        width = _parse_int(width_raw, default=0)
        widths.append(width or None)
    return widths


def _map_word_border_style(value: str) -> str | None:
    mapping = {
        "single": "thin",
        "thick": "thick",
        "double": "double",
        "dashed": "dashed",
        "dotted": "dotted",
        "dashSmallGap": "dashed",
        "dotDash": "dashDot",
        "dotDotDash": "dashDotDot",
    }
    if value in {"nil", "none"}:
        return None
    return mapping.get(value, "thin")


def _build_side(spec: dict[str, str]) -> Side:
    style = _map_word_border_style(spec.get("val", ""))
    if style is None:
        return Side(style=None)
    color = _normalize_color(spec.get("color"))
    return Side(style=style, color=color or "000000")


def _apply_cell_style(ws: openpyxl.worksheet.worksheet.Worksheet, cell: CellData) -> None:
    target = ws.cell(row=cell.row, column=cell.col)
    target.value = cell.text

    style = cell.style
    target.alignment = Alignment(
        horizontal={
            "left": "left",
            "center": "center",
            "right": "right",
            "both": "justify",
            "distribute": "distributed",
        }.get(style.horizontal, None),
        vertical={
            "top": "top",
            "center": "center",
            "bottom": "bottom",
        }.get(style.vertical, None),
        wrap_text=True,
    )

    if style.fill_color:
        target.fill = PatternFill(fill_type="solid", fgColor=style.fill_color)

    if style.border:
        target.border = Border(
            left=_build_side(style.border.get("left", {})),
            right=_build_side(style.border.get("right", {})),
            top=_build_side(style.border.get("top", {})),
            bottom=_build_side(style.border.get("bottom", {})),
        )

    if (
        style.bold
        or style.italic
        or style.underline
        or style.font_color
        or style.font_size is not None
    ):
        target.font = Font(
            bold=style.bold,
            italic=style.italic,
            underline="single" if style.underline else None,
            color=style.font_color,
            size=style.font_size,
        )


def _compute_vertical_spans(pending: dict[int, CellData], table_row_count: int) -> None:
    for cell in pending.values():
        cell.row_span = max(1, table_row_count - cell.row + 1)


def read_docx_tables(path: Path) -> list[TableData]:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))

    tables: list[TableData] = []
    for tbl in root.findall(".//w:tbl", NS):
        grid_widths = _parse_tbl_grid(tbl)
        cells: list[CellData] = []
        pending_vertical: dict[int, CellData] = {}
        n_cols = 0
        row_idx = 0

        for tr in tbl.findall("./w:tr", NS):
            row_idx += 1
            col_idx = 1
            for tc in tr.findall("./w:tc", NS):
                while col_idx in pending_vertical and pending_vertical[col_idx].row + pending_vertical[col_idx].row_span > row_idx:
                    col_idx += 1

                text = _cell_text(tc)
                col_span = _grid_span(tc)
                style = _parse_cell_style(tc)
                width = _cell_width_twips(tc)
                v_merge = _v_merge_state(tc)

                if v_merge == "continue":
                    for offset in range(col_span):
                        parent = pending_vertical.get(col_idx + offset)
                        if parent is not None:
                            parent.row_span += 1
                    col_idx += col_span
                    n_cols = max(n_cols, col_idx - 1)
                    continue

                cell = CellData(
                    text=text,
                    row=row_idx,
                    col=col_idx,
                    row_span=1,
                    col_span=col_span,
                    width_twips=width,
                    style=style,
                )
                cells.append(cell)

                if v_merge == "restart":
                    for offset in range(col_span):
                        pending_vertical[col_idx + offset] = cell
                else:
                    for offset in range(col_span):
                        pending_vertical.pop(col_idx + offset, None)

                col_idx += col_span
                n_cols = max(n_cols, col_idx - 1)

            finished_cols = [
                col
                for col, parent in pending_vertical.items()
                if parent.row + parent.row_span <= row_idx
            ]
            for col in finished_cols:
                pending_vertical.pop(col, None)

        _compute_vertical_spans(pending_vertical, row_idx)
        tables.append(TableData(n_rows=row_idx, n_cols=n_cols, cells=cells, grid_widths=grid_widths))
    return tables


def _safe_sheet_title(index: int) -> str:
    title = f"table_{index}"
    title = INVALID_SHEET_CHARS.sub("_", title)
    return title[:31]


def _set_column_widths(ws: openpyxl.worksheet.worksheet.Worksheet, table: TableData) -> None:
    widths = list(table.grid_widths)
    if len(widths) < table.n_cols:
        widths.extend([None] * (table.n_cols - len(widths)))

    if any(widths):
        for idx, width in enumerate(widths[: table.n_cols], start=1):
            if width:
                excel_width = max(2.0, min(60.0, width / 256.0))
                ws.column_dimensions[get_column_letter(idx)].width = excel_width
        return

    inferred: dict[int, int] = {}
    for cell in table.cells:
        if cell.width_twips and cell.col_span == 1:
            inferred[cell.col] = max(inferred.get(cell.col, 0), cell.width_twips)
    for idx in range(1, table.n_cols + 1):
        width = inferred.get(idx)
        if width:
            ws.column_dimensions[get_column_letter(idx)].width = max(2.0, min(60.0, width / 256.0))


def write_xlsx(tables: list[TableData], output_path: Path) -> None:
    if openpyxl is None:
        raise SystemExit("openpyxl is required for .xlsx output. Please install it first.")

    wb = openpyxl.Workbook()
    first_sheet = True

    for idx, table in enumerate(tables, start=1):
        ws = wb.active if first_sheet else wb.create_sheet()
        first_sheet = False
        ws.title = _safe_sheet_title(idx)

        for cell in table.cells:
            _apply_cell_style(ws, cell)
            if cell.row_span > 1 or cell.col_span > 1:
                ws.merge_cells(
                    start_row=cell.row,
                    start_column=cell.col,
                    end_row=cell.row + cell.row_span - 1,
                    end_column=cell.col + cell.col_span - 1,
                )

        _set_column_widths(ws, table)

    if not tables:
        ws = wb.active
        ws.title = "table_1"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def write_csvs(tables: list[TableData], output_dir: Path, stem: str) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    if not tables:
        out_path = output_dir / f"{stem}_table_1.csv"
        with out_path.open("w", newline="", encoding="utf-8-sig") as handle:
            csv.writer(handle).writerow([])
        return [out_path]

    for idx, table in enumerate(tables, start=1):
        matrix = [["" for _ in range(table.n_cols)] for _ in range(table.n_rows)]
        for cell in table.cells:
            matrix[cell.row - 1][cell.col - 1] = cell.text
        out_path = output_dir / f"{stem}_table_{idx}.csv"
        with out_path.open("w", newline="", encoding="utf-8-sig") as handle:
            csv.writer(handle).writerows(matrix)
        written.append(out_path)
    return written


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Extract tables from a .docx file and export them as .xlsx and/or .csv."
    )
    parser.add_argument("input_docx", type=Path, help="Path to the input .docx file.")
    parser.add_argument(
        "--output",
        type=Path,
        help="Output .xlsx path. Defaults to the input file path with .xlsx extension.",
    )
    parser.add_argument(
        "--format",
        choices=("xlsx", "csv", "both"),
        default="xlsx",
        help="Export format. Default: xlsx.",
    )
    parser.add_argument(
        "--csv-dir",
        type=Path,
        help="Directory for CSV files. Defaults to a sibling directory named after the input file stem.",
    )
    return parser


def main() -> None:
    args = build_parser().parse_args()
    input_docx: Path = args.input_docx

    if input_docx.suffix.lower() != ".docx":
        raise SystemExit(f"Expected a .docx file, got: {input_docx}")
    if not input_docx.is_file():
        raise SystemExit(f"Input file not found: {input_docx}")

    tables = read_docx_tables(input_docx)
    print(f"Found {len(tables)} table(s) in {input_docx.name}")

    if args.format in {"xlsx", "both"}:
        output_path = args.output or input_docx.with_suffix(".xlsx")
        write_xlsx(tables, output_path)
        print(f"Wrote Excel output to: {output_path}")

    if args.format in {"csv", "both"}:
        csv_dir = args.csv_dir or input_docx.with_name(input_docx.stem + "_csv")
        written = write_csvs(tables, csv_dir, input_docx.stem)
        print(f"Wrote {len(written)} CSV file(s) to: {csv_dir}")


if __name__ == "__main__":
    main()
