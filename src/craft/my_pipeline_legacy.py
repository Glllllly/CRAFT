import argparse
from difflib import SequenceMatcher
import json
import os
import re
import shutil
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openai import OpenAI
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries

ROOT = Path(__file__).resolve().parent
RUNTIME_DIR = ROOT / "runtime"

try:
    from .runtime.engine import (
        PipelineError,
        build_form_from_qwen,
        ensure_xlsx,
        run_qwen_lora,
        run_qwen_small,
    )
    from .reflect_plugin import ReflectContext, ReflectTarget, normalize_key, run_reflect_round
except ImportError:
    if str(RUNTIME_DIR) not in sys.path:
        sys.path.insert(0, str(RUNTIME_DIR))
    from engine import (  # type: ignore
        PipelineError,
        build_form_from_qwen,
        ensure_xlsx,
        run_qwen_lora,
        run_qwen_small,
    )
    from reflect_plugin import ReflectContext, ReflectTarget, normalize_key, run_reflect_round  # type: ignore


def build_client(base_url: str | None, api_key: str | None) -> OpenAI:
    key = api_key or os.getenv("OPENAI_API_KEY") or os.getenv("PPCHAT_API_KEY")
    if not key:
        raise RuntimeError("Missing API key. Set OPENAI_API_KEY or PPCHAT_API_KEY.")
    return OpenAI(base_url=base_url or os.getenv("OPENAI_BASE_URL"), api_key=key)


def _enable_windows_ansi() -> None:
    if os.name != "nt":
        return
    try:
        import ctypes

        kernel32 = ctypes.windll.kernel32
        handle = kernel32.GetStdHandle(-11)
        if handle in (0, -1):
            return
        mode = ctypes.c_uint32()
        if kernel32.GetConsoleMode(handle, ctypes.byref(mode)):
            kernel32.SetConsoleMode(handle, mode.value | 0x0004)
    except Exception:
        pass


def _supports_console_color() -> bool:
    if os.getenv("NO_COLOR"):
        return False
    force = str(os.getenv("FORCE_COLOR") or "").strip()
    if force and force != "0":
        return True
    stream = getattr(sys, "stdout", None)
    if stream is None or not hasattr(stream, "isatty") or not stream.isatty():
        return False
    if str(os.getenv("TERM") or "").lower() == "dumb":
        return False
    _enable_windows_ansi()
    return True


_ANSI_RESET = "\033[0m"
_ANSI_BY_KIND = {
    "info": "\033[1;36m",
    "warn": "\033[1;33m",
    "error": "\033[1;31m",
    "ok": "\033[1;32m",
    "run": "\033[1;34m",
    "usage": "\033[1;96m",
    "reflect": "\033[1;35m",
    "reflect_check": "\033[1;95m",
    "assess_ok": "\033[32m",
    "assess_warn": "\033[1;33m",
    "assess_bad": "\033[1;31m",
    "html": "\033[2;37m",
}
_CONSOLE_COLOR_ENABLED = _supports_console_color()
_LAST_CONSOLE_SUMMARY: str | None = None


def _paint_console(text: str, kind: str) -> str:
    if not _CONSOLE_COLOR_ENABLED:
        return text
    prefix = _ANSI_BY_KIND.get(kind, "")
    if not prefix:
        return text
    return f"{prefix}{text}{_ANSI_RESET}"


def _colorize_console_line(msg: str) -> str:
    if not _CONSOLE_COLOR_ENABLED:
        return msg
    if "[REFLECT_PLUGIN][ASSESS]" in msg:
        if "high_risk=True" in msg:
            return _paint_console(msg, "assess_bad")
        if "flags=[]" in msg:
            return _paint_console(msg, "assess_ok")
        return _paint_console(msg, "assess_warn")
    if "[WARN]" in msg:
        return _paint_console(msg, "warn")
    if "FAIL" in msg or " failed" in msg or "failed:" in msg:
        return _paint_console(msg, "error")
    if "[HTML_RENDER]" in msg:
        return _paint_console(msg, "html")
    if "[REFLECT_PLUGIN]" in msg:
        return _paint_console(msg, "reflect")
    if "[REFLECT]" in msg:
        return _paint_console(msg, "reflect_check")
    if "[AGENT]" in msg or "[INFO]" in msg:
        return _paint_console(msg, "info")
    if msg.startswith("Done. Filled file:"):
        return _paint_console(msg, "ok")
    if "[OK]" in msg:
        return _paint_console(msg, "ok")
    if "] RUN " in msg or msg.startswith("[RUN]"):
        return _paint_console(msg, "run")
    if "[USAGE]" in msg or "] USAGE " in msg:
        return _paint_console(msg, "usage")
    return msg


def _safe_console_text(text: str) -> str:
    try:
        encoding = getattr(sys.stdout, "encoding", None) or "utf-8"
        return str(text).encode(encoding, errors="replace").decode(encoding, errors="replace")
    except Exception:
        return str(text)


def _truncate_console(text: str, limit: int = 120) -> str:
    clean = re.sub(r"\s+", " ", str(text or "")).strip()
    if len(clean) <= limit:
        return clean
    return clean[: max(0, limit - 3)] + "..."


def _strip_log_tags(text: str) -> str:
    return re.sub(r"\[[^\]]+\]", "", str(text or "")).strip(" :-")


def _summarize_console_line(msg: str) -> str | None:
    text = str(msg or "").strip()
    if not text:
        return None
    if text.startswith("Done. Filled file:"):
        return text
    if "[USAGE]" in text or "[SOURCE_FILES]" in text:
        return None
    if "[INPUT_PLANNER]" in text:
        return "Step: prepare input"
    if "[SOURCE_ADAPTER]" in text:
        return "Step: read source files"
    if "[HTML_RENDER]" in text:
        return "Step: render worksheet"
    if "Exporting sheet image via Excel COM" in text:
        return None
    if "[AGENT][PLAN]" in text:
        action_match = re.search(r"action=([a-z_]+)", text)
        region_match = re.search(r"region=([A-Z0-9:.-]+)", text)
        action = str(action_match.group(1)) if action_match else ""
        region = str(region_match.group(1)) if region_match else ""
        mapping = {
            "free_fill": "Step: first pass fill",
            "reflect_and_hint": "Step: inspect result",
            "restore_template_region": "Step: restore template cells",
            "guided_refill": "Step: repair workbook",
            "verify_and_deliver": "Step: verify result",
            "deliver": "Step: finalize output",
        }
        summary = mapping.get(action, "Step: process workbook")
        if region and region != "-":
            if action in {"guided_refill", "restore_template_region"}:
                summary += f" ({region})"
        return summary
    if "[GUIDED_REFILL]" in text:
        region_match = re.search(r"region=([A-Z0-9:.-]+)", text)
        region = str(region_match.group(1)) if region_match else ""
        return f"Step: repair workbook{f' ({region})' if region else ''}"
    if "[REFLECT_PLUGIN]" in text or "[REFLECT]" in text or "[VERIFY]" in text:
        if "no risk" in text.lower() or "no local fix needed" in text.lower():
            return "Step: verification passed"
        return "Step: inspect result"
    if "Running slot_detect.py" in text or "Running qwen_small.py" in text:
        return None
    if "Running lora_infer.py" in text or "infer_lora_qwen_vl.py" in text:
        return "Step: match fields"
    if "Converting .xls to .xlsx" in text:
        return "Step: convert spreadsheet"
    if "[WARN]" in text or "failed" in text.lower():
        cleaned = _strip_log_tags(text)
        return f"Warning: {_truncate_console(cleaned, 100)}"
    if "[INFO]" in text:
        cleaned = _strip_log_tags(text)
        if "skip_initial_fill" in cleaned:
            return "Step: reuse existing workbook"
        return None
    return None


def _emit_console_line(msg: str) -> None:
    global _LAST_CONSOLE_SUMMARY
    summary = _summarize_console_line(msg)
    if summary is None:
        return
    if summary == _LAST_CONSOLE_SUMMARY:
        return
    _LAST_CONSOLE_SUMMARY = summary
    print(_safe_console_text(_colorize_console_line(summary)), flush=True)


DEFAULT_DPI = 96


def _col_width_to_px(width: float | None) -> int:
    if width is None:
        width = 8.43
    return int(round(width * 7 + 5))


def _row_height_to_px(height: float | None, dpi: int = DEFAULT_DPI) -> int:
    if height is None:
        height = 15.0
    return int(round(height * dpi / 72.0))


def _get_used_bbox(ws) -> tuple[int, int, int, int]:
    def _cell_has_border(cell) -> bool:
        b = cell.border
        if b is None:
            return False
        return any(getattr(getattr(b, side, None), "style", None) for side in ("top", "bottom", "left", "right"))

    def _cell_has_fill(cell) -> bool:
        f = cell.fill
        if f is None:
            return False
        return bool(getattr(f, "fill_type", None))

    max_r = max(1, int(getattr(ws, "max_row", 1) or 1))
    max_c = max(1, int(getattr(ws, "max_column", 1) or 1))

    merged_cells: set[tuple[int, int]] = set()
    for rng in ws.merged_cells.ranges:
        for rr in range(rng.min_row, rng.max_row + 1):
            for cc in range(rng.min_col, rng.max_col + 1):
                merged_cells.add((rr, cc))

    min_row = min_col = None
    max_row = max_col = None

    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            cell = ws.cell(row=r, column=c)
            has_content = cell.value is not None and str(cell.value).strip() != ""
            structural = _cell_has_border(cell) or _cell_has_fill(cell) or ((r, c) in merged_cells)
            if has_content or structural:
                if min_row is None:
                    min_row, min_col, max_row, max_col = r, c, r, c
                else:
                    min_row = min(min_row, r)
                    min_col = min(min_col, c)
                    max_row = max(max_row, r)
                    max_col = max(max_col, c)

    if min_row is None:
        return 1, 1, 1, 1
    return int(min_row), int(min_col), int(max_row), int(max_col)


def _compute_edges(ws, min_row: int, min_col: int, max_row: int, max_col: int, dpi: int) -> tuple[list[int], list[int]]:
    x_edges = [0]
    for c in range(min_col, max_col + 1):
        letter = get_column_letter(c)
        w = ws.column_dimensions[letter].width
        x_edges.append(x_edges[-1] + _col_width_to_px(w))
    y_edges = [0]
    for r in range(min_row, max_row + 1):
        h = ws.row_dimensions[r].height
        y_edges.append(y_edges[-1] + _row_height_to_px(h, dpi=dpi))
    return x_edges, y_edges


def _rgb_from_color_obj(color_obj) -> str | None:
    if color_obj is None:
        return None
    rgb = getattr(color_obj, "rgb", None)
    if isinstance(rgb, str):
        rgb = rgb.strip()
        if len(rgb) == 8:
            rgb = rgb[2:]
        if len(rgb) == 6:
            return f"#{rgb.upper()}"
    return None


def _css_border_from_side(side) -> str:
    st = getattr(side, "style", None) if side is not None else None
    if not st:
        return "none"
    width_map = {
        "hair": "1px",
        "thin": "1px",
        "medium": "2px",
        "thick": "3px",
        "double": "3px",
        "dashed": "1px",
        "dotted": "1px",
    }
    style_map = {
        "double": "double",
        "dashed": "dashed",
        "dotted": "dotted",
    }
    w = width_map.get(st, "1px")
    s = style_map.get(st, "solid")
    c = _rgb_from_color_obj(getattr(side, "color", None)) or "#444444"
    return f"{w} {s} {c}"


def _merged_anchor_map(ws) -> dict[tuple[int, int], tuple[int, int, int, int]]:
    m: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for rng in ws.merged_cells.ranges:
        r0, c0, r1, c1 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        for rr in range(r0, r1 + 1):
            for cc in range(c0, c1 + 1):
                m[(rr, cc)] = (r0, c0, r1, c1)
    return m


def _html_escape(s: str) -> str:
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )


def _value_to_text(v: object) -> str:
    if v is None:
        return ""
    return str(v)


def _encode_image_base64(image_path: Path) -> tuple[str, str]:
    suffix = image_path.suffix.lower()
    mime = {
        ".png": "png",
        ".jpg": "jpeg",
        ".jpeg": "jpeg",
        ".webp": "webp",
        ".bmp": "bmp",
    }.get(suffix, "png")
    import base64
    import io

    try:
        from PIL import Image
    except Exception:
        data = image_path.read_bytes()
        return base64.b64encode(data).decode("utf-8"), mime

    max_dim = 7900
    with Image.open(image_path) as img:
        width, height = img.size
        if max(width, height) > max_dim:
            scale = min(max_dim / float(max(width, height)), 1.0)
            new_size = (
                max(1, int(width * scale)),
                max(1, int(height * scale)),
            )
            resample = getattr(Image, "Resampling", Image).LANCZOS
            img = img.resize(new_size, resample)
            if mime == "jpeg" and str(img.mode or "").upper() in {"RGBA", "LA", "P"}:
                img = img.convert("RGB")
            buf = io.BytesIO()
            save_format = "JPEG" if mime == "jpeg" else "PNG"
            save_mime = "jpeg" if mime == "jpeg" else "png"
            img.save(buf, format=save_format, optimize=True)
            return base64.b64encode(buf.getvalue()).decode("utf-8"), save_mime

    data = image_path.read_bytes()

    return base64.b64encode(data).decode("utf-8"), mime


def _parse_chat_response(resp: Any) -> dict[str, Any]:
    """
    Normalize various gateway response shapes into:
      {"assistant_message": <obj|dict>, "content": <str>, "tool_calls": <list>}
    """
    candidate = resp
    if isinstance(resp, str):
        s = resp.strip()
        if not s:
            return {"assistant_message": {"role": "assistant", "content": ""}, "content": "", "tool_calls": []}
        try:
            candidate = json.loads(s)
        except Exception:
            return {"assistant_message": {"role": "assistant", "content": s}, "content": s, "tool_calls": []}

    # OpenAI python object path
    if hasattr(candidate, "choices"):
        try:
            msg = candidate.choices[0].message
            content = getattr(msg, "content", "") or ""
            tool_calls = getattr(msg, "tool_calls", None) or []
            return {"assistant_message": msg, "content": content, "tool_calls": tool_calls}
        except Exception:
            pass

    # Dict-like path
    if isinstance(candidate, dict):
        choices = candidate.get("choices")
        if isinstance(choices, list) and choices:
            first = choices[0]
            if isinstance(first, dict):
                m = first.get("message")
                if isinstance(m, dict):
                    content = m.get("content", "")
                    if isinstance(content, list):
                        parts: list[str] = []
                        for it in content:
                            if isinstance(it, dict) and isinstance(it.get("text"), str):
                                parts.append(it.get("text", ""))
                        content = "\n".join([x for x in parts if x.strip()]) if parts else json.dumps(content, ensure_ascii=False)
                    elif not isinstance(content, str):
                        content = json.dumps(content, ensure_ascii=False)
                    tool_calls = m.get("tool_calls", [])
                    if not isinstance(tool_calls, list):
                        tool_calls = []
                    return {
                        "assistant_message": {"role": "assistant", "content": content},
                        "content": content,
                        "tool_calls": tool_calls,
                    }
                txt = first.get("text")
                if isinstance(txt, str):
                    return {"assistant_message": {"role": "assistant", "content": txt}, "content": txt, "tool_calls": []}

        # Some providers put output in top-level fields.
        for k in ("content", "text", "output"):
            v = candidate.get(k)
            if isinstance(v, str):
                return {"assistant_message": {"role": "assistant", "content": v}, "content": v, "tool_calls": []}

    raise PipelineError(f"Unexpected chat response type: {type(resp).__name__}")


def _usage_to_dict(resp: Any) -> dict[str, int]:
    usage = getattr(resp, "usage", None)
    if not usage:
        return {}
    out: dict[str, int] = {}
    for k in ("prompt_tokens", "completion_tokens", "total_tokens"):
        v = getattr(usage, k, None)
        if isinstance(v, int):
            out[k] = v
    return out


def _load_edit_records(edit_log_jsonl: Path) -> list[dict[str, Any]]:
    if not edit_log_jsonl.exists():
        return []
    rows: list[dict[str, Any]] = []
    with edit_log_jsonl.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                item = json.loads(line)
            except Exception:
                continue
            if isinstance(item, dict):
                rows.append(item)
    return rows


def _actual_edit_records_for_round(edit_log_jsonl: Path, round_idx: int) -> list[dict[str, Any]]:
    rows = _load_edit_records(edit_log_jsonl)
    return [
        row
        for row in rows
        if bool(row.get("is_actual_edit")) and int(row.get("round_idx", -1) or -1) == int(round_idx)
    ]


def _format_edit_cell_list(rows: list[dict[str, Any]], limit: int = 12) -> str:
    cells: list[str] = []
    for row in rows:
        cell = str(row.get("target_cell") or "").strip()
        if cell:
            cells.append(cell)
    uniq = list(dict.fromkeys(cells))
    if len(uniq) <= limit:
        return json.dumps(uniq, ensure_ascii=False)
    return json.dumps(uniq[:limit] + [f"... +{len(uniq) - limit} more"], ensure_ascii=False)


def _collect_nonempty_cells(xlsx_path: Path, sheet_name: str | None) -> dict[str, str]:
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb[wb.sheetnames[0]]
        out: dict[str, str] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if text:
                    out[str(cell.coordinate).upper()] = text
        return out
    finally:
        wb.close()


def _collect_nonempty_cells_in_range(xlsx_path: Path, sheet_name: str | None, cell_range: str) -> dict[str, str]:
    normalized = _normalize_a1_range(cell_range)
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb[wb.sheetnames[0]]
        out: dict[str, str] = {}
        for cell_ref in _iter_range_cells(normalized):
            value = ws[cell_ref].value
            if value is None:
                continue
            text = str(value).strip()
            if text:
                out[str(cell_ref).upper()] = text
        return out
    finally:
        wb.close()


def _iter_tool_calls(tool_calls: list[Any]) -> list[tuple[str, str, str]]:
    out: list[tuple[str, str, str]] = []
    for i, call in enumerate(tool_calls):
        call_id = f"call_{i}"
        name = ""
        args_json = "{}"
        if hasattr(call, "function"):
            fn = getattr(call, "function", None)
            call_id = str(getattr(call, "id", call_id))
            name = str(getattr(fn, "name", "") or "")
            args_json = str(getattr(fn, "arguments", "{}") or "{}")
        elif isinstance(call, dict):
            call_id = str(call.get("id", call_id))
            fn = call.get("function", {})
            if isinstance(fn, dict):
                name = str(fn.get("name", "") or "")
                args_json = str(fn.get("arguments", "{}") or "{}")
        if name:
            out.append((call_id, name, args_json))
    return out


def _normalize_a1_range(a1: str) -> str:
    s = str(a1 or "").strip().upper().replace(" ", "")
    if not s:
        raise ValueError("empty range")
    # Accept single-cell and normalize to A1:A1
    if ":" not in s:
        _ = range_boundaries(s)
        return f"{s}:{s}"
    c1, r1, c2, r2 = range_boundaries(s)
    # Canonicalize order if reversed
    min_c, max_c = min(c1, c2), max(c1, c2)
    min_r, max_r = min(r1, r2), max(r1, r2)
    left = f"{get_column_letter(min_c)}{min_r}"
    right = f"{get_column_letter(max_c)}{max_r}"
    return f"{left}:{right}"


_A1_RANGE_PATTERN = re.compile(r"([A-Za-z]{1,4}\d{1,7}\s*:\s*[A-Za-z]{1,4}\d{1,7}|[A-Za-z]{1,4}\d{1,7})")


def _dedup_normalized_ranges(items: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for item in items:
        if not item or item in seen:
            continue
        seen.add(item)
        out.append(item)
    return out


def _extract_all_a1_ranges(raw: str) -> list[str]:
    out: list[str] = []
    for match in _A1_RANGE_PATTERN.finditer(str(raw or "")):
        try:
            out.append(_normalize_a1_range(match.group(1)))
        except Exception:
            continue
    return _dedup_normalized_ranges(out)


def _extract_focus_ranges_from_text(raw: str) -> list[str]:
    s = str(raw or "")
    if not s.strip():
        return []
    try:
        obj = json.loads(s)
        if isinstance(obj, dict):
            out: list[str] = []
            for k in ("bad_ranges", "focus_ranges", "ranges", "error_ranges"):
                items = obj.get(k)
                if isinstance(items, list):
                    for item in items:
                        if isinstance(item, str) and item.strip():
                            out.extend(_extract_all_a1_ranges(item))
            for k in ("bad_range", "focus_range", "range", "error_range"):
                v = obj.get(k)
                if isinstance(v, str) and v.strip():
                    out.extend(_extract_all_a1_ranges(v))
            if out:
                return _dedup_normalized_ranges(out)
    except Exception:
        pass
    return _extract_all_a1_ranges(s)


def _extract_focus_range_from_text(raw: str) -> str | None:
    items = _extract_focus_ranges_from_text(raw)
    return items[0] if items else None


def _extract_json_object_from_text(raw: str) -> dict[str, Any]:
    text = str(raw or "").strip()
    if not text:
        return {}
    if "```" in text:
        m = re.search(r"```(?:json)?\s*(.*?)```", text, flags=re.IGNORECASE | re.DOTALL)
        if m:
            text = m.group(1).strip()
    try:
        obj = json.loads(text)
        return obj if isinstance(obj, dict) else {}
    except Exception:
        pass
    m = re.search(r"\{.*\}", text, flags=re.DOTALL)
    if not m:
        return {}
    try:
        obj = json.loads(m.group(0))
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}


def _extract_focus_ranges_from_payload(payload: dict[str, Any]) -> list[str]:
    out: list[str] = []
    items = payload.get("correct_ranges")
    if not isinstance(items, list):
        return out
    for item in items:
        if not isinstance(item, str) or not item.strip():
            continue
        try:
            out.append(_normalize_a1_range(item))
        except Exception:
            continue
    return _dedup_normalized_ranges(out)


def _extract_bad_ranges_raw_from_payload(payload: dict[str, Any]) -> list[str]:
    out: list[str] = []
    items = payload.get("bad_ranges")
    if isinstance(items, list):
        for item in items:
            if isinstance(item, str) and item.strip():
                out.append(str(item).strip())
    single = payload.get("bad_range")
    if isinstance(single, str) and single.strip():
        out.append(str(single).strip())
    return _dedup_normalized_ranges(out)


def _range_intersects(a: str, b: str) -> bool:
    a1, b1, a2, b2 = range_boundaries(_normalize_a1_range(a))
    c1, d1, c2, d2 = range_boundaries(_normalize_a1_range(b))
    rows_overlap = not (b2 < d1 or d2 < b1)
    cols_overlap = not (a2 < c1 or c2 < a1)
    return rows_overlap and cols_overlap


def _iter_range_cells(range_str: str) -> list[str]:
    out: list[str] = []
    for token in [x.strip() for x in str(range_str or "").split(",") if x.strip()]:
        c1, r1, c2, r2 = range_boundaries(_normalize_a1_range(token))
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                out.append(f"{get_column_letter(cc)}{rr}")
    return out


def _expand_focus_range(
    focus_range: str,
    pad_left_cols: int,
    pad_top_rows: int,
    pad_right_cols: int,
    pad_bottom_rows: int,
    clip_min_row: int,
    clip_min_col: int,
    clip_max_row: int,
    clip_max_col: int,
) -> tuple[int, int, int, int]:
    c1, r1, c2, r2 = range_boundaries(_normalize_a1_range(focus_range))
    min_c = max(int(clip_min_col), min(c1, c2) - max(0, int(pad_left_cols)))
    max_c = min(int(clip_max_col), max(c1, c2) + max(0, int(pad_right_cols)))
    min_r = max(int(clip_min_row), min(r1, r2) - max(0, int(pad_top_rows)))
    max_r = min(int(clip_max_row), max(r1, r2) + max(0, int(pad_bottom_rows)))
    if min_r > max_r or min_c > max_c:
        return int(clip_min_row), int(clip_min_col), int(clip_max_row), int(clip_max_col)
    return int(min_r), int(min_c), int(max_r), int(max_c)


def _filter_form_pairs_by_focus_range(form: dict[str, Any], focus_range: str) -> dict[str, Any]:
    pairs = form.get("pairs", []) if isinstance(form, dict) else []
    invalid_slots = form.get("invalid_slots", []) if isinstance(form, dict) else []
    dropped_suspicious_pairs = form.get("dropped_suspicious_pairs", []) if isinstance(form, dict) else []
    out_pairs: list[dict[str, Any]] = []
    out_invalid_slots: list[dict[str, Any]] = []
    out_dropped_suspicious_pairs: list[dict[str, Any]] = []
    for p in pairs:
        if not isinstance(p, dict):
            continue
        vc = p.get("value_cell")
        if not isinstance(vc, str) or not vc.strip():
            continue
        hit = False
        for tok in [x.strip() for x in vc.split(",") if x.strip()]:
            try:
                if _range_intersects(tok, focus_range):
                    hit = True
                    break
            except Exception:
                continue
        if hit:
            out_pairs.append(p)
    for item in invalid_slots:
        if not isinstance(item, dict):
            continue
        vc = item.get("value_cell") or item.get("candidate_value_cell")
        if not isinstance(vc, str) or not vc.strip():
            continue
        hit = False
        for tok in [x.strip() for x in vc.split(",") if x.strip()]:
            try:
                if _range_intersects(tok, focus_range):
                    hit = True
                    break
            except Exception:
                continue
        if hit:
            out_invalid_slots.append(item)
    for item in dropped_suspicious_pairs:
        if not isinstance(item, dict):
            continue
        vc = item.get("value_cell") or item.get("candidate_value_cell")
        if not isinstance(vc, str) or not vc.strip():
            continue
        hit = False
        for tok in [x.strip() for x in vc.split(",") if x.strip()]:
            try:
                if _range_intersects(tok, focus_range):
                    hit = True
                    break
            except Exception:
                continue
        if hit:
            out_dropped_suspicious_pairs.append(item)
    return {
        "pairs": out_pairs,
        "invalid_slots": out_invalid_slots,
        "dropped_suspicious_pairs": out_dropped_suspicious_pairs,
    }


def _filter_slots_json_by_focus_range(slots_json_path: Path, focus_range: str, out_json_path: Path) -> int:
    payload = json.loads(slots_json_path.read_text(encoding="utf-8"))
    slots = payload.get("slots", []) if isinstance(payload, dict) else []
    out_slots: list[dict[str, Any]] = []
    for s in slots:
        if not isinstance(s, dict):
            continue
        hit = False
        for key in ("value_range_a1", "value_anchor_a1", "value_cell"):
            v = s.get(key)
            if not isinstance(v, str) or not v.strip():
                continue
            try:
                if _range_intersects(v, focus_range):
                    hit = True
                    break
            except Exception:
                continue
        if hit:
            out_slots.append(s)
    new_payload = dict(payload) if isinstance(payload, dict) else {}
    new_payload["slots"] = out_slots
    out_json_path.parent.mkdir(parents=True, exist_ok=True)
    out_json_path.write_text(json.dumps(new_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return len(out_slots)


def _read_anchor_value_from_workbook(xlsx_path: Path, sheet_name: str | None, cell_range: str) -> Any:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        token = [x.strip() for x in str(cell_range or "").split(",") if x.strip()]
        if not token:
            return None
        c1, r1, _, _ = range_boundaries(_normalize_a1_range(token[0]))
        merged = _merged_anchor_map(ws)
        anchor = merged.get((r1, c1), (r1, c1, r1, c1))
        return ws.cell(row=anchor[0], column=anchor[1]).value
    finally:
        wb.close()


def _is_blankish_value(value: Any) -> bool:
    if value is None:
        return True
    return isinstance(value, str) and not value.strip()


def _normalize_compare_text(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def _scan_label_and_input_candidates(ws) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    merged = _merged_anchor_map(ws)

    def _border_tags(cell) -> list[str]:
        b = cell.border
        tags: list[str] = []
        if b is None:
            return tags
        if getattr(getattr(b, "top", None), "style", None):
            tags.append("T")
        if getattr(getattr(b, "bottom", None), "style", None):
            tags.append("B")
        if getattr(getattr(b, "left", None), "style", None):
            tags.append("L")
        if getattr(getattr(b, "right", None), "style", None):
            tags.append("R")
        return tags

    min_row, min_col, max_row, max_col = _get_used_bbox(ws)
    labels: list[dict[str, Any]] = []
    input_candidates: list[dict[str, Any]] = []
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            a1 = f"{get_column_letter(c)}{r}"
            tags = _border_tags(cell)
            is_merged = (r, c) in merged
            if isinstance(cell.value, str) and cell.value.strip():
                labels.append({"a1": a1, "text": cell.value.strip()})
            if not tags and not is_merged:
                continue
            is_boxed = set(tags) >= {"T", "B", "L", "R"}
            has_bottom = "B" in tags
            if is_boxed:
                input_candidates.append({"a1": a1, "type": "boxed"})
            elif has_bottom:
                input_candidates.append({"a1": a1, "type": "bottom_line"})
            elif is_merged:
                anchor = merged[(r, c)]
                if r == anchor[0] and c == anchor[1]:
                    input_candidates.append({"a1": a1, "type": "merged_anchor"})
    return labels, input_candidates


def _detect_template_label_value_overwrite(
    before_xlsx_path: Path,
    after_xlsx_path: Path,
    sheet_name: str | None,
    max_findings: int = 12,
) -> list[dict[str, Any]]:
    before_wb = openpyxl.load_workbook(before_xlsx_path, data_only=True)
    after_wb = openpyxl.load_workbook(after_xlsx_path, data_only=True)
    try:
        before_ws = before_wb[sheet_name] if sheet_name else before_wb.active
        after_ws = after_wb[sheet_name] if sheet_name else after_wb.active
        labels, input_candidates = _scan_label_and_input_candidates(before_ws)
        label_cells = {
            str(item.get("a1") or "").upper()
            for item in labels
            if isinstance(item, dict) and str(item.get("a1") or "").strip()
        }
        input_candidate_cells = {
            str(item.get("a1") or "").upper()
            for item in input_candidates
            if isinstance(item, dict) and str(item.get("a1") or "").strip()
        }
        static_label_cells = label_cells - input_candidate_cells
        findings: list[dict[str, Any]] = []
        for label in labels:
            label_a1 = str(label.get("a1") or "").upper()
            template_text = str(label.get("text") or "").strip()
            if not label_a1 or not template_text or label_a1 not in static_label_cells:
                continue
            current_value = after_ws[label_a1].value
            if current_value is None:
                continue
            current_text = str(current_value).strip()
            if not current_text or current_text == template_text:
                continue
            template_norm = _normalize_compare_text(template_text)
            current_norm = _normalize_compare_text(current_text)
            if not template_norm or current_norm == template_norm:
                continue

            label_appended = False
            extra_text = ""
            if current_norm.startswith(template_norm):
                extra_text = current_text[len(template_text):].strip(" :-\t")
                label_appended = bool(extra_text)
            if not label_appended and template_norm.endswith(":") and current_norm.startswith(template_norm.rstrip(":")):
                raw_prefix = template_text.rstrip()
                extra_text = current_text[len(raw_prefix):].strip(" :-\t")
                label_appended = bool(extra_text)
            if not label_appended:
                continue

            candidate_cells = _rank_candidate_cells_for_label(label_a1, input_candidates, labels, limit=3)
            missing_targets: list[str] = []
            for candidate in candidate_cells:
                after_candidate = _read_anchor_value_from_workbook(after_xlsx_path, sheet_name, candidate)
                before_candidate = _read_anchor_value_from_workbook(before_xlsx_path, sheet_name, candidate)
                if _is_blankish_value(after_candidate) or _normalize_compare_text(after_candidate) == _normalize_compare_text(before_candidate):
                    missing_targets.append(str(candidate).upper())
            bad_ranges = [label_a1] + missing_targets[:2]
            findings.append(
                {
                    "label_cell": label_a1,
                    "template_text": template_text,
                    "current_text": current_text,
                    "extra_text": extra_text,
                    "candidate_value_cells": candidate_cells,
                    "missing_target_cells": missing_targets,
                    "bad_ranges": _dedup_normalized_ranges(bad_ranges),
                    "reason": (
                        f"Template label cell {label_a1} was overwritten with combined label+value text "
                        f"('{current_text}') instead of keeping the static label '{template_text}' and filling a separate value cell."
                    ),
                }
            )
            if len(findings) >= max(1, int(max_findings)):
                break
        return findings
    finally:
        before_wb.close()
        after_wb.close()


def _save_screenshot_playwright(html_path: Path, png_path: Path, width: int, height: int, scale: float = 1.0) -> None:
    try:
        from playwright.sync_api import sync_playwright
    except Exception as exc:
        raise PipelineError(
            f"playwright_not_installed: {exc}. Install with: pip install playwright && playwright install chromium"
        ) from exc
    w = int(max(1, round(float(width) * float(scale))))
    h = int(max(1, round(float(height) * float(scale))))
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": w, "height": h, "device_scale_factor": 1})
        page.goto(html_path.resolve().as_uri(), wait_until="load")
        # Stabilize layout to avoid locator.screenshot waiting forever on "element stable".
        page.add_style_tag(
            content="*{animation:none !important;transition:none !important;} html,body{scroll-behavior:auto !important;}"
        )
        page.wait_for_selector("#sheet", state="visible", timeout=120_000)
        el = page.query_selector("#sheet")
        if el is None:
            raise PipelineError("html_css_render_failed: missing #sheet element")
        last_exc = None
        # Try progressive downscale when chromium capture fails on large surfaces.
        for s in [1.0, 0.8, 0.6, 0.45, 0.33, 0.25]:
            try:
                page.evaluate(
                    """(scale) => {
                        const el = document.querySelector('#sheet');
                        if (!el) return;
                        el.style.transformOrigin = 'top left';
                        el.style.transform = `scale(${scale})`;
                    }""",
                    s,
                )
                box = el.bounding_box()
                if not (box and box.get("width", 0) > 0 and box.get("height", 0) > 0):
                    continue
                page.screenshot(
                    path=str(png_path),
                    clip={
                        "x": float(box["x"]),
                        "y": float(box["y"]),
                        "width": float(box["width"]),
                        "height": float(box["height"]),
                    },
                    timeout=120_000,
                )
                last_exc = None
                break
            except Exception as exc:
                last_exc = exc
                continue
        if last_exc is not None:
            raise PipelineError(f"html_css_screenshot_failed: {last_exc}")
        browser.close()


def preprocess_excel_html_css(
    xlsx_path: Path,
    out_dir: Path,
    sheet_name: str | None,
    dpi: int,
    canvas_bg: str = "#f5f5f5",
    png_scale: float = 1.0,
    max_render_rows: int = 400,
    max_render_cols: int = 120,
    focus_range: str | None = None,
    focus_pad_left_cols: int = 0,
    focus_pad_top_rows: int = 0,
    focus_pad_right_cols: int = 0,
    focus_pad_bottom_rows: int = 0,
) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    sheet_png = out_dir / "sheet.png"
    bounds_json = out_dir / "sheet_bounds.json"
    edges_json = out_dir / "sheet_edges.json"
    html_path = out_dir / "sheet.html"
    cell_boxes_json = out_dir / "cell_boxes.json"

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        used_min_row, used_min_col, used_max_row, used_max_col = _get_used_bbox(ws)
        min_row, min_col, max_row, max_col = used_min_row, used_min_col, used_max_row, used_max_col
        if (focus_range or "").strip():
            try:
                min_row, min_col, max_row, max_col = _expand_focus_range(
                    focus_range=str(focus_range),
                    pad_left_cols=int(focus_pad_left_cols),
                    pad_top_rows=int(focus_pad_top_rows),
                    pad_right_cols=int(focus_pad_right_cols),
                    pad_bottom_rows=int(focus_pad_bottom_rows),
                    clip_min_row=int(used_min_row),
                    clip_min_col=int(used_min_col),
                    clip_max_row=int(used_max_row),
                    clip_max_col=int(used_max_col),
                )
            except Exception:
                min_row, min_col, max_row, max_col = used_min_row, used_min_col, used_max_row, used_max_col
        if int(max_render_rows) > 0:
            max_row = min(max_row, min_row + int(max_render_rows) - 1)
        if int(max_render_cols) > 0:
            max_col = min(max_col, min_col + int(max_render_cols) - 1)
        x_edges, y_edges = _compute_edges(ws, min_row, min_col, max_row, max_col, dpi=dpi)

        merged = _merged_anchor_map(ws)
        elements: list[str] = []
        cell_boxes: dict[str, dict[str, object]] = {}
        total_w = int(x_edges[-1])
        total_h = int(y_edges[-1])

        for r in range(min_row, max_row + 1):
            rr = r - min_row
            for c in range(min_col, max_col + 1):
                cc = c - min_col
                a1 = f"{get_column_letter(c)}{r}"
                if (r, c) in merged:
                    r0, c0, r1, c1 = merged[(r, c)]
                    # For focused/local rendering, merged range may be partially outside current bbox.
                    vr0 = max(r0, min_row)
                    vc0 = max(c0, min_col)
                    vr1 = min(r1, max_row)
                    vc1 = min(c1, max_col)
                    if vr0 > vr1 or vc0 > vc1:
                        continue
                    # Render once at the visible top-left corner of the clipped merged region.
                    if not (r == vr0 and c == vc0):
                        continue
                    rr0, cc0 = vr0 - min_row, vc0 - min_col
                    rr1, cc1 = vr1 - min_row, vc1 - min_col
                    x0, y0 = int(x_edges[cc0]), int(y_edges[rr0])
                    x1, y1 = int(x_edges[cc1 + 1]), int(y_edges[rr1 + 1])
                else:
                    x0, y0 = int(x_edges[cc]), int(y_edges[rr])
                    x1, y1 = int(x_edges[cc + 1]), int(y_edges[rr + 1])
                wpx = max(1, x1 - x0)
                hpx = max(1, y1 - y0)

                cell = ws.cell(row=r, column=c)
                txt = _html_escape(_value_to_text(cell.value)).replace("\n", "<br/>")
                font = cell.font
                align = cell.alignment
                fill = cell.fill
                border = cell.border

                bg = "transparent"
                if getattr(fill, "fill_type", None):
                    fg = _rgb_from_color_obj(getattr(fill, "fgColor", None))
                    if fg:
                        bg = fg
                base_font_size = float(getattr(font, "size", 11) or 11)
                font_size = max(14, int(round(base_font_size * 1.2)))
                font_color = _rgb_from_color_obj(getattr(font, "color", None)) or "#111111"
                font_weight = "700" if bool(getattr(font, "bold", False)) else "400"
                font_style = "italic" if bool(getattr(font, "italic", False)) else "normal"
                text_decoration = "underline" if bool(getattr(font, "underline", False)) else "none"
                h_align = getattr(align, "horizontal", None) or "left"
                v_align = getattr(align, "vertical", None) or "center"
                wrap = bool(getattr(align, "wrap_text", False))
                text_align = {"general": "left", "left": "left", "center": "center", "right": "right"}.get(h_align, "left")
                justify_content = {"top": "flex-start", "center": "center", "bottom": "flex-end"}.get(v_align, "center")
                white_space = "normal" if wrap else "nowrap"

                b_left = _css_border_from_side(getattr(border, "left", None))
                b_right = _css_border_from_side(getattr(border, "right", None))
                b_top = _css_border_from_side(getattr(border, "top", None))
                b_bottom = _css_border_from_side(getattr(border, "bottom", None))

                style = (
                    f"left:{x0}px;top:{y0}px;width:{wpx}px;height:{hpx}px;"
                    f"background:{bg};"
                    f"color:{font_color};font-size:{font_size}px;font-weight:{font_weight};font-style:{font_style};"
                    f"text-decoration:{text_decoration};text-align:{text_align};justify-content:{justify_content};white-space:{white_space};"
                    f"border-left:{b_left};border-right:{b_right};border-top:{b_top};border-bottom:{b_bottom};"
                )
                elements.append(f'<div class="cell" data-a1="{a1}" style="{style}"><div class="inner">{txt}</div></div>')
                cell_boxes[a1] = {"cell": a1, "x0": x0, "y0": y0, "x1": x1, "y1": y1, "w": wpx, "h": hpx}

        html = f"""<!doctype html>
<html><head><meta charset="utf-8" />
<style>
html,body{{margin:0;padding:0;background:{canvas_bg};}}
#sheet{{position:relative;width:{total_w}px;height:{total_h}px;background:{canvas_bg};overflow:hidden;box-sizing:border-box;}}
.cell{{position:absolute;display:flex;align-items:stretch;box-sizing:border-box;overflow:hidden;font-family:"Calibri","Arial",sans-serif;line-height:1.2;}}
.inner{{display:block;width:100%;padding:3px 5px;box-sizing:border-box;overflow:hidden;text-overflow:ellipsis;}}
</style></head><body><div id="sheet">{''.join(elements)}</div></body></html>"""

        html_path.write_text(html, encoding="utf-8")
        bounds_json.write_text(
            json.dumps(
                {
                    "sheet": ws.title,
                    "min_row": int(min_row),
                    "min_col": int(min_col),
                    "max_row": int(max_row),
                    "max_col": int(max_col),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        cell_boxes_json.write_text(
            json.dumps({"image_size": {"w": total_w, "h": total_h}, "cell_boxes": cell_boxes}, ensure_ascii=False),
            encoding="utf-8",
        )
    finally:
        wb.close()

    _emit_console_line(
        f"[HTML_RENDER] bbox=({min_row},{min_col})-({max_row},{max_col}) "
        f"size={total_w}x{total_h} scale={png_scale}"
    )
    _save_screenshot_playwright(html_path, sheet_png, total_w, total_h, scale=png_scale)

    if float(png_scale) != 1.0:
        sx = float(png_scale)
        x_edges = [int(round(float(x) * sx)) for x in x_edges]
        y_edges = [int(round(float(y) * sx)) for y in y_edges]

    edges_json.write_text(
        json.dumps({"dpi": int(dpi), "x_edges": x_edges, "y_edges": y_edges}, ensure_ascii=False),
        encoding="utf-8",
    )
    return {
        "sheet_png": sheet_png,
        "bounds_json": bounds_json,
        "edges_json": edges_json,
        "html": html_path,
        "cell_boxes_json": cell_boxes_json,
    }


def _load_planner_policy(path_arg: str) -> dict[str, Any]:
    path_s = str(path_arg or "").strip()
    if not path_s:
        return {}
    path = Path(path_s).resolve()
    if not path.exists():
        raise FileNotFoundError(f"planner_json_path not found: {path}")
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"planner_json_path must contain a JSON object: {path}")
    return payload


def _sanitize_model_io_value(obj: Any) -> Any:
    if isinstance(obj, dict):
        out: dict[str, Any] = {}
        for key, value in obj.items():
            if key == "image_url" and isinstance(value, dict):
                inner = dict(value)
                url = inner.get("url")
                if isinstance(url, str) and url.startswith("data:image/"):
                    inner["url"] = f"[data_url_omitted length={len(url)}]"
                out[key] = _sanitize_model_io_value(inner)
            else:
                out[key] = _sanitize_model_io_value(value)
        return out
    if isinstance(obj, list):
        return [_sanitize_model_io_value(x) for x in obj]
    if isinstance(obj, tuple):
        return [_sanitize_model_io_value(x) for x in obj]
    if isinstance(obj, Path):
        return str(obj)
    return obj


def _dump_model_io(
    base_dir: Path | None,
    name: str,
    request_payload: Any,
    response_payload: Any = None,
    raw_text: str = "",
    usage: dict[str, Any] | None = None,
    extra: dict[str, Any] | None = None,
) -> None:
    if base_dir is None:
        return
    dump_dir = Path(base_dir) / str(name)
    dump_dir.mkdir(parents=True, exist_ok=True)
    request_obj = _sanitize_model_io_value(request_payload)
    response_obj = _sanitize_model_io_value(response_payload)
    summary = {
        "request": request_obj,
        "response": response_obj,
        "raw_text": str(raw_text or ""),
        "usage": usage or {},
        "extra": _sanitize_model_io_value(extra or {}),
    }
    (dump_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    (dump_dir / "request.json").write_text(
        json.dumps(request_obj, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    if response_obj is not None:
        (dump_dir / "response.json").write_text(
            json.dumps(response_obj, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
    if str(raw_text or "").strip():
        (dump_dir / "response.txt").write_text(str(raw_text), encoding="utf-8")


def _planner_string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    out: list[str] = []
    for item in value:
        text = str(item or "").strip()
        if text:
            out.append(text)
    return out


def _planner_as_bool(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    if isinstance(value, str):
        lowered = value.strip().lower()
        if lowered in {"1", "true", "yes", "y"}:
            return True
        if lowered in {"0", "false", "no", "n"}:
            return False
    return default


def _planner_as_float(value: Any, default: float = 0.0) -> float:
    try:
        out = float(value)
    except Exception:
        return default
    return max(0.0, min(1.0, out))


def _planner_as_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except Exception:
        return default


def _normalize_generated_planner_policy(plan: dict[str, Any]) -> dict[str, Any]:
    global_policies = plan.get("global_policies") if isinstance(plan.get("global_policies"), dict) else {}
    repeated_block_policy = plan.get("repeated_block_policy") if isinstance(plan.get("repeated_block_policy"), dict) else {}
    workflow_policy = plan.get("workflow_policy") if isinstance(plan.get("workflow_policy"), dict) else {}

    field_groups = []
    for item in plan.get("field_groups") or []:
        if not isinstance(item, dict):
            continue
        group_name = str(item.get("group_name") or "").strip()
        fields = _planner_string_list(item.get("fields"))
        if not group_name and not fields:
            continue
        field_groups.append(
            {
                "group_name": group_name or "group",
                "fill_priority": max(1, _planner_as_int(item.get("fill_priority"), len(field_groups) + 1)),
                "fields": fields,
            }
        )
    field_groups.sort(key=lambda item: int(item.get("fill_priority", 999999)))

    checkbox_policies = []
    for item in plan.get("checkbox_policies") or []:
        if not isinstance(item, dict):
            continue
        target = str(item.get("target") or "").strip()
        selection_mode = str(item.get("selection_mode") or "").strip()
        if target or selection_mode:
            checkbox_policies.append({"target": target, "selection_mode": selection_mode or "single"})

    high_confidence_fields = []
    for item in plan.get("high_confidence_fields") or []:
        if not isinstance(item, dict):
            continue
        target = str(item.get("target") or "").strip()
        if not target:
            continue
        high_confidence_fields.append(
            {
                "target": target,
                "value_hint": str(item.get("value_hint") or "").strip() or None,
                "evidence_span": str(item.get("evidence_span") or "").strip() or None,
                "confidence": _planner_as_float(item.get("confidence"), 0.0),
            }
        )

    uncertain_fields = []
    for item in plan.get("uncertain_fields") or []:
        if not isinstance(item, dict):
            continue
        target = str(item.get("target") or "").strip()
        if not target:
            continue
        uncertain_fields.append(
            {
                "target": target,
                "reason": str(item.get("reason") or "").strip(),
                "evidence_span": str(item.get("evidence_span") or "").strip() or None,
                "confidence": _planner_as_float(item.get("confidence"), 0.0),
                "action": str(item.get("action") or "leave_blank").strip() or "leave_blank",
            }
        )

    background_reads = []
    for item in plan.get("background_reads") or []:
        if not isinstance(item, dict):
            continue
        target = str(item.get("target") or "").strip()
        if not target:
            continue
        stage = str(item.get("stage") or "any").strip().lower() or "any"
        if stage in {"all", "*"}:
            stage = "any"
        if stage not in {"any", "first_pass", "reflect", "refill"}:
            stage = "any"
        background_reads.append(
            {
                "stage": stage,
                "target": target,
                "reason": str(item.get("reason") or "").strip(),
                "max_radius": max(1, min(4, _planner_as_int(item.get("max_radius"), 2))),
            }
        )

    return {
        "form_type": str(plan.get("form_type") or "").strip() or "generic_form",
        "global_policies": {
            "leave_unspecified_blank": _planner_as_bool(global_policies.get("leave_unspecified_blank"), True),
            "preserve_existing_content": _planner_as_bool(global_policies.get("preserve_existing_content"), True),
            "avoid_guessing": _planner_as_bool(global_policies.get("avoid_guessing"), True),
        },
        "workflow_policy": {
            "read_background_before_first_pass": _planner_as_bool(
                workflow_policy.get("read_background_before_first_pass"),
                False,
            ),
            "read_background_before_refill": _planner_as_bool(
                workflow_policy.get("read_background_before_refill"),
                False,
            ),
            "prefer_hints_when_available": _planner_as_bool(
                workflow_policy.get("prefer_hints_when_available"),
                True,
            ),
            "run_reflect_after_first_pass": _planner_as_bool(
                workflow_policy.get("run_reflect_after_first_pass"),
                True,
            ),
            "stop_when_no_high_risk_regions": _planner_as_bool(
                workflow_policy.get("stop_when_no_high_risk_regions"),
                True,
            ),
            "stop_when_two_rounds_no_actual_edits": _planner_as_bool(
                workflow_policy.get("stop_when_two_rounds_no_actual_edits"),
                True,
            ),
        },
        "field_groups": field_groups,
        "blank_policies": [],
        "do_not_fill_sections": _planner_string_list(plan.get("do_not_fill_sections")),
        "checkbox_policies": checkbox_policies,
        "repeated_block_policy": {
            "has_repeated_blocks": _planner_as_bool(repeated_block_policy.get("has_repeated_blocks"), False),
            "block_key": str(repeated_block_policy.get("block_key") or "").strip(),
            "expected_count": (
                None
                if repeated_block_policy.get("expected_count") in (None, "", "null")
                else _planner_as_int(repeated_block_policy.get("expected_count"), 0)
            ),
        },
        "high_confidence_fields": high_confidence_fields,
        "uncertain_fields": uncertain_fields,
        "background_reads": background_reads,
        "notes": _planner_string_list(plan.get("notes")),
    }


def _scan_structure_for_planner(xlsx_path: Path, sheet_name: str | None) -> dict[str, Any]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        min_row, min_col, max_row, max_col = _get_used_bbox(ws)
        text_cells = []
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if not text:
                    continue
                text_cells.append({"a1": str(cell.coordinate), "text": text[:120]})
        merged_ranges = [str(rng) for rng in list(ws.merged_cells.ranges)[:200]]
        return {
            "available": True,
            "sheet": ws.title,
            "bbox": {
                "min_row": int(min_row),
                "min_col": int(min_col),
                "max_row": int(max_row),
                "max_col": int(max_col),
            },
            "merged_ranges": merged_ranges,
            "text_cells": text_cells[:200],
        }
    finally:
        wb.close()


def _build_planner_prompt(instruction: str, struct_summary: dict[str, Any], sheet_name: str | None) -> str:
    return (
        "You are planning how an agent should fill a spreadsheet form.\n"
        "Produce a concise, execution-oriented controller plan in JSON only.\n"
        "Do not fill the form. Do not invent values. Work only from the instruction and workbook structure.\n"
        "The plan must act as a control layer for a later executor.\n"
        "Decide when the executor should read local background context, when hints should be preferred, and whether reflect should run after first pass.\n"
        "Focus on workflow control, fill order, repeated-block awareness, and local background inspection.\n\n"
        "Return strict JSON with this schema:\n"
        "{\n"
        '  "form_type": <string>,\n'
        '  "workflow_policy": {"read_background_before_first_pass": <bool>, "read_background_before_refill": <bool>, "prefer_hints_when_available": <bool>, "run_reflect_after_first_pass": <bool>, "stop_when_no_high_risk_regions": <bool>, "stop_when_two_rounds_no_actual_edits": <bool>},\n'
        '  "field_groups": [{"group_name": <string>, "fill_priority": <int>, "fields": [<string>, ...]}],\n'
        '  "repeated_block_policy": {"has_repeated_blocks": <bool>, "block_key": <string>, "expected_count": <int|null>},\n'
        '  "high_confidence_fields": [{"target": <string>, "value_hint": <string|null>, "evidence_span": <string|null>, "confidence": <0..1>}],\n'
        '  "background_reads": [{"stage": <\"first_pass\"|\"reflect\"|\"refill\"|\"any\">, "target": <string>, "reason": <string>, "max_radius": <1..4>}],\n'
        '  "notes": [<string>, ...]\n'
        "}\n\n"
        "Rules:\n"
        "- Prefer short machine-actionable targets over long prose.\n"
        "- Use workflow_policy to control executor behavior, not just summarize the form.\n"
        "- Set read_background_before_first_pass/read_background_before_refill=true when local labels, repeated sections, or ambiguous layout should be inspected before writing.\n"
        "- Set prefer_hints_when_available=false only when hint usage is likely to be misleading.\n"
        "- Set run_reflect_after_first_pass=false only when the form is simple and the instruction is already sufficient.\n"
        "- background_reads should name short targets or A1 ranges the executor should inspect before writing.\n"
        "- If repeated blocks exist, state the expected count if you can infer it; otherwise use null.\n"
        "- high_confidence_fields should only include fields strongly supported by the instruction.\n"
        "- When a dark/filled label or section header is followed by an empty merged cell/range, treat the empty merged range as the likely write target, not the label/header cell.\n"
        "- For topic-style labels, section titles, or document-library forms, the later executor may summarize the matching source section into the nearby blank value range when the source section is clearly matched.\n"
        "- Keep notes short and operational.\n\n"
        f"Target sheet: {sheet_name or 'auto'}\n"
        f"Instruction:\n{instruction.strip()}\n\n"
        f"Workbook structure summary:\n{json.dumps(struct_summary, ensure_ascii=False, indent=2)}\n"
    )


def _request_planner_policy(
    instruction: str,
    struct_summary: dict[str, Any],
    sheet_name: str | None,
    model: str,
    base_url: str | None,
    api_key: str | None,
    temperature: float,
    model_io_dir: Path | None = None,
) -> dict[str, Any]:
    prompt = _build_planner_prompt(instruction, struct_summary, sheet_name)
    messages = [
        {"role": "system", "content": "You are a spreadsheet form-filling planner. Return only valid JSON."},
        {"role": "user", "content": prompt},
    ]
    client = build_client(base_url, api_key)
    resp = client.chat.completions.create(
        model=model,
        temperature=float(temperature),
        messages=messages,
    )
    raw = str(resp.choices[0].message.content or "").strip()
    usage = _usage_to_dict(resp)
    plan = _normalize_generated_planner_policy(_extract_json_object_from_text(raw))
    plan["_raw"] = raw
    plan["_usage"] = usage
    _dump_model_io(
        model_io_dir,
        "planner",
        request_payload={"model": model, "temperature": float(temperature), "messages": messages},
        response_payload={"content": raw},
        raw_text=raw,
        usage=usage,
        extra={"sheet_name": sheet_name, "struct_summary": struct_summary},
    )
    return plan


def _build_planner_policy_preview(plan: dict[str, Any]) -> str:
    payload = {k: v for k, v in plan.items() if not str(k).startswith("_")}
    policies: list[str] = []
    workflow = payload.get("workflow_policy") or {}
    if workflow.get("read_background_before_first_pass"):
        policies.append("- Controller: read local background before first-pass writing.")
    if workflow.get("read_background_before_refill"):
        policies.append("- Controller: read local background before refill/patch writing.")
    if not workflow.get("prefer_hints_when_available", True):
        policies.append("- Controller: do not automatically trust detector or reflect hints.")
    if not workflow.get("run_reflect_after_first_pass", True):
        policies.append("- Controller: skip reflect after first pass unless explicitly forced elsewhere.")
    for item in payload.get("background_reads") or []:
        if not isinstance(item, dict):
            continue
        target = str(item.get("target") or "").strip()
        stage = str(item.get("stage") or "any").strip()
        reason = str(item.get("reason") or "").strip()
        if target:
            line = f"- Background read ({stage}): inspect '{target}'."
            if reason:
                line += f" Reason: {reason}."
            policies.append(line)
    text = "\n".join(policies) if policies else "(none)"
    return (
        "[Execution policies - high priority]\n"
        f"{text}\n\n"
        "[Structured fill hints]\n"
        f"{json.dumps(payload, ensure_ascii=False, indent=2)}\n"
    )


def _planner_policy_stage_payload(planner_policy: dict[str, Any] | None, stage: str) -> dict[str, Any]:
    plan = planner_policy if isinstance(planner_policy, dict) else {}
    if not plan:
        return {}
    payload: dict[str, Any] = {
        "form_type": plan.get("form_type"),
        "workflow_policy": plan.get("workflow_policy", {}),
    }
    if stage == "first_pass":
        keys = (
            "field_groups",
            "repeated_block_policy",
            "high_confidence_fields",
            "notes",
        )
    elif stage == "reflect":
        keys = (
            "repeated_block_policy",
            "notes",
        )
    else:
        keys = (
            "field_groups",
            "repeated_block_policy",
            "high_confidence_fields",
            "notes",
        )
    for key in keys:
        if key in plan:
            payload[key] = plan.get(key)
    background_reads = _planner_stage_background_reads(plan, stage)
    if background_reads:
        payload["background_reads"] = background_reads
    return payload


def _planner_policy_lines(planner_policy: dict[str, Any] | None, stage: str) -> list[str]:
    payload = _planner_policy_stage_payload(planner_policy, stage)
    if not payload:
        return []
    lines: list[str] = []
    workflow_policy = payload.get("workflow_policy") if isinstance(payload.get("workflow_policy"), dict) else {}
    if stage == "first_pass" and workflow_policy.get("read_background_before_first_pass"):
        lines.append("Before first-pass writing, inspect planner-selected local background context.")
    if stage == "refill" and workflow_policy.get("read_background_before_refill"):
        lines.append("Before refill, inspect planner-selected local background context for this region.")
    if not workflow_policy.get("prefer_hints_when_available", True):
        lines.append("Do not automatically rely on detector or reflect hints.")
    if stage == "first_pass" and not workflow_policy.get("run_reflect_after_first_pass", True):
        lines.append("Planner requests skipping reflect after first pass if the result looks acceptable.")

    repeated = payload.get("repeated_block_policy") if isinstance(payload.get("repeated_block_policy"), dict) else {}
    if repeated.get("has_repeated_blocks"):
        block_key = str(repeated.get("block_key") or "").strip()
        expected_count = repeated.get("expected_count")
        if block_key and expected_count not in (None, "", 0):
            lines.append(f"Repeated block policy: '{block_key}' appears about {expected_count} times.")
        elif block_key:
            lines.append(f"Repeated block policy: repeated block key is '{block_key}'.")
        else:
            lines.append("Repeated block policy: the form contains repeated blocks.")

    field_groups = payload.get("field_groups") if isinstance(payload.get("field_groups"), list) else []
    if stage in {"first_pass", "refill"} and field_groups:
        ordered: list[str] = []
        for item in field_groups[:8]:
            if not isinstance(item, dict):
                continue
            name = str(item.get("group_name") or "").strip()
            priority = int(item.get("fill_priority") or 0)
            if name:
                ordered.append(f"{priority}:{name}" if priority else name)
        if ordered:
            lines.append("Suggested fill order by group: " + " -> ".join(ordered) + ".")

    background_reads = payload.get("background_reads") if isinstance(payload.get("background_reads"), list) else []
    for item in background_reads[:6]:
        if not isinstance(item, dict):
            continue
        target = str(item.get("target") or "").strip()
        reason = str(item.get("reason") or "").strip()
        if not target:
            continue
        line = f"Read local background around '{target}' before writing."
        if reason:
            line += f" Reason: {reason}."
        lines.append(line)

    return lines


def _planner_policy_prompt_block(planner_policy: dict[str, Any] | None, stage: str) -> str:
    payload = _planner_policy_stage_payload(planner_policy, stage)
    if not payload:
        return ""
    lines = _planner_policy_lines(planner_policy, stage)
    block = "Planner controller summary:\n"
    if lines:
        block += "\n".join(f"- {line}" for line in lines) + "\n"
    else:
        block += "(none)\n"
    block += (
        "\nPlanner structured controller object:\n"
        f"{json.dumps(payload, ensure_ascii=False)}\n"
    )
    return block


def _planner_blank_like_targets(planner_policy: dict[str, Any] | None) -> list[dict[str, str]]:
    return []


def _planner_workflow_policy(planner_policy: dict[str, Any] | None) -> dict[str, bool]:
    plan = planner_policy if isinstance(planner_policy, dict) else {}
    workflow = plan.get("workflow_policy") if isinstance(plan.get("workflow_policy"), dict) else {}
    return {
        "read_background_before_first_pass": _planner_as_bool(
            workflow.get("read_background_before_first_pass"),
            False,
        ),
        "read_background_before_refill": _planner_as_bool(
            workflow.get("read_background_before_refill"),
            False,
        ),
        "prefer_hints_when_available": _planner_as_bool(
            workflow.get("prefer_hints_when_available"),
            True,
        ),
        "run_reflect_after_first_pass": _planner_as_bool(
            workflow.get("run_reflect_after_first_pass"),
            True,
        ),
        "stop_when_no_high_risk_regions": _planner_as_bool(
            workflow.get("stop_when_no_high_risk_regions"),
            True,
        ),
        "stop_when_two_rounds_no_actual_edits": _planner_as_bool(
            workflow.get("stop_when_two_rounds_no_actual_edits"),
            True,
        ),
    }


def _planner_stage_background_reads(planner_policy: dict[str, Any] | None, stage: str) -> list[dict[str, Any]]:
    plan = planner_policy if isinstance(planner_policy, dict) else {}
    out: list[dict[str, Any]] = []
    for item in plan.get("background_reads", []) or []:
        if not isinstance(item, dict):
            continue
        item_stage = str(item.get("stage") or "any").strip().lower() or "any"
        if item_stage in {"all", "*"}:
            item_stage = "any"
        if item_stage not in {"any", stage}:
            continue
        target = str(item.get("target") or "").strip()
        if not target:
            continue
        out.append(
            {
                "stage": item_stage,
                "target": target,
                "reason": str(item.get("reason") or "").strip(),
                "max_radius": max(1, min(4, _planner_as_int(item.get("max_radius"), 2))),
            }
        )
    return out


def _a1_coord(a1: str) -> tuple[int, int] | None:
    try:
        col_s, row = coordinate_from_string(str(a1 or "").upper())
        return int(row), int(column_index_from_string(col_s))
    except Exception:
        return None


def _a1_range_from_bounds(min_row: int, min_col: int, max_row: int, max_col: int) -> str:
    return f"{get_column_letter(int(min_col))}{int(min_row)}:{get_column_letter(int(max_col))}{int(max_row)}"


def _soft_key_match_score(a: str, b: str) -> float:
    aa = normalize_key(str(a or ""))
    bb = normalize_key(str(b or ""))
    if not aa or not bb:
        return 0.0
    if aa == bb:
        return 1.0
    ratio = SequenceMatcher(None, aa, bb).ratio()
    ta = set(aa.split())
    tb = set(bb.split())
    overlap = (len(ta & tb) / max(1, len(ta | tb))) if (ta or tb) else 0.0
    contains = 1.0 if (aa in bb or bb in aa) and min(len(aa), len(bb)) >= 4 else 0.0
    return max(ratio, overlap, contains)


def _rank_candidate_cells_for_label(
    label_a1: str,
    input_candidates: list[dict[str, Any]],
    labels: list[dict[str, Any]],
    limit: int = 2,
) -> list[str]:
    label_pos = _a1_coord(label_a1)
    if label_pos is None:
        return []
    lr, lc = label_pos
    label_cells = {
        str(item.get("a1") or "").upper()
        for item in labels
        if isinstance(item, dict) and str(item.get("a1") or "").strip()
    }
    ranked: list[tuple[float, str]] = []
    seen: set[str] = set()
    for item in input_candidates:
        if not isinstance(item, dict):
            continue
        a1 = str(item.get("a1") or "").upper()
        if not a1 or a1 == label_a1.upper() or a1 in label_cells:
            continue
        pos = _a1_coord(a1)
        if pos is None:
            continue
        r, c = pos
        dr = abs(r - lr)
        dc = c - lc
        if dc >= 1:
            score = 2.2 / (1.0 + 0.35 * dr + dc)
        elif r > lr:
            score = 1.3 / (1.0 + 0.45 * abs(c - lc) + (r - lr))
        else:
            score = 0.5 / (1.0 + dr + abs(dc))
        ctype = str(item.get("type") or "")
        if ctype == "boxed":
            score += 0.12
        elif ctype == "bottom_line":
            score += 0.08
        elif ctype == "merged_anchor":
            score += 0.05
        if score <= 0.24:
            continue
        if a1 in seen:
            continue
        seen.add(a1)
        ranked.append((score, a1))
    ranked.sort(key=lambda x: x[0], reverse=True)
    return [a1 for _, a1 in ranked[: max(1, int(limit))]]


def _resolve_planner_guard_cells(
    struct_summary: dict[str, Any] | None,
    planner_policy: dict[str, Any] | None,
) -> tuple[dict[str, str], list[dict[str, Any]]]:
    struct = struct_summary if isinstance(struct_summary, dict) else {}
    labels = struct.get("labels")
    input_candidates = struct.get("input_candidates")
    if not isinstance(labels, list) or not isinstance(input_candidates, list):
        return {}, []
    targets = _planner_blank_like_targets(planner_policy)
    if not targets:
        return {}, []
    protected: dict[str, str] = {}
    matches: list[dict[str, Any]] = []
    for item in targets:
        target_text = str(item.get("target") or "").strip()
        if not target_text:
            continue
        best_label = None
        best_score = 0.0
        for label in labels:
            if not isinstance(label, dict):
                continue
            label_text = str(label.get("text") or "").strip()
            label_a1 = str(label.get("a1") or "").strip()
            if not label_text or not label_a1:
                continue
            score = _soft_key_match_score(target_text, label_text)
            if score > best_score:
                best_score = score
                best_label = label
        if not isinstance(best_label, dict) or best_score < 0.74:
            continue
        label_a1 = str(best_label.get("a1") or "").upper()
        candidate_cells = _rank_candidate_cells_for_label(label_a1, input_candidates, labels, limit=2)
        if not candidate_cells:
            continue
        reason_base = (
            f"planner {item.get('kind')} target '{target_text}' should remain blank"
        )
        reason_detail = str(item.get("reason") or "").strip()
        if reason_detail:
            reason_base += f" ({reason_detail})"
        for cell in candidate_cells:
            protected.setdefault(cell.upper(), reason_base)
        matches.append(
            {
                "target": target_text,
                "kind": item.get("kind"),
                "matched_label_a1": label_a1,
                "matched_label_text": str(best_label.get("text") or ""),
                "match_score": round(best_score, 4),
                "protected_cells": candidate_cells,
            }
        )
    return protected, matches


def _collect_planner_background_context(
    xlsx_path: Path,
    sheet_name: str | None,
    planner_policy: dict[str, Any] | None,
    stage: str,
    focus_range: str | None = None,
    max_items: int = 4,
) -> dict[str, Any]:
    workflow = _planner_workflow_policy(planner_policy)
    stage_reads = _planner_stage_background_reads(planner_policy, stage)
    should_read = bool(stage_reads)
    if stage == "first_pass":
        should_read = should_read or bool(workflow.get("read_background_before_first_pass"))
    elif stage == "refill":
        should_read = should_read or bool(workflow.get("read_background_before_refill"))
    if not should_read:
        return {}

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active
        min_row, min_col, max_row, max_col = _get_used_bbox(ws)
        merged = _merged_anchor_map(ws)
        labels: list[dict[str, Any]] = []
        input_candidates: list[dict[str, Any]] = []
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                value = cell.value
                text = str(value).strip() if value is not None else ""
                a1 = f"{get_column_letter(c)}{r}"
                if text:
                    labels.append({"a1": a1, "text": text[:160]})
                tags: list[str] = []
                border = cell.border
                if border is not None:
                    if getattr(getattr(border, "top", None), "style", None):
                        tags.append("T")
                    if getattr(getattr(border, "bottom", None), "style", None):
                        tags.append("B")
                    if getattr(getattr(border, "left", None), "style", None):
                        tags.append("L")
                    if getattr(getattr(border, "right", None), "style", None):
                        tags.append("R")
                is_merged = (r, c) in merged
                if not tags and not is_merged:
                    continue
                if set(tags) >= {"T", "B", "L", "R"}:
                    input_candidates.append({"a1": a1, "type": "boxed"})
                elif "B" in tags:
                    input_candidates.append({"a1": a1, "type": "bottom_line"})
                elif is_merged:
                    anchor = merged[(r, c)]
                    if r == anchor[0] and c == anchor[1]:
                        input_candidates.append({"a1": a1, "type": "merged_anchor"})

        if not stage_reads:
            auto_reads: list[dict[str, Any]] = []
            for item in (planner_policy or {}).get("high_confidence_fields", []) or []:
                if not isinstance(item, dict):
                    continue
                target = str(item.get("target") or "").strip()
                if target:
                    auto_reads.append({"stage": stage, "target": target, "reason": "planner high_confidence field", "max_radius": 2})
            for item in (planner_policy or {}).get("uncertain_fields", []) or []:
                if not isinstance(item, dict):
                    continue
                target = str(item.get("target") or "").strip()
                if target:
                    auto_reads.append({"stage": stage, "target": target, "reason": "planner uncertain field", "max_radius": 2})
            if focus_range and not auto_reads:
                auto_reads.append({"stage": stage, "target": focus_range, "reason": "local focus range context", "max_radius": 2})
            stage_reads = []
            seen_targets: set[str] = set()
            for item in auto_reads:
                target = str(item.get("target") or "").strip()
                if not target or target in seen_targets:
                    continue
                seen_targets.add(target)
                stage_reads.append(item)

        items: list[dict[str, Any]] = []
        focus_norm = _normalize_a1_range(focus_range) if (focus_range or "").strip() else ""
        focus_guard = ""
        if focus_norm:
            fg_min_row, fg_min_col, fg_max_row, fg_max_col = _expand_focus_range(
                focus_norm,
                pad_left_cols=2,
                pad_top_rows=1,
                pad_right_cols=2,
                pad_bottom_rows=1,
                clip_min_row=min_row,
                clip_min_col=min_col,
                clip_max_row=max_row,
                clip_max_col=max_col,
            )
            focus_guard = _a1_range_from_bounds(fg_min_row, fg_min_col, fg_max_row, fg_max_col)
        for item in stage_reads[: max(1, int(max_items))]:
            target = str(item.get("target") or "").strip()
            reason = str(item.get("reason") or "").strip()
            radius = max(1, min(4, _planner_as_int(item.get("max_radius"), 2)))
            if not target:
                continue
            context_range = ""
            matched_label: dict[str, Any] | None = None
            candidate_cells: list[str] = []
            match_score = 0.0
            match_kind = "target_text"
            if re.fullmatch(r"[A-Za-z]+[0-9]+(?::[A-Za-z]+[0-9]+)?", target):
                ctx_min_row, ctx_min_col, ctx_max_row, ctx_max_col = _expand_focus_range(
                    target,
                    pad_left_cols=max(2, radius * 2),
                    pad_top_rows=radius,
                    pad_right_cols=max(2, radius * 2),
                    pad_bottom_rows=radius,
                    clip_min_row=min_row,
                    clip_min_col=min_col,
                    clip_max_row=max_row,
                    clip_max_col=max_col,
                )
                context_range = _a1_range_from_bounds(ctx_min_row, ctx_min_col, ctx_max_row, ctx_max_col)
                match_kind = "a1_range"
            else:
                for label in labels:
                    score = _soft_key_match_score(target, str(label.get("text") or ""))
                    if score > match_score:
                        match_score = score
                        matched_label = label
                if matched_label is None or match_score < 0.60:
                    continue
                label_a1 = str(matched_label.get("a1") or "").upper()
                candidate_cells = _rank_candidate_cells_for_label(label_a1, input_candidates, labels, limit=2)
                label_pos = _a1_coord(label_a1)
                if label_pos is None:
                    continue
                lr, lc = label_pos
                right_cols = []
                for cand in candidate_cells:
                    pos = _a1_coord(cand)
                    if pos is not None:
                        right_cols.append(pos[1])
                max_context_col = max([lc + max(4, radius * 2)] + right_cols)
                context_range = _a1_range_from_bounds(
                    max(min_row, lr - radius),
                    max(min_col, lc - 2),
                    min(max_row, lr + radius),
                    min(max_col, max_context_col),
                )
            if focus_guard:
                try:
                    if not _range_intersects(context_range, focus_guard):
                        continue
                except Exception:
                    pass
            c1, r1, c2, r2 = range_boundaries(_normalize_a1_range(context_range))
            nonempty_cells: list[dict[str, Any]] = []
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    cell = ws.cell(row=rr, column=cc)
                    if cell.value is None:
                        continue
                    text = str(cell.value).strip()
                    if not text:
                        continue
                    nonempty_cells.append(
                        {
                            "a1": f"{get_column_letter(cc)}{rr}",
                            "value": text[:160],
                        }
                    )
            items.append(
                {
                    "target": target,
                    "reason": reason,
                    "match_kind": match_kind,
                    "matched_label_a1": str(matched_label.get("a1") or "") if isinstance(matched_label, dict) else None,
                    "matched_label_text": str(matched_label.get("text") or "") if isinstance(matched_label, dict) else None,
                    "match_score": round(match_score, 4),
                    "context_range": context_range,
                    "candidate_value_cells": candidate_cells,
                    "nonempty_cells": nonempty_cells[:40],
                }
            )
        if not items:
            return {}
        return {
            "stage": stage,
            "focus_range": focus_norm or None,
            "workflow_policy": workflow,
            "items": items,
        }
    finally:
        wb.close()


def agent_fill(
    template_xlsx: Path,
    output_xlsx: Path,
    instruction: str,
    form: dict,
    model: str,
    base_url: str | None,
    api_key: str | None,
    tool_mode: str,
    max_steps: int,
    skills_text: str = "",
    mapping_mode: str = "off",
    sheet_name: str | None = None,
    allowed_write_range: str | None = None,
    edit_log_jsonl: Path | None = None,
    round_idx: int = 0,
    reflect_hints: list[dict[str, Any]] | None = None,
    context_summary: dict[str, Any] | None = None,
    compare_before_after_pngs: dict[str, str] | None = None,
    history_dump_path: Path | None = None,
    temperature: float = 0.0,
    protected_write_cells: dict[str, str] | None = None,
    planner_policy: dict[str, Any] | None = None,
    planner_stage: str = "first_pass",
    model_io_dir: Path | None = None,
    agent_profile: str = "guided",
    enforce_label_guard: bool = True,
    enforce_bottom_border_guard: bool = False,
    enforce_protected_cells: bool = True,
    role_template_xlsx: Path | None = None,
    rendered_html_path: Path | None = None,
    rendered_html_max_chars: int = 0,
    form_context_mode: str = "struct",
) -> dict[str, Any]:
    wb = openpyxl.load_workbook(template_xlsx)
    ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active
    role_wb = None
    role_ws = ws
    if role_template_xlsx is not None:
        try:
            role_wb = openpyxl.load_workbook(role_template_xlsx)
            role_ws = role_wb[sheet_name] if (sheet_name and sheet_name in role_wb.sheetnames) else role_wb.active
        except Exception:
            role_wb = None
            role_ws = ws
    dirty = False
    code_executed = False
    code_exec_errors: list[str] = []
    write_log: list[dict[str, Any]] = []
    write_attempt_log: list[dict[str, Any]] = []
    current_step_idx = 0
    usage_totals = {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0, "calls_with_usage": 0}
    normalized_allowed_range = _normalize_a1_range(allowed_write_range) if (allowed_write_range or "").strip() else None
    normalized_protected_cells = (
        {str(k).upper(): str(v) for k, v in (protected_write_cells or {}).items() if str(k).strip()}
        if enforce_protected_cells
        else {}
    )
    if edit_log_jsonl is not None:
        edit_log_jsonl.parent.mkdir(parents=True, exist_ok=True)

    def _resolve_write_target(cell_ref: str) -> str:
        target = str(cell_ref or "").upper()
        for merged_range in ws.merged_cells.ranges:
            try:
                if target in merged_range:
                    return str(merged_range.start_cell.coordinate).upper()
            except Exception:
                continue
        return target

    def _border_tags(cell) -> list[str]:
        b = cell.border
        tags: list[str] = []
        if b is None:
            return tags
        if getattr(getattr(b, "top", None), "style", None):
            tags.append("T")
        if getattr(getattr(b, "bottom", None), "style", None):
            tags.append("B")
        if getattr(getattr(b, "left", None), "style", None):
            tags.append("L")
        if getattr(getattr(b, "right", None), "style", None):
            tags.append("R")
        return tags

    def _color_token(color) -> str:
        if color is None:
            return ""
        color_type = str(getattr(color, "type", "") or "").strip()
        rgb = str(getattr(color, "rgb", "") or "").strip()
        indexed = getattr(color, "indexed", None)
        theme = getattr(color, "theme", None)
        tint = getattr(color, "tint", None)
        if color_type == "rgb" and rgb:
            return rgb
        if color_type == "indexed" and indexed is not None:
            return f"indexed:{indexed}"
        if color_type == "theme" and theme is not None:
            token = f"theme:{theme}"
            if tint not in (None, 0, 0.0):
                token += f":tint={tint}"
            return token
        if rgb:
            return rgb
        return ""

    def _cell_style_summary(cell) -> dict[str, Any]:
        style: dict[str, Any] = {}

        fill = getattr(cell, "fill", None)
        if fill is not None:
            fill_pattern = str(getattr(fill, "patternType", "") or getattr(fill, "fill_type", "") or "").strip()
            fill_fg = _color_token(getattr(fill, "fgColor", None))
            fill_bg = _color_token(getattr(fill, "bgColor", None))
            fill_payload: dict[str, Any] = {}
            if fill_pattern and fill_pattern.lower() != "none":
                fill_payload["pattern"] = fill_pattern
                fill_payload["is_filled"] = True
            elif fill_fg or fill_bg:
                fill_payload["is_filled"] = bool(fill_fg or fill_bg)
            if fill_fg:
                fill_payload["fgColor"] = fill_fg
            if fill_bg:
                fill_payload["bgColor"] = fill_bg
            if fill_payload:
                style["fill"] = fill_payload

        font = getattr(cell, "font", None)
        if font is not None:
            font_payload: dict[str, Any] = {}
            if getattr(font, "name", None):
                font_payload["name"] = str(font.name)
            if getattr(font, "sz", None) is not None:
                try:
                    font_payload["size"] = float(font.sz)
                except Exception:
                    font_payload["size"] = font.sz
            if bool(getattr(font, "b", False)):
                font_payload["bold"] = True
            if bool(getattr(font, "i", False)):
                font_payload["italic"] = True
            if bool(getattr(font, "strike", False)):
                font_payload["strike"] = True
            underline = str(getattr(font, "u", "") or "").strip()
            if underline:
                font_payload["underline"] = underline
            font_color = _color_token(getattr(font, "color", None))
            if font_color:
                font_payload["color"] = font_color
            if font_payload:
                style["font"] = font_payload

        alignment = getattr(cell, "alignment", None)
        if alignment is not None:
            alignment_payload: dict[str, Any] = {}
            horizontal = str(getattr(alignment, "horizontal", "") or "").strip()
            vertical = str(getattr(alignment, "vertical", "") or "").strip()
            if horizontal:
                alignment_payload["horizontal"] = horizontal
            if vertical:
                alignment_payload["vertical"] = vertical
            if bool(getattr(alignment, "wrap_text", False) or getattr(alignment, "wrapText", False)):
                alignment_payload["wrap_text"] = True
            if bool(getattr(alignment, "shrink_to_fit", False) or getattr(alignment, "shrinkToFit", False)):
                alignment_payload["shrink_to_fit"] = True
            text_rotation = getattr(alignment, "textRotation", None)
            if text_rotation not in (None, 0, 0.0):
                alignment_payload["text_rotation"] = text_rotation
            if horizontal in {"center", "centerContinuous", "distributed", "justify"}:
                alignment_payload["is_centered"] = True
            if alignment_payload:
                style["alignment"] = alignment_payload

        number_format = str(getattr(cell, "number_format", "") or "").strip()
        if number_format and number_format != "General":
            style["number_format"] = number_format

        try:
            row_height = cell.parent.row_dimensions[cell.row].height
            if row_height not in (None, 0, 0.0):
                style["row_height"] = row_height
        except Exception:
            pass
        try:
            col_width = cell.parent.column_dimensions[get_column_letter(cell.column)].width
            if col_width not in (None, 0, 0.0):
                style["col_width"] = col_width
        except Exception:
            pass

        return style

    merged = _merged_anchor_map(ws)

    def _scan_structure_for_worksheet(target_ws) -> dict[str, Any]:
        target_merged = _merged_anchor_map(target_ws)
        min_row, min_col, max_row, max_col = _get_used_bbox(target_ws)
        merged_ranges = []
        for rng in target_ws.merged_cells.ranges:
            merged_ranges.append({"range": str(rng), "anchor": rng.start_cell.coordinate})

        input_candidates = []
        boxed_cells = []
        labels = []
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = target_ws.cell(row=r, column=c)
                a1 = f"{get_column_letter(c)}{r}"
                tags = _border_tags(cell)
                is_merged = (r, c) in target_merged
                style_summary = _cell_style_summary(cell)
                if isinstance(cell.value, str) and cell.value.strip():
                    label_payload = {"a1": a1, "text": cell.value.strip()}
                    if style_summary:
                        label_payload["style"] = style_summary
                    labels.append(label_payload)
                if not tags and not is_merged:
                    continue
                is_boxed = set(tags) >= {"T", "B", "L", "R"}
                has_bottom = "B" in tags
                if is_boxed:
                    candidate_payload = {"a1": a1, "type": "boxed"}
                    if style_summary:
                        candidate_payload["style"] = style_summary
                    boxed_cells.append(candidate_payload)
                    input_candidates.append(dict(candidate_payload))
                elif has_bottom:
                    candidate_payload = {"a1": a1, "type": "bottom_line"}
                    if style_summary:
                        candidate_payload["style"] = style_summary
                    input_candidates.append(candidate_payload)
                elif is_merged:
                    # merged anchors often represent input regions in forms
                    anchor = target_merged[(r, c)]
                    if r == anchor[0] and c == anchor[1]:
                        candidate_payload = {"a1": a1, "type": "merged_anchor"}
                        if style_summary:
                            candidate_payload["style"] = style_summary
                        input_candidates.append(candidate_payload)

        return {
            "bbox": {"min_row": min_row, "min_col": min_col, "max_row": max_row, "max_col": max_col},
            "merged_ranges": merged_ranges,
            "input_candidates": input_candidates,
            "boxed_cells": boxed_cells,
            "labels": labels,
        }

    def read_range(range_str: str, include_style: bool = False):
        cells = ws[range_str]
        if not include_style:
            return [[c.value for c in row] for row in cells]
        out = []
        for row in cells:
            out_row = []
            for c in row:
                out_row.append(
                    {
                        "a1": c.coordinate,
                        "value": c.value,
                        "border": _border_tags(c),
                        "merged": (c.row, c.column) in merged,
                        "style": _cell_style_summary(c),
                    }
                )
            out.append(out_row)
        return out

    def scan_structure():
        return _scan_structure_for_worksheet(ws)

    def _norm_text(s: object) -> str:
        if s is None:
            return ""
        t = str(s).strip().lower()
        t = re.sub(r"\s+", " ", t)
        return t

    def _coord_of(a1: str) -> tuple[int, int] | None:
        try:
            col_s, row = coordinate_from_string(a1)
            return int(row), int(column_index_from_string(col_s))
        except Exception:
            return None

    def _text_sim(a: str, b: str) -> float:
        aa = _norm_text(a)
        bb = _norm_text(b)
        if not aa or not bb:
            return 0.0
        if aa == bb:
            return 1.0
        ratio = SequenceMatcher(None, aa, bb).ratio()
        # Token overlap improves robustness on punctuation/format variants.
        ta = set(aa.replace("/", " ").replace("-", " ").split())
        tb = set(bb.replace("/", " ").replace("-", " ").split())
        overlap = (len(ta & tb) / max(1, len(ta | tb))) if (ta or tb) else 0.0
        return max(ratio, overlap)

    def build_fused_hints(struct: dict[str, object], qwen_pairs: list[dict[str, object]]) -> list[dict[str, object]]:
        labels = struct.get("labels", []) if isinstance(struct, dict) else []
        cands = struct.get("input_candidates", []) if isinstance(struct, dict) else []
        if not isinstance(labels, list) or not isinstance(cands, list):
            return []

        out: list[dict[str, object]] = []
        for p in qwen_pairs:
            key_obj = p.get("key") if isinstance(p, dict) else None
            key_text = key_obj.get("text") if isinstance(key_obj, dict) else None
            if not isinstance(key_text, str) or not key_text.strip():
                continue
            qwen_cell = p.get("value_cell") if isinstance(p.get("value_cell"), str) else ""

            # 1) Best label match by text similarity.
            best_label = None
            best_label_score = 0.0
            for lb in labels:
                if not isinstance(lb, dict):
                    continue
                lb_text = lb.get("text")
                lb_a1 = lb.get("a1")
                if not isinstance(lb_text, str) or not isinstance(lb_a1, str):
                    continue
                s = _text_sim(key_text, lb_text)
                if s > best_label_score:
                    best_label_score = s
                    best_label = lb

            lb_pos = _coord_of(best_label["a1"]) if isinstance(best_label, dict) and isinstance(best_label.get("a1"), str) else None
            q_pos = _coord_of(qwen_cell) if qwen_cell else None

            ranked = []
            for c in cands:
                if not isinstance(c, dict):
                    continue
                a1 = c.get("a1")
                ctype = c.get("type", "")
                if not isinstance(a1, str):
                    continue
                pos = _coord_of(a1)
                if pos is None:
                    continue
                r, col = pos

                # Structural score: right-of-label + proximity + candidate type prior.
                structural_score = 0.0
                if lb_pos is not None:
                    lr, lc = lb_pos
                    dr = abs(r - lr)
                    dc = col - lc
                    right_bonus = 1.0 if dc >= 1 else 0.5
                    prox = 1.0 / (1.0 + float(dr) + 0.35 * float(abs(dc)))
                    structural_score = 0.65 * right_bonus + 0.35 * prox
                type_bonus = 0.0
                if ctype == "boxed":
                    type_bonus = 0.12
                elif ctype == "bottom_line":
                    type_bonus = 0.08
                elif ctype == "merged_anchor":
                    type_bonus = 0.05
                structural_score = min(1.0, structural_score + type_bonus)

                # Qwen score: near qwen value_cell if available.
                qwen_score = 0.0
                if q_pos is not None:
                    qr, qc = q_pos
                    d = abs(r - qr) + abs(col - qc)
                    qwen_score = 1.0 / (1.0 + float(d))

                final_score = 0.70 * structural_score + 0.30 * qwen_score
                ranked.append(
                    {
                        "a1": a1,
                        "type": ctype,
                        "final_score": round(final_score, 4),
                        "structural_score": round(structural_score, 4),
                        "qwen_score": round(qwen_score, 4),
                    }
                )

            ranked.sort(key=lambda x: x["final_score"], reverse=True)
            out.append(
                {
                    "key": key_text,
                    "best_label": best_label.get("a1") if isinstance(best_label, dict) else None,
                    "best_label_text": best_label.get("text") if isinstance(best_label, dict) else None,
                    "label_score": round(best_label_score, 4),
                    "qwen_value_cell": qwen_cell or None,
                    "candidates_top3": ranked[:3],
                }
            )
        return out

    def _append_edit_record(requested_cell: str, target_cell: str, old_value: Any, new_value: Any) -> None:
        record = {
            "sheet": ws.title,
            "requested_cell": str(requested_cell).upper(),
            "target_cell": str(target_cell).upper(),
            "old_value": old_value,
            "new_value": new_value,
            "round_idx": int(round_idx),
            "step_idx": int(current_step_idx),
            "source": "agent_write",
            "is_actual_edit": old_value != new_value,
        }
        write_log.append(record)
        if edit_log_jsonl is not None:
            with edit_log_jsonl.open("a", encoding="utf-8") as f:
                f.write(json.dumps(record, ensure_ascii=False, default=str) + "\n")

    def _append_write_attempt(
        requested_cell: str,
        target_cell: str,
        requested_value: Any,
        old_value: Any,
        status: str,
        message: str,
        is_actual_edit: bool = False,
    ) -> None:
        write_attempt_log.append(
            {
                "sheet": ws.title,
                "requested_cell": str(requested_cell).upper(),
                "target_cell": str(target_cell).upper(),
                "requested_value": requested_value,
                "old_value": old_value,
                "status": status,
                "message": message,
                "round_idx": int(round_idx),
                "step_idx": int(current_step_idx),
                "is_actual_edit": bool(is_actual_edit),
            }
        )

    def write_cell(cell: str, value: str):
        nonlocal dirty
        target = cell
        for merged in ws.merged_cells.ranges:
            if cell in merged:
                target = merged.start_cell.coordinate
                break
        target = str(target).upper()
        try:
            old_value = ws[target].value
        except Exception:
            message = f"WARNING: invalid target {target}"
            _append_write_attempt(
                requested_cell=cell,
                target_cell=target,
                requested_value=value,
                old_value=None,
                status="invalid_target",
                message=message,
                is_actual_edit=False,
            )
            return message
        target_cell_obj = ws[target]
        if enforce_bottom_border_guard and not getattr(getattr(getattr(target_cell_obj, "border", None), "bottom", None), "style", None):
            message = (
                f"WARNING: {target} does not have a bottom border, so it is not an allowed writable field. "
                "Choose a cell with a bottom border."
            )
            _append_write_attempt(
                requested_cell=cell,
                target_cell=target,
                requested_value=value,
                old_value=old_value,
                status="blocked_missing_bottom_border",
                message=message,
                is_actual_edit=False,
            )
            return message
        if enforce_label_guard and target in hard_block_label_cells:
            message = (
                f"WARNING: {target} appears to be a key/label cell, not a value cell. "
                "Do not write field values into label positions."
            )
            _append_write_attempt(
                requested_cell=cell,
                target_cell=target,
                requested_value=value,
                old_value=old_value,
                status="blocked_label_guard",
                message=message,
                is_actual_edit=False,
            )
            return message
        if enforce_protected_cells and normalized_protected_cells:
            if target in normalized_protected_cells and old_value != value:
                reason = normalized_protected_cells[target]
                message = f"WARNING: {target} is protected ({reason}). Re-check before modifying."
                _append_write_attempt(
                    requested_cell=cell,
                    target_cell=target,
                    requested_value=value,
                    old_value=old_value,
                    status="blocked_protected_cell",
                    message=message,
                    is_actual_edit=False,
                )
                return message
        if guarded_write_candidates and target not in guarded_write_candidates:
            message = (
                f"WARNING: {target} is outside the current approved local write-candidate set. "
                "Choose a controller-approved candidate cell instead."
            )
            _append_write_attempt(
                requested_cell=cell,
                target_cell=target,
                requested_value=value,
                old_value=old_value,
                status="blocked_unapproved_candidate",
                message=message,
                is_actual_edit=False,
            )
            return message
        if normalized_allowed_range:
            try:
                if not _range_intersects(target, normalized_allowed_range):
                    message = (
                        f"WARNING: {target} is outside suggested focus range {normalized_allowed_range}. "
                        "Re-check before modifying."
                    )
                    _append_write_attempt(
                        requested_cell=cell,
                        target_cell=target,
                        requested_value=value,
                        old_value=old_value,
                        status="blocked_outside_focus_range",
                        message=message,
                        is_actual_edit=False,
                    )
                    return message
            except Exception:
                message = f"WARNING: invalid target {target}"
                _append_write_attempt(
                    requested_cell=cell,
                    target_cell=target,
                    requested_value=value,
                    old_value=old_value,
                    status="invalid_target",
                    message=message,
                    is_actual_edit=False,
                )
                return message
        ws[target] = value
        _append_edit_record(requested_cell=cell, target_cell=target, old_value=old_value, new_value=value)
        is_actual_edit = old_value != value
        dirty = dirty or is_actual_edit
        message = "OK" if is_actual_edit else f"NOOP: {target} already had the requested value."
        _append_write_attempt(
            requested_cell=cell,
            target_cell=target,
            requested_value=value,
            old_value=old_value,
            status="ok" if is_actual_edit else "noop",
            message=message,
            is_actual_edit=is_actual_edit,
        )
        return message

    def extract_python_code(text: str) -> str:
        s = (text or "").strip()
        if not s:
            return ""
        # Support ReAct-style output:
        # Think: ...
        # Action: Python
        # Action Input:
        # <code>
        m_react = re.search(
            r"Action\s*:\s*Python\s*[\r\n]+Action\s*Input\s*:\s*(.*)$",
            s,
            flags=re.IGNORECASE | re.DOTALL,
        )
        if m_react:
            s = m_react.group(1).strip()
        if "```" in s:
            m = re.search(r"```(?:python)?\s*(.*?)```", s, flags=re.IGNORECASE | re.DOTALL)
            if m:
                s = m.group(1).strip()
        # Also support JSON payload: {"code": "..."}.
        try:
            obj = json.loads(s)
            if isinstance(obj, dict) and isinstance(obj.get("code"), str):
                return obj["code"].strip()
            # JSON actions/non-code payload should not be executed as Python.
            if isinstance(obj, dict) and isinstance(obj.get("actions"), list):
                return ""
            if isinstance(obj, (dict, list)):
                return ""
        except Exception:
            pass
        # Heuristic: likely JSON-ish reply, not executable python.
        if s[:1] in "{[" and s[-1:] in "}]":
            return ""
        # Fallback: strip common meta lines and keep potential code body.
        lines = []
        for ln in s.splitlines():
            if re.match(r"^\s*(Think|Action|Observation|Final)\s*:", ln, flags=re.IGNORECASE):
                continue
            lines.append(ln)
        return "\n".join(lines).strip()

    tools = [
        {
            "type": "function",
            "function": {
                "name": "read_range",
                "description": "Read a rectangular range from Excel",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "range_str": {"type": "string"},
                        "include_style": {"type": "boolean"},
                    },
                    "required": ["range_str"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "scan_structure",
                "description": "Scan workbook structure: merged ranges, border-based input candidates and text labels.",
                "parameters": {
                    "type": "object",
                    "properties": {},
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "write_cell",
                "description": "Write a value to a specific cell",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "cell": {"type": "string"},
                        "value": {"type": "string"},
                    },
                    "required": ["cell", "value"],
                },
            },
        },
    ]

    pairs = form.get("pairs", []) or []
    invalid_slots = form.get("invalid_slots", []) or []
    guarded_write_candidates_raw = form.get("guarded_write_candidates", []) or []
    positive_reflect_hints = [
        hint
        for hint in (reflect_hints or [])
        if isinstance(hint, dict) and not bool(hint.get("invalid_slot")) and not bool(hint.get("avoid_cell"))
    ]
    negative_reflect_hints = [
        hint
        for hint in (reflect_hints or [])
        if isinstance(hint, dict) and (bool(hint.get("invalid_slot")) or bool(hint.get("avoid_cell")))
    ]
    mapping_lines = []
    for p in pairs:
        key = p.get("key")
        value_cell = p.get("value_cell")
        if key and value_cell:
            mapping_lines.append(f"- {key} -> {value_cell}")

    json_protocol = (
        "You must respond with ONLY valid JSON. No markdown, no extra text.\n"
        "Format:\n"
        "{\n"
        '  "actions": [\n'
        '    {"tool": "read_range", "args": {"range_str": "A1:C10"}},\n'
        '    {"tool": "write_cell", "args": {"cell": "B2", "value": "Alice"}}\n'
        "  ],\n"
        '  "final": "optional short confirmation"\n'
        "}\n"
        "If you are done, return {\"actions\": [], \"final\": \"done\"}.\n"
    )
    code_protocol = (
        "You may write Python code to fill the sheet.\n"
        "Return either:\n"
        "1) a Python code block, or\n"
        "2) JSON {\"code\": \"...\", \"final\": \"optional\"}.\n"
        "Available runtime objects/functions:\n"
        "- scan_structure() -> dict\n"
        "- read_range(range_str: str, include_style: bool=False) -> list[list] | styled grid\n"
        "- write_cell(cell: str, value: str) -> 'OK' | 'WARNING: ...' | 'NOOP: ...'\n"
        "Rules:\n"
        "- Do NOT use file/network/process operations.\n"
        "- Use scan_structure() and read_range(include_style=True) whenever local workbook verification is needed before writing.\n"
        "- scan_structure() and read_range(include_style=True) expose formatting signals such as fill/background color, font, alignment/centering, borders, and number format. Use them to distinguish labels from writable value areas.\n"
        "- You MUST use write_cell() for every edit. Do not assign workbook/sheet cells directly.\n"
        "- The raw workbook/worksheet objects are intentionally not exposed; all writes must go through write_cell().\n"
        "- Prefer read_range/write_cell for deterministic writes.\n"
        "- If a step returns code_result=warning/partial or write_attempts with WARNING/NOOP, use that feedback to choose different cells before returning done.\n"
        "- After successful writes, stop immediately and return final done. Do not add narrative text after code.\n"
        "- If done, return JSON {\"code\": \"\", \"final\": \"done\"}.\n"
    )
    if enforce_bottom_border_guard:
        code_protocol += "- Only write into cells that have a bottom border. If write_cell warns that a cell has no bottom border, choose another target.\n"

    system_prompt = (
        "You are a deterministic Excel form filling agent.\n"
        "You can only inspect/write workbook state via tools.\n"
        "Workflow MUST follow:\n"
        "1) Inspect the provided form context and determine candidate input cells.\n"
        "2) If uncertain, call scan_structure() and/or read_range() on a small nearby area with include_style=true.\n"
        "3) Then write values using write_cell().\n"
        "All edits MUST go through write_cell() so they can be audited.\n"
        "Use formatting evidence such as fill/background color, font emphasis, alignment/centering, borders, and merged layout to separate static labels from writable value areas.\n"
        "A blank merged cell/range immediately below or beside a filled/dark/static label is usually the writable value area for that label; do not write into the label cell when such a blank merged value area exists.\n"
        "When the label is a topic/section title and the source contains a clearly matching section, you may write a concise evidence-grounded summary of that source section into the corresponding blank value area.\n"
    )
    if enforce_bottom_border_guard:
        system_prompt += "Only write into cells that have a bottom border.\n"
    if enforce_label_guard:
        system_prompt += (
            "Never write a field value into a static template key/label cell. "
            "If a template shows field text like 'Time', 'Date', '(First)', or similar label text, preserve that cell and write the actual field value into the neighboring input/value cell instead. "
            "If write_cell warns that a target looks like a label cell, choose a different value cell.\n"
        )
    if agent_profile == "free_fill":
        system_prompt += (
            "This is a conservative free first-pass fill. Work directly from the instruction and workbook structure, "
            "but only write when the target cell is strongly supported by local evidence.\n"
            "If there is meaningful ambiguity about which nearby cell is the writable value area, inspect a small local region first.\n"
            "If local structure shows a dark/filled label cell with an adjacent or following blank merged range, prefer the blank merged range as the writable value area after verification.\n"
            "For forms organized by topic headings, matching a source section title to the heading is sufficient evidence to write a short summary into that heading's blank value area; preserve unsupported headings as blank.\n"
            "If uncertainty remains after local inspection, leave that field blank for now rather than forcing a write.\n"
            "Prefer missing a low-confidence field over writing into the wrong cell.\n"
        )
    else:
        system_prompt += (
            "This is a guided repair pass. Use the provided focus region, hints, and protection rules as the main repair frame.\n"
        )
    system_prompt += (
        "If negative detector/reflect hints mark a cell or slot as invalid, treat that location as an avoid-set. "
        "Do not write the value back into that cell unless strong local workbook evidence clearly overrides the hint.\n"
        "Never guess coordinates without structural confirmation.\n"
        "After you finish writing, stop and return done. Do not continue with explanatory prose.\n"
    )
    if planner_policy:
        system_prompt += (
            "\nIf planner controller information is provided, use it to sequence your work.\n"
            "Treat it as workflow guidance, not as a hard content constraint.\n"
        )
    if skills_text.strip():
        system_prompt += (
            "\nAgent skills to follow (higher priority than default style, but never violate tool constraints):\n"
            f"{skills_text.strip()}\n"
        )
    if tool_mode == "json":
        system_prompt += json_protocol
    elif tool_mode == "code":
        system_prompt += code_protocol

    struct_summary = scan_structure()
    role_struct_summary = _scan_structure_for_worksheet(role_ws)
    planner_guard_cells, planner_guard_matches = _resolve_planner_guard_cells(role_struct_summary, planner_policy)
    if enforce_protected_cells:
        for cell, reason in planner_guard_cells.items():
            normalized_protected_cells.setdefault(str(cell).upper(), str(reason))
    fused_hints = build_fused_hints(struct_summary, pairs)
    current_label_cells = {
        str(item.get("a1") or "").upper()
        for item in (struct_summary.get("labels") or [])
        if isinstance(item, dict) and str(item.get("a1") or "").strip()
    }
    role_label_cells = {
        str(item.get("a1") or "").upper()
        for item in (role_struct_summary.get("labels") or [])
        if isinstance(item, dict) and str(item.get("a1") or "").strip()
    }
    role_input_candidate_cells = {
        str(item.get("a1") or "").upper()
        for item in (role_struct_summary.get("input_candidates") or [])
        if isinstance(item, dict) and str(item.get("a1") or "").strip()
    }
    explicit_value_cells: set[str] = set()
    for pair in pairs:
        if not isinstance(pair, dict):
            continue
        value_cell = str(pair.get("value_cell") or "").strip()
        if not value_cell:
            continue
        for cell_ref in _iter_range_cells(value_cell):
            explicit_value_cells.add(str(cell_ref).upper())
    for hint in positive_reflect_hints:
        if not isinstance(hint, dict):
            continue
        hint_cell = str(hint.get("candidate_value_cell") or "").strip()
        if hint_cell:
            for cell_ref in _iter_range_cells(hint_cell):
                explicit_value_cells.add(str(cell_ref).upper())
    guarded_write_candidates: set[str] = set()
    for item in guarded_write_candidates_raw:
        text = str(item or "").strip()
        if not text:
            continue
        try:
            expanded_cells = _iter_range_cells(text)
        except Exception:
            expanded_cells = [text]
        for cell_ref in expanded_cells:
            normalized = _resolve_write_target(str(cell_ref))
            if normalized:
                guarded_write_candidates.add(normalized)
    template_static_label_cells = role_label_cells - role_input_candidate_cells
    if not template_static_label_cells:
        template_static_label_cells = current_label_cells - role_input_candidate_cells
    hard_block_label_cells = (
        template_static_label_cells - explicit_value_cells - guarded_write_candidates
        if enforce_label_guard
        else set()
    )

    user_prompt = (
        "Fill the form based on this instruction.\n"
    )
    planner_block = _planner_policy_prompt_block(planner_policy, planner_stage)
    if planner_block:
        user_prompt += planner_block + "\n"
    user_prompt += f"Instruction:\n{instruction}\n\nFields to fill should be extracted from the instruction text.\n"
    if agent_profile == "free_fill":
        user_prompt += (
            "\nConservative first-pass policy:\n"
            "- Only fill fields whose target cell is high-confidence from local layout evidence.\n"
            "- If a nearby target might be a static label/header or you are unsure between neighboring cells, do not write yet.\n"
            "- Empty merged cells/ranges directly below or beside filled/dark label cells are strong writable-value candidates; verify locally and write there instead of the label cell.\n"
            "- For topic/section labels, if the source contains a clearly matching section, summarize that source section concisely into the corresponding blank merged value range.\n"
            "- Read a small local range before writing when the layout is even slightly ambiguous.\n"
            "- It is acceptable to leave uncertain fields blank in first pass.\n"
        )
    mode = str(form_context_mode or "struct").strip().lower()
    if mode not in {"struct", "html", "struct_html"}:
        mode = "struct"
    if mode in {"struct", "struct_html"}:
        user_prompt += (
            "\nPre-scanned structure summary (authoritative):\n"
            f"{json.dumps(struct_summary, ensure_ascii=False)}\n"
        )
    if rendered_html_path is not None:
        try:
            rendered_html_text = Path(rendered_html_path).read_text(encoding="utf-8")
            max_chars = int(rendered_html_max_chars or 0)
            if max_chars > 0:
                rendered_html_text = rendered_html_text[:max_chars]
            html_header = (
                "\nRendered sheet HTML (primary form context):\n"
                if mode == "html"
                else "\nRendered sheet HTML (supplementary, exact layout/cell-coordinate context):\n"
            )
            user_prompt += (
                html_header
                + f"{rendered_html_text}\n"
                + (
                    "Treat the HTML as the main table-understanding context for this fill step. "
                    "Use workbook tools to verify uncertain local areas before writing.\n"
                    if mode == "html"
                    else "The HTML includes rendered cell layout and may contain direct coordinate markers such as data-a1 attributes. "
                    "Use it as additional evidence for local field-to-cell mapping, while keeping workbook tool observations authoritative.\n"
                )
            )
        except Exception as exc:
            user_prompt += f"\n[WARN] rendered_html_unavailable: {exc}\n"
    if mapping_mode == "hint":
        user_prompt += (
            "\nOptional low-confidence detector hints (key -> candidate value_cell, may be noisy):\n"
            + ("\n".join(mapping_lines) if mapping_lines else "(none)")
        )
        if fused_hints:
            user_prompt += (
                "\n\nFused structure+qwen candidate ranking (use this as soft guidance; structure is primary):\n"
                f"{json.dumps(fused_hints, ensure_ascii=False)}"
            )
        if positive_reflect_hints:
            user_prompt += (
                "\n\nReflect correction hints (high priority within the allowed patch region):\n"
                f"{json.dumps(positive_reflect_hints, ensure_ascii=False)}\n"
                "Use these hints first when revising cells. Do not modify cells outside the allowed patch region."
            )
    negative_detector_hints = [item for item in invalid_slots if isinstance(item, dict)]
    if negative_detector_hints:
        user_prompt += (
            "\n\nNegative detector hints (cells/slots judged invalid by Qwen; treat as avoid-set):\n"
            f"{json.dumps(negative_detector_hints[:20], ensure_ascii=False)}\n"
            "Do not write values into these candidate cells unless strong local workbook evidence clearly overrides the detector."
        )
    if negative_reflect_hints:
        user_prompt += (
            "\n\nNegative reflect hints (high priority avoid-set inside the current review/patch context):\n"
            f"{json.dumps(negative_reflect_hints[:20], ensure_ascii=False)}\n"
            "These cells/slots were judged invalid or unsafe. Avoid reusing them unless direct local evidence strongly supports them."
        )
    if normalized_allowed_range:
        user_prompt += (
            f"\n\nCurrent focus region: {normalized_allowed_range}.\n"
            "Prefer edits inside this focus region. If you decide to edit outside it, re-check structure first."
        )
    if enforce_protected_cells and normalized_protected_cells:
        protected_examples = list(normalized_protected_cells.items())[:8]
        user_prompt += (
            f"\n\nProtected cells: {len(normalized_protected_cells)} total.\n"
            "These cells either contain original template content or were previously assessed as correct. "
            "If write_cell returns a WARNING for one of them, reconsider before modifying.\n"
            f"Examples: {json.dumps(protected_examples, ensure_ascii=False)}"
        )
    if enforce_label_guard and hard_block_label_cells:
        label_examples = list(sorted(hard_block_label_cells))[:10]
        user_prompt += (
            f"\n\nDetected static template label cells that should not receive field values: {len(hard_block_label_cells)} total.\n"
            "These are derived from the original template role, not just from whichever cells currently contain text.\n"
            "If write_cell warns that a target is a label cell, keep that label/static text in place and choose a nearby value/input cell instead.\n"
            f"Examples: {json.dumps(label_examples, ensure_ascii=False)}"
        )
    if guarded_write_candidates:
        guarded_examples = list(sorted(guarded_write_candidates))[:20]
        user_prompt += (
            f"\n\nApproved local write-candidate cells for this repair pass: {len(guarded_write_candidates)} total.\n"
            "Treat this as a hard local whitelist during guided repair. "
            "If write_cell warns that a target is outside this set, pick another candidate from this whitelist.\n"
            f"Examples: {json.dumps(guarded_examples, ensure_ascii=False)}"
        )
    if planner_guard_matches:
        user_prompt += (
            "\n\nPlanner-resolved blank/uncertain field guards:\n"
            f"{json.dumps(planner_guard_matches[:10], ensure_ascii=False)}\n"
            "Treat these matched candidate cells as do-not-fill targets unless the instruction provides stronger direct evidence."
        )

    user_content: Any = user_prompt
    if context_summary or compare_before_after_pngs:
        blocks: list[dict[str, Any]] = [{"type": "text", "text": user_prompt}]
        if context_summary:
            blocks.append(
                {
                    "type": "text",
                    "text": "Prior fill context summary:\n" + json.dumps(context_summary, ensure_ascii=False),
                }
            )
        if compare_before_after_pngs:
            before_path = compare_before_after_pngs.get("before")
            after_path = compare_before_after_pngs.get("after")
            screenshot_summary = compare_before_after_pngs.get("summary", "")
            if screenshot_summary:
                blocks.append({"type": "text", "text": f"Before/after screenshot summary:\n{screenshot_summary}"})
            if before_path:
                b64_before, mime_before = _encode_image_base64(Path(before_path))
                blocks.append({"type": "text", "text": "BEFORE screenshot (template / pre-fill):"})
                blocks.append({"type": "image_url", "image_url": {"url": f"data:image/{mime_before};base64,{b64_before}"}})
            if after_path:
                b64_after, mime_after = _encode_image_base64(Path(after_path))
                blocks.append({"type": "text", "text": "AFTER screenshot (current filled result):"})
                blocks.append({"type": "image_url", "image_url": {"url": f"data:image/{mime_after};base64,{b64_after}"}})
        user_content = blocks

    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_content},
    ]

    client = build_client(base_url, api_key)

    output_dir = output_xlsx.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    extra_retry_after_block_budget = 1
    step_limit = int(max_steps)

    def _maybe_extend_after_block(step_attempt_start: int) -> None:
        nonlocal step_limit, extra_retry_after_block_budget
        if extra_retry_after_block_budget <= 0:
            return
        if current_step_idx < step_limit:
            return
        step_write_attempts = write_attempt_log[step_attempt_start:]
        if not step_write_attempts:
            return
        actual_step_edits = sum(1 for item in step_write_attempts if bool(item.get("is_actual_edit")))
        blocked_step_attempts = [
            item
            for item in step_write_attempts
            if str(item.get("status") or "").startswith("blocked_") or str(item.get("status") or "") == "invalid_target"
        ]
        if actual_step_edits != 0 or not blocked_step_attempts:
            return
        step_limit += 1
        extra_retry_after_block_budget -= 1
        messages.append(
            {
                "role": "user",
                "content": (
                    "Your last write attempt was rejected by workbook guards and produced no actual edits. "
                    "Try one more time: inspect nearby structure again, then choose a different writable cell that satisfies the guard constraints."
                ),
            }
        )

    def _dump_agent_step_io(step_idx: int, mode_label: str, resp_obj: Any, parsed_resp: dict[str, Any], messages_snapshot: Any) -> None:
        raw_text = str(parsed_resp.get("content") or "")
        _dump_model_io(
            model_io_dir,
            f"step_{int(step_idx):03d}_{mode_label}",
            request_payload={
                "model": model,
                "temperature": float(temperature),
                "tool_mode": tool_mode,
                "mapping_mode": mapping_mode,
                "planner_stage": planner_stage,
                "messages": messages_snapshot,
            },
            response_payload={
                "content": raw_text,
                "tool_calls": parsed_resp.get("tool_calls", []),
            },
            raw_text=raw_text,
            usage=_usage_to_dict(resp_obj),
            extra={
                "allowed_write_range": normalized_allowed_range,
                "protected_write_cells_count": len(normalized_protected_cells),
            },
        )

    step_idx = 0
    while step_idx < step_limit:
        step_idx += 1
        current_step_idx = step_idx
        if tool_mode == "native":
            step_attempt_start = len(write_attempt_log)
            request_messages = list(messages)
            resp = client.chat.completions.create(
                model=model,
                messages=request_messages,
                temperature=float(temperature),
                tools=tools,
                tool_choice="auto",
            )
            usage = _usage_to_dict(resp)
            if usage:
                usage_totals["prompt_tokens"] += int(usage.get("prompt_tokens", 0))
                usage_totals["completion_tokens"] += int(usage.get("completion_tokens", 0))
                usage_totals["total_tokens"] += int(usage.get("total_tokens", 0))
                usage_totals["calls_with_usage"] += 1
            parsed_resp = _parse_chat_response(resp)
            _dump_agent_step_io(step_idx, "native", resp, parsed_resp, request_messages)
            msg = parsed_resp["assistant_message"]
            messages.append(msg)
            tool_calls = parsed_resp["tool_calls"]
            if not tool_calls:
                break
            for call_id, name, args_raw in _iter_tool_calls(tool_calls):
                args = json.loads(args_raw)
                if name == "read_range":
                    result = read_range(**args)
                elif name == "scan_structure":
                    result = scan_structure()
                elif name == "write_cell":
                    result = write_cell(**args)
                else:
                    result = f"Unknown tool: {name}"
                messages.append(
                    {
                        "role": "tool",
                        "tool_call_id": call_id,
                        "content": json.dumps(result),
                    }
                )
            _maybe_extend_after_block(step_attempt_start)
        elif tool_mode == "json":
            step_attempt_start = len(write_attempt_log)
            request_messages = list(messages)
            resp = client.chat.completions.create(model=model, messages=request_messages, temperature=float(temperature))
            usage = _usage_to_dict(resp)
            if usage:
                usage_totals["prompt_tokens"] += int(usage.get("prompt_tokens", 0))
                usage_totals["completion_tokens"] += int(usage.get("completion_tokens", 0))
                usage_totals["total_tokens"] += int(usage.get("total_tokens", 0))
                usage_totals["calls_with_usage"] += 1
            parsed_resp = _parse_chat_response(resp)
            _dump_agent_step_io(step_idx, "json", resp, parsed_resp, request_messages)
            msg = parsed_resp["assistant_message"]
            messages.append(msg)
            try:
                payload = json.loads(parsed_resp["content"] or "{}")
            except json.JSONDecodeError:
                messages.append(
                    {
                        "role": "user",
                        "content": "Invalid JSON. Respond ONLY with valid JSON per the required format.",
                    }
                )
                continue
            actions = payload.get("actions", [])
            if not actions:
                break
            results = []
            for action in actions:
                name = action.get("tool")
                args = action.get("args", {})
                if name == "read_range":
                    results.append({"tool": name, "result": read_range(**args)})
                elif name == "scan_structure":
                    results.append({"tool": name, "result": scan_structure()})
                elif name == "write_cell":
                    results.append({"tool": name, "result": write_cell(**args)})
                else:
                    results.append({"tool": name, "result": f"Unknown tool: {name}"})
            messages.append({"role": "user", "content": json.dumps({"results": results})})
            _maybe_extend_after_block(step_attempt_start)
        else:
            request_messages = list(messages)
            resp = client.chat.completions.create(model=model, messages=request_messages, temperature=float(temperature))
            usage = _usage_to_dict(resp)
            if usage:
                usage_totals["prompt_tokens"] += int(usage.get("prompt_tokens", 0))
                usage_totals["completion_tokens"] += int(usage.get("completion_tokens", 0))
                usage_totals["total_tokens"] += int(usage.get("total_tokens", 0))
                usage_totals["calls_with_usage"] += 1
            parsed_resp = _parse_chat_response(resp)
            _dump_agent_step_io(step_idx, "code", resp, parsed_resp, request_messages)
            msg = parsed_resp["assistant_message"]
            messages.append(msg)
            raw_content = parsed_resp["content"] or ""
            code = extract_python_code(raw_content)
            if not code:
                # Compatibility fallback: some models return JSON actions even in code mode.
                parsed = None
                try:
                    parsed = json.loads(raw_content)
                except Exception:
                    parsed = None
                if isinstance(parsed, dict):
                    final_text = str(parsed.get("final") or "").strip().lower()
                    parsed_code = str(parsed.get("code") or "").strip()
                    parsed_actions = parsed.get("actions")
                    if final_text in {"done", "no_op", "noop"} and not parsed_code and not parsed_actions:
                        break
                if isinstance(parsed, dict) and isinstance(parsed.get("actions"), list):
                    results = []
                    for action in parsed.get("actions", []):
                        name = action.get("tool")
                        args = action.get("args", {})
                        if name == "read_range":
                            results.append({"tool": name, "result": read_range(**args)})
                        elif name == "scan_structure":
                            results.append({"tool": name, "result": scan_structure()})
                        elif name == "write_cell":
                            results.append({"tool": name, "result": write_cell(**args)})
                        else:
                            results.append({"tool": name, "result": f"Unknown tool: {name}"})
                    messages.append({"role": "user", "content": json.dumps({"results": results})})
                    continue
                # Ask model to return executable code instead of terminating early.
                messages.append(
                    {
                        "role": "user",
                        "content": "No executable Python code detected. Return ONLY Python code (or JSON {\"code\": \"...\"}) and call write_cell(...) for every edit.",
                    }
                )
                continue

            safe_builtins = {
                "len": len,
                "range": range,
                "min": min,
                "max": max,
                "sum": sum,
                "str": str,
                "int": int,
                "float": float,
                "bool": bool,
                "list": list,
                "dict": dict,
                "set": set,
                "tuple": tuple,
                "enumerate": enumerate,
                "zip": zip,
                "sorted": sorted,
                "print": print,
            }
            env = {
                "__builtins__": safe_builtins,
                "scan_structure": scan_structure,
                "read_range": read_range,
                "write_cell": write_cell,
                "openpyxl": openpyxl,
            }
            step_attempt_start = len(write_attempt_log)
            try:
                exec(code, env, {})
                code_executed = True
                step_write_attempts = write_attempt_log[step_attempt_start:]
                actual_step_edits = sum(1 for item in step_write_attempts if bool(item.get("is_actual_edit")))
                blocked_step_attempts = [
                    item
                    for item in step_write_attempts
                    if str(item.get("status") or "").startswith("blocked_") or str(item.get("status") or "") == "invalid_target"
                ]
                noop_step_attempts = [item for item in step_write_attempts if str(item.get("status") or "") == "noop"]
                if step_write_attempts and actual_step_edits == 0 and blocked_step_attempts:
                    code_result = "warning"
                    summary_message = "All write attempts were blocked. Choose different writable cells."
                elif step_write_attempts and actual_step_edits == 0 and noop_step_attempts:
                    code_result = "warning"
                    summary_message = "Code ran, but it produced no actual edits. Re-check the target cells."
                elif blocked_step_attempts:
                    code_result = "partial"
                    summary_message = "Some writes succeeded, but other write attempts were blocked."
                else:
                    code_result = "ok"
                    summary_message = "Code executed."
                messages.append(
                    {
                        "role": "user",
                        "content": json.dumps(
                            {
                                "code_result": code_result,
                                "message": summary_message,
                                "actual_edits": actual_step_edits,
                                "write_attempts": step_write_attempts[-10:],
                            },
                            ensure_ascii=False,
                            default=str,
                        ),
                    }
                )
                _maybe_extend_after_block(step_attempt_start)
            except Exception as exc:
                step_write_attempts = write_attempt_log[step_attempt_start:]
                messages.append(
                    {
                        "role": "user",
                        "content": json.dumps(
                            {
                                "code_result": "error",
                                "error": str(exc),
                                "write_attempts": step_write_attempts[-10:],
                            },
                            ensure_ascii=False,
                            default=str,
                        ),
                    }
                )
                if step_write_attempts:
                    messages.append(
                        {
                            "role": "user",
                            "content": "Some write attempts may have been blocked before the error. Inspect the returned write_attempts and choose different writable cells if needed.",
                        }
                    )
                code_exec_errors.append(str(exc))
                _maybe_extend_after_block(step_attempt_start)

    no_code_generated = False
    no_code_generated_reason = ""
    if (tool_mode == "code") and (not dirty) and (not code_executed):
        no_code_generated = True
        no_code_generated_reason = "; ".join(code_exec_errors[-3:]) if code_exec_errors else "no write_cell call produced"

    if dirty or (not output_xlsx.exists()):
        wb.save(output_xlsx)
    if history_dump_path is not None:
        history_dump_path.parent.mkdir(parents=True, exist_ok=True)
        history_dump_path.write_text(json.dumps(messages, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    if role_wb is not None:
        role_wb.close()
    wb.close()
    return {
        "write_log": write_log,
        "write_attempt_log": write_attempt_log,
        "dirty": dirty,
        "sheet": ws.title,
        "struct_summary": struct_summary,
        "role_struct_summary": role_struct_summary,
        "planner_guard_cells": normalized_protected_cells,
        "planner_guard_matches": planner_guard_matches,
        "hard_block_label_cells": sorted(hard_block_label_cells),
        "guarded_write_candidates": sorted(guarded_write_candidates),
        "no_code_generated": no_code_generated,
        "no_code_generated_reason": no_code_generated_reason,
        "code_exec_errors": code_exec_errors,
        "history_dump_path": str(history_dump_path) if history_dump_path is not None else "",
        "usage": usage_totals,
    }


def reflect_check_with_screenshot(
    before_xlsx_path: Path,
    after_xlsx_path: Path,
    sheet_name: str | None,
    instruction: str,
    model: str,
    base_url: str | None,
    api_key: str | None,
    out_dir: Path,
    dpi: int = 96,
    checker_context_summary: dict[str, Any] | None = None,
    planner_policy: dict[str, Any] | None = None,
    model_io_dir: Path | None = None,
) -> dict[str, Any]:
    def _refine_bad_range(reason_text: str, initial_range: str) -> tuple[str, str, dict[str, int]]:
        refine_prompt = (
            "You are refining a previously detected Excel form issue location.\n"
            "The issue itself is already known. Your only task is to localize the smallest precise Excel A1 range for that issue.\n"
            "Return JSON only with schema:\n"
            "{\"bad_ranges\": [<A1 range>, ...], \"bad_range\": <first A1 range string or empty>, \"reason\": <short string>}.\n"
            "Rules:\n"
            "- bad_ranges should usually contain exactly one item for this refinement step.\n"
            "- bad_range must be the same as the first item of bad_ranges, or empty if no valid range exists.\n"
            "- bad_ranges must be the smallest precise A1 ranges for the described issue.\n"
            "- Prefer a single cell when the error is only one cell.\n"
            "- Ignore other possible issues.\n"
            "- Never return field names or prose in bad_ranges or bad_range.\n\n"
            f"Instruction:\n{instruction}\n\n"
            f"Known issue:\n{reason_text}\n\n"
            f"Initial guessed range:\n{initial_range or '(none)'}\n"
        )
        planner_block = _planner_policy_prompt_block(planner_policy, "reflect")
        if planner_block:
            refine_prompt += f"\n{planner_block}\n"
        if checker_context_summary:
            refine_prompt += (
                "\nPrior workbook-read context summary:\n"
                f"{json.dumps(checker_context_summary, ensure_ascii=False)}\n"
                "Use this auxiliary structure/read context when localizing bad_range.\n"
            )
        refine_messages = [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": refine_prompt},
                    {"type": "text", "text": "BEFORE screenshot (template / pre-fill):"},
                    {"type": "image_url", "image_url": {"url": f"data:image/{mime_before};base64,{b64_before}"}},
                    {"type": "text", "text": "AFTER screenshot (current filled result):"},
                    {"type": "image_url", "image_url": {"url": f"data:image/{mime_after};base64,{b64_after}"}},
                ],
            }
        ]
        refine_resp = client.chat.completions.create(
            model=model,
            temperature=0.0,
            messages=refine_messages,
        )
        refine_parsed = _parse_chat_response(refine_resp)
        refine_raw = str(refine_parsed["content"] or "").strip()
        refine_usage = _usage_to_dict(refine_resp)
        _dump_model_io(
            model_io_dir,
            "stage1_refine_bad_range",
            request_payload={"model": model, "temperature": 0.0, "messages": refine_messages},
            response_payload={"content": refine_raw},
            raw_text=refine_raw,
            usage=refine_usage,
            extra={"initial_range": initial_range, "reason_text": reason_text},
        )
        refine_payload = _extract_json_object_from_text(refine_raw)
        refined_list = _extract_focus_ranges_from_text(
            json.dumps(refine_payload, ensure_ascii=False) if refine_payload else refine_raw
        )
        refined = refined_list[0] if refined_list else ""
        return refined, refine_raw, refine_usage

    prep_before = preprocess_excel_html_css(
        xlsx_path=before_xlsx_path,
        out_dir=out_dir / "before",
        sheet_name=sheet_name,
        dpi=dpi,
    )
    prep_after = preprocess_excel_html_css(
        xlsx_path=after_xlsx_path,
        out_dir=out_dir / "after",
        sheet_name=sheet_name,
        dpi=dpi,
    )
    before_png = prep_before["sheet_png"]
    after_png = prep_after["sheet_png"]
    b64_before, mime_before = _encode_image_base64(before_png)
    b64_after, mime_after = _encode_image_base64(after_png)
    prompt = (
        "You are checking an Excel form filling result with BEFORE/AFTER screenshots.\n"
        "Compare AFTER against instruction, while using BEFORE as baseline template context.\n"
        "Return JSON only with schema:\n"
        "{\"need_fix\": <bool>, \"bad_ranges\": [<A1 range>, ...], \"bad_range\": <first A1 range string or empty>, \"correct_ranges\": [<A1 range>, ...], \"reason\": <short string>, \"confidence\": <0..1>}.\n"
        "Set need_fix=true only when the AFTER sheet shows an obvious wrong or missing filled region.\n"
        "According to the instruction, if a form field is expected but either its key/label or its corresponding value is missing, you may mark that area as bad_ranges.\n"
        "Also return correct_ranges for regions that are clearly already correct and should be protected from later modification, including fields that are correctly left blank.\n"
        "Because this file is a form, treat a region as correct only when the form shows both the correct key/label and the correct corresponding value together; do not mark a region correct when only the key is right or only the value is right.\n"
        "If a filled value appears inside a title cell, header cell, key/label cell, or any other static template cell from BEFORE, that is wrong even if the text itself is plausible.\n"
        "If a combined value is written into one sibling field while the neighboring sibling value field remains blank or still shows template label text, that is wrong; mark both the wrong source area and the missing target area as bad.\n"
        "Do not call a region correct just because the right words appear somewhere nearby; the words must appear in the correct role and cell type.\n"
        "Be conservative with correct_ranges: only include regions that are clearly correct from screenshots and prior workbook-read context.\n"
        "Do not mark a region as bad_ranges for merely suspicious structural/pattern risks; those are handled in a later risk-scoring stage.\n"
        "bad_ranges must contain only valid Excel A1 ranges, such as \"A4:B7\" or \"C12\".\n"
        "bad_range must be empty or equal to the first item in bad_ranges.\n"
        "When the issue is a single missing/wrong filled cell, return that exact single cell as the only bad_ranges item.\n"
        "Prefer the smallest precise rectangles that cover the actual wrong/missing content, not larger surrounding blocks.\n"
        "correct_ranges must contain only valid Excel A1 ranges.\n"
        "Never put field names, labels, prose, or descriptions into bad_ranges or bad_range.\n"
        "If you cannot localize the issue to concrete A1 ranges, set need_fix=false and bad_ranges=[].\n"
        "If mistakes are local, each bad_ranges item should be a compact rectangle such as \"A1:F8\".\n"
        "If no fix needed, bad_ranges must be empty and bad_range must be empty string.\n\n"
        "Few-shot guidance:\n"
        "Example 1: BEFORE shows a static label cell for a vendor/account name, but AFTER places an account number into that same label/static cell while the real value cell stays blank. This is wrong. need_fix=true, and bad_ranges should include the overwritten static/label cell and the missing value cell.\n"
        "Example 2: BEFORE has sibling fields for last name and first name. AFTER writes a combined string like 'Doe, Jane' into the last-name field, while the first-name field is still blank or still shows template label text such as '(First)'. This is wrong. need_fix=true, and bad_ranges should include both the wrongly used source field and the missing first-name field.\n"
        "Example 3: If AFTER still shows the template label text in a supposed value field, that field is not correctly filled. Template label text visible where a value should be counts as missing/wrong fill, not as correct.\n\n"
        "Example 4: BEFORE shows a static label cell like 'CONTACT NAME:'. AFTER changes that same cell to 'CONTACT NAME: Alex Rivera' while the adjacent value area is still empty. This is wrong. need_fix=true, and bad_ranges should include the overwritten label cell and the missing neighboring value cell.\n\n"
        f"Instruction:\n{instruction}\n"
    )
    planner_block = _planner_policy_prompt_block(planner_policy, "reflect")
    if planner_block:
        prompt += f"\n{planner_block}\n"
    if checker_context_summary:
        prompt += (
            "\nPrior workbook-read context summary:\n"
            f"{json.dumps(checker_context_summary, ensure_ascii=False)}\n"
            "When deciding bad_range, also use this auxiliary read/structure context from earlier workbook inspection.\n"
        )
    client = build_client(base_url, api_key)
    checker_messages = [
        {
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {"type": "text", "text": "BEFORE screenshot (template / pre-fill):"},
                {"type": "image_url", "image_url": {"url": f"data:image/{mime_before};base64,{b64_before}"}},
                {"type": "text", "text": "AFTER screenshot (current filled result):"},
                {"type": "image_url", "image_url": {"url": f"data:image/{mime_after};base64,{b64_after}"}},
            ],
        }
    ]
    resp = client.chat.completions.create(
        model=model,
        temperature=0.0,
        messages=checker_messages,
    )
    parsed_resp = _parse_chat_response(resp)
    usage = _usage_to_dict(resp)
    raw = str(parsed_resp["content"] or "").strip()
    _dump_model_io(
        model_io_dir,
        "stage1_checker",
        request_payload={"model": model, "temperature": 0.0, "messages": checker_messages},
        response_payload={"content": raw},
        raw_text=raw,
        usage=usage,
        extra={"before_sheet_png": str(before_png), "after_sheet_png": str(after_png)},
    )
    payload = _extract_json_object_from_text(raw)
    bad_ranges_raw = _extract_bad_ranges_raw_from_payload(payload) if payload else []
    correct_ranges = _extract_focus_ranges_from_payload(payload) if payload else []
    bad_ranges: list[str] = []
    if bad_ranges_raw:
        for raw_item in bad_ranges_raw:
            bad_ranges.extend(_extract_focus_ranges_from_text(json.dumps({"bad_range": raw_item}, ensure_ascii=False)))
        bad_ranges = _dedup_normalized_ranges(bad_ranges)
    elif not payload:
        bad_ranges = _extract_focus_ranges_from_text(raw)
    need_fix = bool(payload.get("need_fix")) if payload else bool(bad_ranges)
    bad_range_parse_failed = bool(need_fix) and bool(bad_ranges_raw) and not bool(bad_ranges)
    reason = str(payload.get("reason", "")) if payload else ""
    conf_raw = payload.get("confidence", 0.0) if payload else 0.0
    try:
        conf = float(conf_raw)
    except Exception:
        conf = 0.0
    heuristic_findings = _detect_template_label_value_overwrite(before_xlsx_path, after_xlsx_path, sheet_name)
    heuristic_bad_ranges = _dedup_normalized_ranges(
        [rng for item in heuristic_findings for rng in (item.get("bad_ranges") or []) if str(rng or "").strip()]
    )
    if heuristic_bad_ranges:
        need_fix = True
        bad_ranges = _dedup_normalized_ranges(bad_ranges + heuristic_bad_ranges)
        if correct_ranges:
            filtered_correct: list[str] = []
            for rng in correct_ranges:
                try:
                    if any(_range_intersects(rng, bad_rng) for bad_rng in heuristic_bad_ranges):
                        continue
                except Exception:
                    pass
                filtered_correct.append(rng)
            correct_ranges = filtered_correct
        heuristic_reason = "; ".join(str(item.get("reason") or "").strip() for item in heuristic_findings[:2] if str(item.get("reason") or "").strip())
        if heuristic_reason:
            reason = heuristic_reason if not reason.strip() else f"{reason} Heuristic check also found: {heuristic_reason}"
        conf = max(conf, 0.95)
    refine_raw = ""
    if bool(need_fix) and reason.strip() and len(bad_ranges) == 1:
        try:
            refined_focus, refine_raw, refine_usage = _refine_bad_range(reason, bad_ranges[0] or (bad_ranges_raw[0] if bad_ranges_raw else ""))
            if refined_focus:
                bad_ranges = [refined_focus]
            for k in ("prompt_tokens", "completion_tokens", "total_tokens"):
                usage[k] = int(usage.get(k, 0)) + int(refine_usage.get(k, 0))
        except Exception:
            pass
    if not bool(need_fix):
        bad_ranges = []
    result = {
        "need_fix": bool(need_fix),
        "bad_ranges": bad_ranges,
        "bad_ranges_raw": bad_ranges_raw,
        "bad_range": bad_ranges[0] if bad_ranges else "",
        "bad_range_raw": bad_ranges_raw[0] if bad_ranges_raw else "",
        "correct_ranges": correct_ranges,
        "bad_range_parse_failed": bool(bad_range_parse_failed),
        "reason": reason,
        "confidence": conf,
        "heuristic_bad_ranges": heuristic_bad_ranges,
        "heuristic_findings": heuristic_findings,
        "raw": raw,
        "refine_raw": refine_raw,
        "before_sheet_png": str(before_png),
        "after_sheet_png": str(after_png),
        "usage": usage,
    }
    return result


def main() -> None:
    ap = argparse.ArgumentParser(description="Pipeline with agent-based filling step.")
    ap.add_argument("--input", required=True, help="Input Excel (.xlsx/.xls)")
    ap.add_argument("--instruction", required=True, help="Natural language instruction")
    ap.add_argument("--planner_json_path", default="", help="Optional structured planner policy JSON.")
    ap.add_argument("--agentic_planner_enable", action="store_true", help="Generate planner policy inside my-pipeline before first pass.")
    ap.add_argument("--planner_model", default="", help="Planner model. Defaults to agent_model.")
    ap.add_argument("--planner_base_url", default="", help="Planner base URL. Defaults to agent_base_url.")
    ap.add_argument("--planner_api_key", default="", help="Planner API key. Defaults to agent_api_key.")
    ap.add_argument("--planner_temperature", type=float, default=0.0, help="Planner temperature.")
    ap.add_argument("--planner_debug_dir", default="", help="Optional planner debug directory. Default: runs/agent_pipeline/planner")
    ap.add_argument("--save_model_io", action="store_true", help="Save per-call model input/output logs.")
    ap.add_argument("--model_io_dump_dir", default="", help="Optional model I/O dump directory. Default: runs/agent_pipeline/model_io")
    ap.add_argument("--output", required=True, help="Output filled Excel")
    ap.add_argument("--sheet", default=None, help="Sheet name (default: active)")
    ap.add_argument("--dpi", type=int, default=96)
    ap.add_argument("--ckpt", default="", help="SegFormer ckpt path")
    ap.add_argument(
        "--preprocess_backend",
        default="html_css",
        choices=["html_css", "libreoffice", "excel_com"],
        help="Preprocess mode. html_css is the new cross-platform renderer.",
    )
    ap.add_argument("--soffice", default="", help="Deprecated; kept for CLI compatibility. Ignored in html_css mode.")
    ap.add_argument("--render_canvas_bg", default="#f5f5f5", help="Background color for html_css render.")
    ap.add_argument("--render_png_scale", type=float, default=1.0, help="Scale factor for html_css screenshot.")
    ap.add_argument("--render_max_rows", type=int, default=400, help="Max rows to render from used-range (0 means no limit).")
    ap.add_argument("--render_max_cols", type=int, default=120, help="Max cols to render from used-range (0 means no limit).")

    ap.add_argument("--qwen_mode", default="lora", choices=["api", "lora"])
    ap.add_argument("--qwen_base_url", default="")
    ap.add_argument("--qwen_model", default="")
    ap.add_argument("--qwen_api_key", default="")
    ap.add_argument("--qwen_lora_model", default="")
    ap.add_argument("--qwen_lora_path", default="")
    ap.add_argument("--qwen_lora_device_map", default="auto")
    ap.add_argument("--qwen_lora_max_new_tokens", type=int, default=40960)
    ap.add_argument("--qwen_lora_temperature", type=float, default=0.0)
    ap.add_argument("--qwen_lora_slots_chunk_size", type=int, default=8)
    ap.add_argument("--qwen_lora_max_image_size", type=int, default=1600)
    ap.add_argument("--qwen_lora_crop_left_pad", type=int, default=220)
    ap.add_argument("--qwen_lora_crop_top_pad", type=int, default=140)
    ap.add_argument("--qwen_lora_crop_right_pad", type=int, default=80)
    ap.add_argument("--qwen_lora_crop_bottom_pad", type=int, default=80)
    ap.add_argument("--qwen_lora_crop_by_slot_chunk", dest="qwen_lora_crop_by_slot_chunk", action="store_true")
    ap.add_argument("--qwen_lora_no_crop_by_slot_chunk", dest="qwen_lora_crop_by_slot_chunk", action="store_false")
    ap.set_defaults(qwen_lora_crop_by_slot_chunk=True)
    ap.add_argument("--qwen_lora_resize_after_crop", dest="qwen_lora_resize_after_crop", action="store_true")
    ap.add_argument("--qwen_lora_no_resize_after_crop", dest="qwen_lora_resize_after_crop", action="store_false")
    ap.set_defaults(qwen_lora_resize_after_crop=True)
    ap.add_argument("--qwen_lora_annotate_slot_id", dest="qwen_lora_annotate_slot_id", action="store_true")
    ap.add_argument("--qwen_lora_no_annotate_slot_id", dest="qwen_lora_annotate_slot_id", action="store_false")
    ap.set_defaults(qwen_lora_annotate_slot_id=True)

    ap.add_argument("--agent_model", default="gemini-3-flash-preview")
    ap.add_argument("--agent_base_url", default="")
    ap.add_argument("--agent_api_key", default="")
    ap.add_argument("--agent_temperature", type=float, default=0.0, help="Temperature for agent API inference.")
    ap.add_argument("--tool_mode", default="json", choices=["json", "native", "code"])
    ap.add_argument("--max_steps", type=int, default=10)
    ap.add_argument("--external_qwen_pairs_json", default="", help="Optional precomputed qwen_pairs.json path. If set, skip preprocess/slot/qwen inference.")
    ap.add_argument("--agent_skills", default="", help="Comma-separated lightweight skills for agent behavior tuning.")
    ap.add_argument("--agent_skills_file", default="", help="Path to skills text/json file injected into system prompt.")
    ap.add_argument("--skip_initial_fill", action="store_true", help="Skip first-pass fill and run reflect on existing output workbook.")
    ap.add_argument(
        "--first_pass_output_dir",
        default="",
        help="Optional directory to save the first-pass filled workbook before reflect patches.",
    )
    ap.add_argument(
        "--initial_hint_mode",
        default="off",
        choices=["off", "global_qwen"],
        help="off: first-pass agent fill runs without global qwen hints; global_qwen: keep old behavior and run full-sheet qwen before first fill.",
    )
    ap.add_argument(
        "--agent_mapping_mode",
        default="off",
        choices=["off", "hint"],
        help="off: do not provide qwen key->cell hints to agent; hint: provide as low-confidence hints.",
    )
    ap.add_argument("--reflect_enable", action="store_true", help="Enable screenshot reflection and local re-fill patch.")
    ap.add_argument("--reflect_plugin_enable", action="store_true", help="Enable edit-driven reflect plugin.")
    ap.add_argument("--reflect_max_rounds", type=int, default=1, help="Max reflection rounds.")
    ap.add_argument("--reflect_patch_steps", type=int, default=6, help="Max code-agent steps for each local patch round.")
    ap.add_argument("--reflect_checker_model", default="", help="Checker model for screenshot reflection (default: agent_model).")
    ap.add_argument("--reflect_edit_log_jsonl", default="", help="Optional edit log JSONL path. Default: runs/agent_pipeline/edit_log.jsonl")
    ap.add_argument("--reflect_risk_threshold", type=int, default=1, help="High-risk edit threshold.")
    ap.add_argument("--reflect_expand_left_cols", type=int, default=3, help="Reflect plugin local expand cols to the left.")
    ap.add_argument("--reflect_expand_right_cols", type=int, default=3, help="Reflect plugin local expand cols to the right.")
    ap.add_argument("--reflect_expand_top_rows", type=int, default=1, help="Reflect plugin local expand rows to the top.")
    ap.add_argument("--reflect_expand_bottom_rows", type=int, default=1, help="Reflect plugin local expand rows to the bottom.")
    ap.add_argument("--reflect_focus_pad_left_cols", type=int, default=5, help="Reflect local crop: expand cols to the left.")
    ap.add_argument("--reflect_focus_pad_top_rows", type=int, default=2, help="Reflect local crop: expand rows to the top.")
    ap.add_argument("--reflect_focus_pad_right_cols", type=int, default=3, help="Reflect local crop: expand cols to the right.")
    ap.add_argument("--reflect_focus_pad_bottom_rows", type=int, default=2, help="Reflect local crop: expand rows to the bottom.")

    args = ap.parse_args()

    log_dir = Path(__file__).resolve().parent
    log_path = log_dir / "pipeline.log"
    log_dir.mkdir(parents=True, exist_ok=True)

    def log_fn(msg: str) -> None:
        _emit_console_line(msg)
        try:
            with open(log_path, "a", encoding="utf-8") as f:
                f.write(msg.rstrip() + "\n")
        except Exception:
            pass

    def load_skills_text(skills_inline: str, skills_file: str) -> str:
        parts: list[str] = []
        inline = (skills_inline or "").strip()
        if inline:
            items = [x.strip() for x in inline.split(",") if x.strip()]
            if items:
                parts.append("Inline skills:")
                parts.extend([f"- {x}" for x in items])
        fp = (skills_file or "").strip()
        if fp:
            p = Path(fp)
            if not p.is_absolute():
                p = Path.cwd() / p
            if not p.exists():
                raise PipelineError(f"agent_skills_file not found: {p}")
            raw = p.read_text(encoding="utf-8").strip()
            if not raw:
                return "\n".join(parts)
            try:
                obj: Any = json.loads(raw)
                if isinstance(obj, dict):
                    parts.append("File skills (json object):")
                    parts.append(json.dumps(obj, ensure_ascii=False, indent=2))
                elif isinstance(obj, list):
                    parts.append("File skills (json list):")
                    for it in obj:
                        parts.append(f"- {it}")
                else:
                    parts.append("File skills:")
                    parts.append(str(obj))
            except Exception:
                parts.append("File skills:")
                parts.append(raw)
        return "\n".join(parts)

    job_dir = Path.cwd() / "runs" / "agent_pipeline"
    job_dir.mkdir(parents=True, exist_ok=True)
    edit_log_jsonl = Path(args.reflect_edit_log_jsonl).resolve() if args.reflect_edit_log_jsonl else (job_dir / "edit_log.jsonl")
    model_io_dir = (
        Path(args.model_io_dump_dir).resolve()
        if str(args.model_io_dump_dir).strip()
        else (job_dir / "model_io" if args.save_model_io else None)
    )
    if model_io_dir is not None:
        model_io_dir.mkdir(parents=True, exist_ok=True)

    xlsx_path = job_dir / "input.xlsx"
    xlsx_path = ensure_xlsx(Path(args.input), xlsx_path, log_fn)
    before_fill_xlsx = job_dir / "before_fill.xlsx"
    shutil.copyfile(xlsx_path, before_fill_xlsx)
    usage_summary: dict[str, Any] = {
        "first_pass_agent": {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0},
        "reflect_checker": {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0},
        "reflect_model_assess": {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0},
        "patch_agents": [],
        "reflect_plugin_stats": {
            "rounds_started": 0,
            "stage1_need_fix_count": 0,
            "stage1_parse_failed_count": 0,
            "assessed_cells_total": 0,
            "edited_cells_total": 0,
            "scan_only_cells_total": 0,
            "high_risk_cells_total": 0,
            "targets_total": 0,
            "blank_targets_total": 0,
            "patch_jobs_total": 0,
            "patch_written_actual_edits_total": 0,
        },
    }

    run_initial_global_qwen = bool(args.external_qwen_pairs_json) or (args.initial_hint_mode == "global_qwen")
    form: dict[str, Any] = {"pairs": []}
    if args.external_qwen_pairs_json:
        qwen_pairs_path = Path(args.external_qwen_pairs_json).resolve()
        if not qwen_pairs_path.exists():
            raise PipelineError(f"external_qwen_pairs_json not found: {qwen_pairs_path}")
        log_fn(f"[INFO] Using external qwen pairs: {qwen_pairs_path}")
        form = build_form_from_qwen(qwen_pairs_path, xlsx_path=xlsx_path, sheet_name=args.sheet)
        if isinstance(form, dict) and isinstance(form.get("dropped_suspicious_pairs"), list) and form.get("dropped_suspicious_pairs"):
            log_fn(f"[INFO] dropped suspicious qwen pairs: {len(form.get('dropped_suspicious_pairs') or [])}")
        if isinstance(form, dict) and isinstance(form.get("invalid_slots"), list) and form.get("invalid_slots"):
            log_fn(f"[INFO] qwen invalid slots available as negative hints: {len(form.get('invalid_slots') or [])}")
    elif run_initial_global_qwen:
        preprocess_dir = job_dir / "preprocess"
        if args.preprocess_backend != "html_css":
            log_fn(f"[WARN] preprocess_backend={args.preprocess_backend} is deprecated in my-pipeline.py; forcing html_css")
        preprocess_outputs = preprocess_excel_html_css(
            xlsx_path=xlsx_path,
            out_dir=preprocess_dir,
            sheet_name=args.sheet,
            dpi=args.dpi,
            canvas_bg=args.render_canvas_bg,
            png_scale=args.render_png_scale,
            max_render_rows=args.render_max_rows,
            max_render_cols=args.render_max_cols,
        )

        qwen_out_dir = job_dir / "qwen_out"
        qwen_outputs = run_qwen_small(
            image_path=preprocess_outputs["sheet_png"],
            xlsx_path=xlsx_path,
            out_dir=qwen_out_dir,
            sheet_name=args.sheet,
            ckpt_path=args.ckpt or None,
            edges_json=preprocess_outputs["edges_json"],
            bounds_json=preprocess_outputs["bounds_json"],
            qwen_base_url=args.qwen_base_url or None,
            qwen_model=args.qwen_model or None,
            qwen_api_key=args.qwen_api_key or None,
            skip_qwen=(args.qwen_mode == "lora"),
            log_fn=log_fn,
        )

        if args.qwen_mode == "lora":
            if not args.qwen_lora_model or not args.qwen_lora_path:
                raise PipelineError("qwen_lora_model and qwen_lora_path are required for qwen_mode=lora")
            qwen_pairs_path = run_qwen_lora(
                image_path=preprocess_outputs["sheet_png"],
                slots_json=qwen_outputs["slots_json"],
                out_dir=qwen_out_dir,
                model_name=args.qwen_lora_model,
                lora_path=args.qwen_lora_path,
                device_map=args.qwen_lora_device_map,
                max_new_tokens=args.qwen_lora_max_new_tokens,
                temperature=args.qwen_lora_temperature,
                slots_chunk_size=args.qwen_lora_slots_chunk_size,
                crop_by_slot_chunk=args.qwen_lora_crop_by_slot_chunk,
                crop_left_pad=args.qwen_lora_crop_left_pad,
                crop_top_pad=args.qwen_lora_crop_top_pad,
                crop_right_pad=args.qwen_lora_crop_right_pad,
                crop_bottom_pad=args.qwen_lora_crop_bottom_pad,
                max_image_size=args.qwen_lora_max_image_size,
                resize_after_crop=args.qwen_lora_resize_after_crop,
                annotate_slot_id=args.qwen_lora_annotate_slot_id,
                log_fn=log_fn,
                debug_dir=qwen_out_dir / "lora_debug",
            )
        else:
            qwen_pairs_path = qwen_outputs["qwen_pairs"]
        form = build_form_from_qwen(qwen_pairs_path, xlsx_path=xlsx_path, sheet_name=args.sheet)
        if isinstance(form, dict) and isinstance(form.get("dropped_suspicious_pairs"), list) and form.get("dropped_suspicious_pairs"):
            log_fn(f"[INFO] dropped suspicious qwen pairs: {len(form.get('dropped_suspicious_pairs') or [])}")
        if isinstance(form, dict) and isinstance(form.get("invalid_slots"), list) and form.get("invalid_slots"):
            log_fn(f"[INFO] qwen invalid slots available as negative hints: {len(form.get('invalid_slots') or [])}")
    else:
        log_fn("[INFO] initial_hint_mode=off; first-pass agent fill will run without global qwen hints.")
    skills_text = load_skills_text(args.agent_skills, args.agent_skills_file)
    planner_policy: dict[str, Any] = {}
    planner_debug_root: Path | None = None
    if str(args.planner_json_path).strip():
        planner_policy = _load_planner_policy(args.planner_json_path)
        planner_debug_root = Path(args.planner_json_path).resolve().parent
    elif args.agentic_planner_enable:
        planner_model = str(args.planner_model or args.agent_model).strip()
        planner_base_url = str(args.planner_base_url or args.agent_base_url).strip() or None
        planner_api_key = str(args.planner_api_key or args.agent_api_key).strip() or None
        planner_debug_dir = (
            Path(args.planner_debug_dir).resolve()
            if str(args.planner_debug_dir).strip()
            else (job_dir / "planner")
        )
        planner_debug_root = planner_debug_dir
        planner_debug_dir.mkdir(parents=True, exist_ok=True)
        planner_struct_summary = _scan_structure_for_planner(xlsx_path, args.sheet)
        try:
            planner_policy = _request_planner_policy(
                instruction=args.instruction,
                struct_summary=planner_struct_summary,
                sheet_name=args.sheet,
                model=planner_model,
                base_url=planner_base_url,
                api_key=planner_api_key,
                temperature=float(args.planner_temperature),
                model_io_dir=(model_io_dir / "planner" if model_io_dir is not None else None),
            )
            planner_failed = False
        except Exception as exc:
            planner_failed = True
            planner_policy = {
                "form_type": "generic_form",
                "global_policies": {
                    "leave_unspecified_blank": True,
                    "preserve_existing_content": True,
                    "avoid_guessing": True,
                },
                "field_groups": [],
                "blank_policies": [],
                "do_not_fill_sections": [],
                "checkbox_policies": [],
                "repeated_block_policy": {
                    "has_repeated_blocks": False,
                    "block_key": "",
                    "expected_count": None,
                },
                "high_confidence_fields": [],
                "uncertain_fields": [],
                "notes": [f"planner_failed: {exc}"],
            }
        planner_payload = {k: v for k, v in planner_policy.items() if not str(k).startswith("_")}
        (planner_debug_dir / "struct_summary.json").write_text(
            json.dumps(planner_struct_summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (planner_debug_dir / "plan.json").write_text(
            json.dumps(planner_policy, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (planner_debug_dir / "plan.normalized.json").write_text(
            json.dumps(planner_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (planner_debug_dir / "planner_policy_preview.txt").write_text(
            _build_planner_policy_preview(planner_policy),
            encoding="utf-8",
        )
        (planner_debug_dir / "instruction_original.txt").write_text(
            str(args.instruction or ""),
            encoding="utf-8",
        )
        log_fn(
            f"[INFO] planner_status={'failed' if planner_failed else 'ok'} "
            f"plan={planner_debug_dir / 'plan.normalized.json'}"
        )
    workflow_policy = _planner_workflow_policy(planner_policy)
    if planner_policy:
        log_fn(
            "[INFO] planner policy loaded: "
            f"field_groups={len(planner_policy.get('field_groups') or [])} "
            f"high_confidence_fields={len(planner_policy.get('high_confidence_fields') or [])} "
            f"background_reads={len(planner_policy.get('background_reads') or [])} "
            f"read_bg_first={workflow_policy.get('read_background_before_first_pass')} "
            f"read_bg_refill={workflow_policy.get('read_background_before_refill')} "
            f"prefer_hints={workflow_policy.get('prefer_hints_when_available')} "
            f"run_reflect={workflow_policy.get('run_reflect_after_first_pass')}"
        )
    output_path = Path(args.output)
    first_pass_result: dict[str, Any] = {}
    first_pass_sheet_png = ""

    if args.skip_initial_fill:
        if not output_path.exists():
            raise PipelineError(
                f"skip_initial_fill requires existing output file: {output_path}. "
                "Prepare/copy prefilled workbook before running reflect."
            )
        log_fn(f"[INFO] skip_initial_fill enabled; reusing existing output: {output_path}")
    else:
        if edit_log_jsonl.exists():
            edit_log_jsonl.unlink()
        initial_mapping_mode = args.agent_mapping_mode
        if initial_mapping_mode == "hint" and not workflow_policy.get("prefer_hints_when_available", True):
            log_fn("[INFO] planner controller disabled first-pass hint mode.")
            initial_mapping_mode = "off"
        if (not form.get("pairs")) and initial_mapping_mode == "hint":
            log_fn("[INFO] first-pass has no global qwen pairs; forcing initial mapping_mode=off.")
            initial_mapping_mode = "off"
        first_pass_context_summary: dict[str, Any] | None = None
        first_pass_background_context = _collect_planner_background_context(
            xlsx_path=xlsx_path,
            sheet_name=args.sheet,
            planner_policy=planner_policy,
            stage="first_pass",
        )
        if first_pass_background_context:
            first_pass_context_summary = {"planner_background_context": first_pass_background_context}
            if planner_debug_root is not None:
                (planner_debug_root / "background_context_first_pass.json").write_text(
                    json.dumps(first_pass_background_context, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
            log_fn(
                f"[INFO] planner controller read background before first pass: "
                f"{len(first_pass_background_context.get('items') or [])} context windows"
            )
        first_pass_result = agent_fill(
            template_xlsx=xlsx_path,
            output_xlsx=output_path,
            instruction=args.instruction,
            form=form,
            model=args.agent_model,
            base_url=args.agent_base_url,
            api_key=args.agent_api_key or None,
            tool_mode=args.tool_mode,
            max_steps=args.max_steps,
            skills_text=skills_text,
            mapping_mode=initial_mapping_mode,
            sheet_name=args.sheet,
            edit_log_jsonl=edit_log_jsonl,
            round_idx=0,
            context_summary=first_pass_context_summary,
            history_dump_path=job_dir / "first_pass_history.json",
            temperature=float(args.agent_temperature),
            planner_policy=planner_policy,
            planner_stage="first_pass",
            model_io_dir=(model_io_dir / "first_pass_agent" if model_io_dir is not None else None),
        )
        usage_summary["first_pass_agent"] = first_pass_result.get("usage", usage_summary["first_pass_agent"])
        if str(args.first_pass_output_dir).strip():
            first_pass_dir = Path(args.first_pass_output_dir).resolve()
            first_pass_dir.mkdir(parents=True, exist_ok=True)
            first_pass_snapshot = first_pass_dir / output_path.name
            shutil.copyfile(output_path, first_pass_snapshot)
            first_pass_result["first_pass_output_xlsx"] = str(first_pass_snapshot)
            log_fn(f"[INFO] saved first-pass workbook: {first_pass_snapshot}")
        try:
            first_pass_render = preprocess_excel_html_css(
                xlsx_path=output_path,
                out_dir=job_dir / "first_pass_visual",
                sheet_name=args.sheet,
                dpi=args.dpi,
                canvas_bg=args.render_canvas_bg,
                png_scale=args.render_png_scale,
                max_render_rows=args.render_max_rows,
                max_render_cols=args.render_max_cols,
            )
            first_pass_sheet_png = str(first_pass_render.get("sheet_png") or "")
            if first_pass_sheet_png:
                first_pass_result["first_pass_sheet_png"] = first_pass_sheet_png
                log_fn(f"[INFO] saved first-pass screenshot: {first_pass_sheet_png}")
        except Exception as exc:
            log_fn(f"[WARN] failed to render first-pass screenshot: {exc}")
    if not first_pass_sheet_png and output_path.exists():
        try:
            first_pass_render = preprocess_excel_html_css(
                xlsx_path=output_path,
                out_dir=job_dir / "first_pass_visual",
                sheet_name=args.sheet,
                dpi=args.dpi,
                canvas_bg=args.render_canvas_bg,
                png_scale=args.render_png_scale,
                max_render_rows=args.render_max_rows,
                max_render_cols=args.render_max_cols,
            )
            first_pass_sheet_png = str(first_pass_render.get("sheet_png") or "")
            if first_pass_sheet_png:
                first_pass_result["first_pass_sheet_png"] = first_pass_sheet_png
        except Exception:
            pass

    def _is_noisy_reflect_key(key: str) -> bool:
        text = str(key or "").strip()
        if not text:
            return True
        compact = re.sub(r"\s+", " ", text)
        if re.fullmatch(r"[\d\W_]+", compact):
            return True
        alnum = re.sub(r"[^A-Za-z0-9]+", "", compact)
        if len(alnum) <= 1:
            return True
        if re.fullmatch(r"\d+[A-Za-z]?", alnum):
            return True
        return False

    def _keys_soft_match(a: str, b: str) -> bool:
        na = normalize_key(a)
        nb = normalize_key(b)
        if not na or not nb:
            return False
        if na == nb:
            return True
        if na in nb or nb in na:
            return True
        ta = set(na.split())
        tb = set(nb.split())
        if not ta or not tb:
            return False
        overlap = len(ta & tb) / max(1, min(len(ta), len(tb)))
        return overlap >= 0.75

    def _hint_matches_target(
        key: str,
        candidate_value_cell: str,
        region_targets: list[ReflectTarget],
    ) -> bool:
        for target in region_targets:
            target_key = str(getattr(target, "key", "") or "")
            if _keys_soft_match(key, target_key):
                return True
            try:
                if candidate_value_cell and _range_intersects(target.target_range, candidate_value_cell):
                    return True
            except Exception:
                continue
        return False

    def _select_patch_targets(patch_job: dict[str, Any]) -> list[dict[str, Any]]:
        all_targets = [x for x in (patch_job.get("targets", []) or []) if isinstance(x, dict)]
        prioritized = [
            x
            for x in all_targets
            if str(x.get("kind") or "") in {"region_blank_candidate", "screenshot_bad_range"}
        ]
        return prioritized or all_targets

    def _select_patch_region_hints(
        region_hints: list[dict[str, Any]],
        patch_targets: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        selected: list[dict[str, Any]] = []
        target_keys = {
            normalize_key(str(x.get("key") or ""))
            for x in patch_targets
            if str(x.get("key") or "").strip()
        }
        for hint in region_hints:
            if not isinstance(hint, dict):
                continue
            key = str(hint.get("key") or "").strip()
            value_cell = str(hint.get("candidate_value_cell") or "").strip()
            norm_key = normalize_key(key)
            matched = bool(norm_key and norm_key in target_keys)
            if not matched:
                for target in patch_targets:
                    try:
                        if value_cell and _range_intersects(str(target.get("target_range") or ""), value_cell):
                            matched = True
                            break
                    except Exception:
                        continue
            if matched:
                selected.append(hint)
        return selected

    def _build_reflect_qwen_prompt_context(region: str, region_targets: list[ReflectTarget]) -> str:
        flag_order = [
            "label_area",
            "duplicate_pattern",
            "template_inconsistency",
            "high_risk_pattern",
            "special_char_anomaly",
        ]
        region_flags: list[str] = []
        for target in region_targets:
            for flag in getattr(target, "risk_flags", []) or []:
                if flag in flag_order and flag not in region_flags:
                    region_flags.append(flag)
        if not region_flags:
            return ""
        flag_notes = {
            "label_area": (
                "Do not treat title/header/label text as the value key for a slot. "
                "The correct output should point to a nearby fillable value field, not the label cell itself."
            ),
            "duplicate_pattern": (
                "This region may be part of repeated template blocks. Match the key to the slot in the current local block, "
                "not to a similar key from another repeated block."
            ),
            "template_inconsistency": (
                "Prefer the key/value pattern that is visually consistent with neighboring fields in this same section."
            ),
            "high_risk_pattern": (
                "Be careful about vertical misalignment. The correct result should avoid shifting a value to the row above or below its true label."
            ),
            "special_char_anomaly": (
                "Preserve accented or language-specific characters when identifying the key text. Do not normalize them away."
            ),
        }
        few_shots = {
            "label_area": (
                "Example: if a slot sits under the label 'Project Name:', the correct key is 'Project Name', "
                "not a nearby section title like 'Project Information'."
            ),
            "duplicate_pattern": (
                "Example: if 'Date' appears in multiple repeated rows, select the 'Date' from the same repeated row as the slot, "
                "not another row's 'Date'."
            ),
            "template_inconsistency": (
                "Example: if a section contains 'City / State / Zip' style keys, prefer a candidate key matching that local section pattern."
            ),
            "high_risk_pattern": (
                "Example: if a value slot is directly below 'Incident Date', do not assign it to 'Incident Time' from the next row."
            ),
            "special_char_anomaly": (
                "Example: keep 'José' aligned with 'José', not 'Jose'."
            ),
        }
        parts = [
            f"Reflect risk context for region {region}: {', '.join(region_flags)}.",
            "Use these flags as constraints while deciding the correct key for each slot.",
        ]
        for flag in region_flags:
            parts.append(f"{flag}: {flag_notes[flag]}")
            parts.append(f"Few-shot: {few_shots[flag]}")
        return "\n".join(parts)

    def build_reflect_region_hints(region: str, region_targets: list[ReflectTarget], out_dir: Path) -> list[dict[str, Any]]:
        out_dir.mkdir(parents=True, exist_ok=True)
        try:
            prep = preprocess_excel_html_css(
                xlsx_path=before_fill_xlsx,
                out_dir=out_dir / "preprocess",
                sheet_name=args.sheet,
                dpi=args.dpi,
                canvas_bg=args.render_canvas_bg,
                png_scale=args.render_png_scale,
                max_render_rows=args.render_max_rows,
                max_render_cols=args.render_max_cols,
                focus_range=region,
                focus_pad_left_cols=0,
                focus_pad_top_rows=0,
                focus_pad_right_cols=0,
                focus_pad_bottom_rows=0,
            )
            qwen_out_dir = out_dir / "qwen_out"
            qwen_outputs = run_qwen_small(
                image_path=prep["sheet_png"],
                xlsx_path=before_fill_xlsx,
                out_dir=qwen_out_dir,
                sheet_name=args.sheet,
                ckpt_path=args.ckpt or None,
                edges_json=prep["edges_json"],
                bounds_json=prep["bounds_json"],
                qwen_base_url=args.qwen_base_url or None,
                qwen_model=args.qwen_model or None,
                qwen_api_key=args.qwen_api_key or None,
                skip_qwen=(args.qwen_mode == "lora"),
                log_fn=log_fn,
            )
            if args.qwen_mode == "lora":
                if not args.qwen_lora_model or not args.qwen_lora_path:
                    raise PipelineError("reflect plugin requires qwen_lora_model and qwen_lora_path when qwen_mode=lora")
                prompt_context_text = _build_reflect_qwen_prompt_context(region, region_targets)
                reflect_slots_json = qwen_out_dir / "slots_in_region.json"
                kept_slots = _filter_slots_json_by_focus_range(
                    slots_json_path=qwen_outputs["slots_json"],
                    focus_range=region,
                    out_json_path=reflect_slots_json,
                )
                if kept_slots <= 0:
                    log_fn(f"[REFLECT_PLUGIN] region={region} kept_slots=0 after focus-range filtering; no Qwen hints.")
                    return []
                reflect_pairs_path = run_qwen_lora(
                    image_path=prep["sheet_png"],
                    slots_json=reflect_slots_json,
                    out_dir=qwen_out_dir,
                    model_name=args.qwen_lora_model,
                    lora_path=args.qwen_lora_path,
                    device_map=args.qwen_lora_device_map,
                    max_new_tokens=args.qwen_lora_max_new_tokens,
                    temperature=args.qwen_lora_temperature,
                    slots_chunk_size=args.qwen_lora_slots_chunk_size,
                    crop_by_slot_chunk=args.qwen_lora_crop_by_slot_chunk,
                    crop_left_pad=args.qwen_lora_crop_left_pad,
                    crop_top_pad=args.qwen_lora_crop_top_pad,
                    crop_right_pad=args.qwen_lora_crop_right_pad,
                    crop_bottom_pad=args.qwen_lora_crop_bottom_pad,
                    max_image_size=args.qwen_lora_max_image_size,
                    resize_after_crop=args.qwen_lora_resize_after_crop,
                    annotate_slot_id=args.qwen_lora_annotate_slot_id,
                    prompt_context_text=prompt_context_text,
                    log_fn=log_fn,
                    debug_dir=qwen_out_dir / "lora_debug",
                )
            else:
                reflect_pairs_path = qwen_outputs["qwen_pairs"]

            region_form_all = build_form_from_qwen(reflect_pairs_path, xlsx_path=output_path, sheet_name=args.sheet)
            region_form = _filter_form_pairs_by_focus_range(region_form_all, region)
            qwen_pairs_total = len(region_form_all.get("pairs", []) or []) if isinstance(region_form_all, dict) else 0
            region_pairs_total = len(region_form.get("pairs", []) or []) if isinstance(region_form, dict) else 0
            qwen_invalid_total = len(region_form_all.get("invalid_slots", []) or []) if isinstance(region_form_all, dict) else 0
            region_invalid_total = len(region_form.get("invalid_slots", []) or []) if isinstance(region_form, dict) else 0
            hints: list[dict[str, Any]] = []
            dropped_noisy_blank = 0
            dropped_unrelated = 0
            for pair in region_form.get("pairs", []):
                if not isinstance(pair, dict):
                    continue
                key = str(pair.get("key") or "").strip()
                value_cell = str(pair.get("value_cell") or "").strip()
                if not key or not value_cell:
                    continue
                candidate_value = _read_anchor_value_from_workbook(output_path, args.sheet, value_cell)
                is_blank_candidate = _is_blankish_value(candidate_value)
                is_target_related = _hint_matches_target(key, value_cell, region_targets)
                risk_flags: list[str] = []
                for target in region_targets:
                    if _keys_soft_match(target.key, key):
                        risk_flags.extend(target.risk_flags)
                        continue
                    try:
                        if _range_intersects(target.target_range, value_cell):
                            risk_flags.extend(target.risk_flags)
                    except Exception:
                        continue
                if is_blank_candidate and _is_noisy_reflect_key(key):
                    dropped_noisy_blank += 1
                    continue
                if not is_target_related and not risk_flags:
                    dropped_unrelated += 1
                    continue
                hints.append(
                    {
                        "key": key,
                        "target_range": region,
                        "candidate_value_cell": value_cell,
                        "candidate_value": candidate_value,
                        "is_blank_candidate": is_blank_candidate,
                        "is_target_related": is_target_related,
                        "risk_flags": sorted(set(risk_flags)),
                        "hint_source": "reflect_blank_slot" if is_blank_candidate else "reflect_slot_kv",
                    }
                )
            invalid_hint_count = 0
            for invalid in region_form.get("invalid_slots", []):
                if not isinstance(invalid, dict):
                    continue
                key = str(invalid.get("key") or "").strip()
                value_cell = str(invalid.get("value_cell") or invalid.get("candidate_value_cell") or "").strip()
                if not value_cell:
                    continue
                candidate_value = _read_anchor_value_from_workbook(output_path, args.sheet, value_cell)
                is_blank_candidate = _is_blankish_value(candidate_value)
                is_target_related = _hint_matches_target(key, value_cell, region_targets)
                risk_flags: list[str] = []
                for target in region_targets:
                    if key and _keys_soft_match(target.key, key):
                        risk_flags.extend(target.risk_flags)
                        continue
                    try:
                        if _range_intersects(target.target_range, value_cell):
                            risk_flags.extend(target.risk_flags)
                    except Exception:
                        continue
                if not is_target_related and not risk_flags:
                    dropped_unrelated += 1
                    continue
                invalid_hint_count += 1
                hints.append(
                    {
                        "key": key,
                        "target_range": region,
                        "candidate_value_cell": value_cell,
                        "candidate_value": candidate_value,
                        "is_blank_candidate": is_blank_candidate,
                        "is_target_related": is_target_related,
                        "risk_flags": sorted(set(risk_flags)),
                        "hint_source": "reflect_invalid_slot",
                        "invalid_slot": True,
                        "avoid_cell": True,
                        "invalid_reason": str(invalid.get("reason") or "qwen_marked_invalid_slot"),
                    }
                )
            hints.sort(
                key=lambda item: (
                    not bool(item.get("invalid_slot")),
                    not bool(item.get("is_blank_candidate")),
                    not bool(item.get("is_target_related")),
                    -len(item.get("risk_flags", []) or []),
                    str(item.get("candidate_value_cell") or ""),
                )
            )
            blank_hint_count = sum(1 for item in hints if bool(item.get("is_blank_candidate")))
            log_fn(
                f"[REFLECT_PLUGIN] region={region} kept_slots={kept_slots if args.qwen_mode == 'lora' else 'n/a'} "
                f"qwen_pairs_total={qwen_pairs_total} qwen_pairs_in_region={region_pairs_total} "
                f"qwen_invalid_total={qwen_invalid_total} qwen_invalid_in_region={region_invalid_total} "
                f"qwen_hints={len(hints)} invalid_hints={invalid_hint_count} "
                f"blank_candidates={blank_hint_count} dropped_noisy_blank={dropped_noisy_blank} "
                f"dropped_unrelated={dropped_unrelated} "
                f"region_flags={json.dumps(sorted({flag for target in region_targets for flag in (getattr(target, 'risk_flags', []) or []) if flag in {'label_area', 'duplicate_pattern', 'template_inconsistency', 'high_risk_pattern', 'special_char_anomaly'}}), ensure_ascii=False)}"
            )
            return hints
        except Exception as exc:
            log_fn(f"[REFLECT_PLUGIN] failed to build region hints for region={region}: {exc}")
            return []

    def build_protected_write_cells(
        patch_job: dict[str, Any],
        reflect_result: Any,
        before_nonempty_cells: dict[str, str],
    ) -> dict[str, str]:
        protected: dict[str, str] = {}
        target_cells: set[str] = set()
        for target in patch_job.get("targets", []) or []:
            if not isinstance(target, dict):
                continue
            cell = str(target.get("source_cell") or target.get("target_range") or "").strip().upper()
            if cell:
                try:
                    target_cells.update(iter_range_cells(cell))
                except Exception:
                    target_cells.add(cell)
        for cell, value in before_nonempty_cells.items():
            if cell in target_cells:
                continue
            protected[cell] = f"template/original content: {value[:80]}"
        for rng in getattr(reflect_result, "correct_ranges", []) or []:
            text = str(rng or "").strip()
            if not text:
                continue
            try:
                cells_in_range = _range_cells_set([text])
            except Exception:
                continue
            for cell in cells_in_range:
                if cell in target_cells:
                    continue
                protected.setdefault(cell, "reflect marked this region as correct/safe_blank")
        for assess in getattr(reflect_result, "risk_assessments", []) or []:
            if not isinstance(assess, dict):
                continue
            if bool(assess.get("is_high_risk", False)):
                continue
            if (not bool(assess.get("is_correct_region", False))) and (not bool(assess.get("is_safe_blank", False))) and int(assess.get("score", 0) or 0) != 0:
                continue
            protect_ranges: list[str] = []
            target_range = str(assess.get("target_range") or "").strip()
            target_cell = str(assess.get("target_cell") or "").strip()
            candidate_value_cell = str(assess.get("candidate_value_cell") or "").strip()
            if target_range:
                protect_ranges.append(target_range)
            elif target_cell:
                protect_ranges.append(target_cell)
            if candidate_value_cell:
                protect_ranges.append(candidate_value_cell)
            for rng in protect_ranges:
                try:
                    cells_in_range = _range_cells_set([rng])
                except Exception:
                    cells_in_range = {str(rng).upper()}
                for cell in cells_in_range:
                    if not cell or cell in target_cells:
                        continue
                    protected.setdefault(cell, "previously checked and considered okay")
        return protected

    def _range_cells_set(range_list: list[str]) -> set[str]:
        cells: set[str] = set()
        for item in range_list:
            text = str(item or "").strip()
            if not text:
                continue
            try:
                cells.update(_iter_range_cells(text))
            except Exception:
                continue
        return {str(cell).upper() for cell in cells}

    def _rollback_region_from_template(
        template_xlsx: Path,
        target_xlsx: Path,
        sheet_name: str | None,
        region: str,
        preserve_cells: set[str] | None = None,
    ) -> int:
        preserve = {str(x).upper() for x in (preserve_cells or set())}
        src_wb = openpyxl.load_workbook(template_xlsx)
        dst_wb = openpyxl.load_workbook(target_xlsx)
        try:
            src_ws = src_wb[sheet_name] if sheet_name else src_wb.active
            dst_ws = dst_wb[sheet_name] if sheet_name else dst_wb.active
            src_merged = _merged_anchor_map(src_ws)
            dst_merged = _merged_anchor_map(dst_ws)
            touched_dst_anchors: set[str] = set()
            changed = 0
            for cell_ref in _iter_range_cells(region):
                target_ref = str(cell_ref).upper()
                if target_ref in preserve:
                    continue
                try:
                    col_idx, row_idx = range_boundaries(_normalize_a1_range(target_ref))[:2]
                except Exception:
                    continue
                src_anchor = src_merged.get((row_idx, col_idx), (row_idx, col_idx, row_idx, col_idx))
                dst_anchor = dst_merged.get((row_idx, col_idx), (row_idx, col_idx, row_idx, col_idx))
                src_anchor_a1 = f"{get_column_letter(src_anchor[1])}{src_anchor[0]}"
                dst_anchor_a1 = f"{get_column_letter(dst_anchor[1])}{dst_anchor[0]}"
                if dst_anchor_a1 in preserve or dst_anchor_a1 in touched_dst_anchors:
                    continue
                touched_dst_anchors.add(dst_anchor_a1)
                src_value = src_ws[src_anchor_a1].value
                if dst_ws[dst_anchor_a1].value != src_value:
                    dst_ws[dst_anchor_a1] = src_value
                    changed += 1
            if changed > 0:
                dst_wb.save(target_xlsx)
            return changed
        finally:
            src_wb.close()
            dst_wb.close()

    def _write_region_debug_meta(debug_dir: Path, payload: dict[str, Any]) -> None:
        debug_dir.mkdir(parents=True, exist_ok=True)
        (debug_dir / "meta.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )

    planner_allows_reflect = bool(workflow_policy.get("run_reflect_after_first_pass", True))
    if args.reflect_enable and not planner_allows_reflect:
        log_fn("[INFO] planner controller skipped reflect after first pass.")

    if args.reflect_enable and planner_allows_reflect:
        checker_model = args.reflect_checker_model.strip() or args.agent_model
        if args.reflect_plugin_enable:
            for ridx in range(1, max(1, int(args.reflect_max_rounds)) + 1):
                reflect_dir = job_dir / "reflect_plugin" / f"round_{ridx}"
                round_edits_before = len(_actual_edit_records_for_round(edit_log_jsonl, ridx))
                usage_summary["reflect_plugin_stats"]["rounds_started"] += 1
                screenshot_check: dict[str, Any] = {}
                checker_context_summary = {
                    "struct_summary": first_pass_result.get("struct_summary", {}),
                    "first_pass_write_log": (first_pass_result.get("write_log", []) or [])[:40],
                }
                try:
                    screenshot_check = reflect_check_with_screenshot(
                        before_xlsx_path=before_fill_xlsx,
                        after_xlsx_path=output_path,
                        sheet_name=args.sheet,
                        instruction=args.instruction,
                        model=checker_model,
                        base_url=args.agent_base_url or None,
                        api_key=args.agent_api_key or None,
                        out_dir=reflect_dir / "check",
                        dpi=args.dpi,
                        checker_context_summary=checker_context_summary,
                        planner_policy=planner_policy,
                        model_io_dir=((model_io_dir / f"reflect_round_{ridx}" / "stage1_checker") if model_io_dir is not None else None),
                    )
                    log_fn(f"[REFLECT_PLUGIN] round={ridx} screenshot_check={json.dumps(screenshot_check, ensure_ascii=False)}")
                    if bool(screenshot_check.get("need_fix")):
                        usage_summary["reflect_plugin_stats"]["stage1_need_fix_count"] += 1
                    if bool(screenshot_check.get("bad_range_parse_failed")):
                        usage_summary["reflect_plugin_stats"]["stage1_parse_failed_count"] += 1
                        log_fn(
                            "[REFLECT_PLUGIN][WARN] "
                            f"round={ridx} screenshot checker returned non-A1 bad_ranges_raw="
                            f"{json.dumps(screenshot_check.get('bad_ranges_raw', screenshot_check.get('bad_range_raw', '')), ensure_ascii=False)} "
                            "so Stage 1 cannot enqueue a patch region."
                        )
                    if isinstance(screenshot_check.get("usage"), dict):
                        usage_summary["reflect_checker"]["prompt_tokens"] += int(screenshot_check["usage"].get("prompt_tokens", 0))
                        usage_summary["reflect_checker"]["completion_tokens"] += int(screenshot_check["usage"].get("completion_tokens", 0))
                        usage_summary["reflect_checker"]["total_tokens"] += int(screenshot_check["usage"].get("total_tokens", 0))
                except Exception as exc:
                    log_fn(f"[REFLECT_PLUGIN] round={ridx} screenshot_check_failed: {exc}")
                try:
                    reflect_result = run_reflect_round(
                        ReflectContext(
                            workbook_path=output_path,
                            instruction=args.instruction,
                            form=form,
                            edit_log_jsonl=edit_log_jsonl,
                            out_dir=reflect_dir,
                            model=checker_model,
                            base_url=args.agent_base_url or None,
                            api_key=args.agent_api_key or None,
                            round_idx=ridx,
                            sheet_name=args.sheet,
                            risk_threshold=args.reflect_risk_threshold,
                            expand_left_cols=args.reflect_expand_left_cols,
                            expand_right_cols=args.reflect_expand_right_cols,
                            expand_top_rows=args.reflect_expand_top_rows,
                            expand_bottom_rows=args.reflect_expand_bottom_rows,
                            hint_builder=build_reflect_region_hints,
                            prior_context_summary={
                                "write_log": first_pass_result.get("write_log", []),
                                "history_dump_path": first_pass_result.get("history_dump_path", ""),
                            },
                            screenshot_check=screenshot_check,
                            first_pass_sheet_png=first_pass_sheet_png,
                            planner_policy_summary=_planner_policy_prompt_block(planner_policy, "reflect"),
                            planner_policy=planner_policy,
                            model_io_dir=((model_io_dir / f"reflect_round_{ridx}" / "stage2_assess") if model_io_dir is not None else None),
                        )
                    )
                except Exception as exc:
                    log_fn(f"[REFLECT_PLUGIN] round={ridx} failed: {exc}")
                    break
                before_nonempty_cells = _collect_nonempty_cells(before_fill_xlsx, args.sheet)
                assessed_cells = reflect_result.edits_considered if isinstance(reflect_result.edits_considered, list) else []
                high_risk_cells = [
                    x for x in (reflect_result.risk_assessments or [])
                    if isinstance(x, dict) and bool(x.get("is_high_risk"))
                ]
                correct_cells = [
                    x for x in (reflect_result.risk_assessments or [])
                    if isinstance(x, dict) and bool(x.get("is_correct_region"))
                ]
                safe_blank_cells = [
                    x for x in (reflect_result.risk_assessments or [])
                    if isinstance(x, dict) and bool(x.get("is_safe_blank"))
                ]
                blank_targets = [
                    x for x in (reflect_result.reflect_targets or [])
                    if isinstance(x, dict) and str(x.get("kind") or "") == "region_blank_candidate"
                ]
                edited_cells = [
                    x for x in assessed_cells
                    if isinstance(x, dict) and bool(x.get("is_actual_edit"))
                ]
                scan_only_cells = [
                    x for x in assessed_cells
                    if isinstance(x, dict) and not bool(x.get("is_actual_edit"))
                ]
                usage_summary["reflect_plugin_stats"]["assessed_cells_total"] += len(assessed_cells)
                usage_summary["reflect_plugin_stats"]["edited_cells_total"] += len(edited_cells)
                usage_summary["reflect_plugin_stats"]["scan_only_cells_total"] += len(scan_only_cells)
                usage_summary["reflect_plugin_stats"]["high_risk_cells_total"] += len(high_risk_cells)
                usage_summary["reflect_plugin_stats"]["targets_total"] += len(reflect_result.reflect_targets)
                usage_summary["reflect_plugin_stats"]["blank_targets_total"] += len(blank_targets)
                usage_summary["reflect_plugin_stats"]["patch_jobs_total"] += len(reflect_result.patch_jobs)

                log_fn(
                    f"[REFLECT_PLUGIN] round={ridx} assessed={len(reflect_result.edits_considered)} "
                    f"targets={len(reflect_result.reflect_targets)} regions={len(reflect_result.regions)}"
                )
                log_fn(
                    f"[REFLECT_PLUGIN][SUMMARY] round={ridx} "
                    f"stage1_need_fix={bool(screenshot_check.get('need_fix'))} "
                    f"parse_failed={bool(screenshot_check.get('bad_range_parse_failed'))} "
                    f"assessed={len(assessed_cells)} edited={len(edited_cells)} scan_only={len(scan_only_cells)} "
                    f"correct={len(correct_cells)} safe_blank={len(safe_blank_cells)} high_risk={len(high_risk_cells)} "
                    f"blank_targets={len(blank_targets)} patch_jobs={len(reflect_result.patch_jobs)} "
                    f"correct_ranges={len(getattr(reflect_result, 'correct_ranges', []) or [])}"
                )
                if isinstance(reflect_result.model_usage, dict):
                    usage_summary["reflect_model_assess"]["prompt_tokens"] += int(reflect_result.model_usage.get("prompt_tokens", 0))
                    usage_summary["reflect_model_assess"]["completion_tokens"] += int(reflect_result.model_usage.get("completion_tokens", 0))
                    usage_summary["reflect_model_assess"]["total_tokens"] += int(reflect_result.model_usage.get("total_tokens", 0))
                if reflect_result.risk_assessments:
                    for assess in reflect_result.risk_assessments:
                        if not isinstance(assess, dict):
                            continue
                        log_fn(
                            "[REFLECT_PLUGIN][ASSESS] "
                            f"round={ridx} cell={assess.get('target_cell')} "
                            f"key={json.dumps(assess.get('key', ''), ensure_ascii=False)} "
                            f"score={assess.get('score', 0)} "
                            f"correct={assess.get('is_correct_region', False)} "
                            f"safe_blank={assess.get('is_safe_blank', False)} "
                            f"high_risk={assess.get('is_high_risk', False)} "
                            f"flags={json.dumps(assess.get('risk_flags', []), ensure_ascii=False)} "
                            f"reason={json.dumps(assess.get('reason', ''), ensure_ascii=False)}"
                        )
                else:
                    log_fn(f"[REFLECT_PLUGIN] round={ridx} no risk assessments")
                if not reflect_result.patch_jobs:
                    log_fn(f"[REFLECT_PLUGIN] round={ridx} no patch jobs; stop.")
                    break

                any_patch = False
                patch_attempted_regions = 0
                region_debug_index: list[dict[str, Any]] = []
                for job_idx, patch_job in enumerate(reflect_result.patch_jobs, start=1):
                    region = str(patch_job.get("region") or "").strip()
                    region_hints = patch_job.get("region_hints", []) or []
                    if not region:
                        continue
                    region_debug_dir = reflect_dir / f"region_{job_idx}_debug"
                    region_debug_dir.mkdir(parents=True, exist_ok=True)
                    patch_attempted_regions += 1
                    correct_cells_in_region = {
                        cell
                        for cell in _range_cells_set(getattr(reflect_result, "correct_ranges", []) or [])
                        if _range_intersects(cell, region)
                    }
                    patch_targets = _select_patch_targets(patch_job)
                    selected_region_hints = _select_patch_region_hints(region_hints, patch_targets)
                    patch_pairs = []
                    for hint in selected_region_hints:
                        if not isinstance(hint, dict):
                            continue
                        if bool(hint.get("invalid_slot")) or bool(hint.get("avoid_cell")):
                            continue
                        key = str(hint.get("key") or "").strip()
                        candidate_value_cell = str(hint.get("candidate_value_cell") or "").strip()
                        if key and candidate_value_cell:
                            patch_pairs.append({"key": key, "value_cell": candidate_value_cell})
                    patch_mapping_mode = "hint" if patch_pairs else "off"
                    if patch_mapping_mode == "hint" and not workflow_policy.get("prefer_hints_when_available", True):
                        patch_mapping_mode = "off"
                        log_fn(
                            f"[REFLECT_PLUGIN] round={ridx} region={region} "
                            "planner controller disabled hint mode; using local background/context only."
                        )
                    if not patch_pairs:
                        log_fn(
                            f"[REFLECT_PLUGIN] round={ridx} region={region} no region hints/pairs; "
                            "falling back to local free patch."
                        )
                    log_fn(
                        f"[REFLECT_PLUGIN] round={ridx} region={region} "
                        f"patch_targets={len(patch_targets)} selected_hints={len(selected_region_hints)} "
                        f"patch_pairs={len(patch_pairs)} preserve_correct_cells={len(correct_cells_in_region)}"
                    )
                    refill_instruction = (
                        f"{args.instruction}\n\n"
                        f"[Reflect plugin refill round {ridx}, region {job_idx}]\n"
                        f"Focus on refilling region: {region}.\n"
                        "Cells in this region that were not trusted have been reverted back to the original template state.\n"
                        "Important: do not blindly trust the first-pass placement inside this region.\n"
                        "If Stage 1 or Stage 2 questioned a cell/range, assume that prior placement may be wrong.\n"
                        "Do not write a value back to the same original location unless the local labels, row/column structure, and nearby evidence clearly support that exact location.\n"
                        "If the evidence is ambiguous, prefer leaving the cell blank rather than copying the first-pass placement back into the same spot.\n"
                        "Keep already-correct cells unchanged, and refill only the missing/wrong cells in this region.\n"
                        "You must only write cells inside this focus region.\n"
                        "Only write cells directly supported by the reflect targets and reflect hints below.\n"
                        "Do not recompute or rewrite unrelated derived/formula-like cells unless they are explicitly listed as targets.\n"
                        "If write_cell returns a WARNING for a protected cell or outside-focus cell, do not try again on that cell.\n"
                        "First inspect the region with scan_structure/read_range, then write the needed cells, then stop immediately with final done.\n"
                    )
                    if patch_targets:
                        refill_instruction += (
                            "Refill targets:\n"
                            f"{json.dumps(patch_targets, ensure_ascii=False)}\n"
                        )
                        refill_instruction += (
                            "Treat every listed refill target as a previously questioned location. "
                            "Re-validate each one from the current region context instead of assuming the original location was correct.\n"
                        )
                    if selected_region_hints:
                        refill_instruction += (
                            "Prioritize the structured reflect hints below.\n"
                            f"Reflect hints:\n{json.dumps(selected_region_hints, ensure_ascii=False)}"
                        )
                    else:
                        refill_instruction += (
                            "No structured reflect hints are available for this region. "
                            "Use the instruction, current workbook state, and screenshot comparison context "
                            "to re-check and locally correct this region."
                        )
                    protected_write_cells = build_protected_write_cells(patch_job, reflect_result, before_nonempty_cells)
                    refill_instruction += (
                        f"\nProtected cell count: {len(protected_write_cells)}.\n"
                        "Protected cells include original template-content cells and cells previously assessed as okay. "
                        "If you need to change one of them, first verify with structure and nearby reads."
                    )
                    region_backup = reflect_dir / f"region_{job_idx}_before_refill.xlsx"
                    shutil.copyfile(output_path, region_backup)
                    region_before_snapshot = region_debug_dir / "before_refill.xlsx"
                    shutil.copyfile(output_path, region_before_snapshot)
                    reverted_cells = _rollback_region_from_template(
                        template_xlsx=before_fill_xlsx,
                        target_xlsx=output_path,
                        sheet_name=args.sheet,
                        region=region,
                        preserve_cells=correct_cells_in_region,
                    )
                    region_after_rollback_snapshot = region_debug_dir / "after_rollback.xlsx"
                    shutil.copyfile(output_path, region_after_rollback_snapshot)
                    debug_meta: dict[str, Any] = {
                        "round": ridx,
                        "region_index": job_idx,
                        "region": region,
                        "patch_targets": patch_targets,
                        "selected_region_hints": selected_region_hints,
                        "patch_pairs": patch_pairs,
                        "patch_mapping_mode": patch_mapping_mode,
                        "preserve_correct_cells": sorted(correct_cells_in_region),
                        "reverted_from_template": int(reverted_cells),
                        "protected_write_cells_count": len(protected_write_cells),
                        "screenshot_check": screenshot_check,
                        "first_pass_write_log_preview": (first_pass_result.get("write_log", []) or [])[:40],
                        "input_files": {
                            "before_fill_xlsx": str(before_fill_xlsx),
                            "region_before_refill_xlsx": str(region_before_snapshot),
                            "region_after_rollback_xlsx": str(region_after_rollback_snapshot),
                            "patch_history_json": str(reflect_dir / f"patch_region_{job_idx}_history.json"),
                        },
                    }
                    _write_region_debug_meta(region_debug_dir, debug_meta)
                    region_debug_index.append(
                        {
                            "region_index": job_idx,
                            "region": region,
                            "debug_dir": str(region_debug_dir),
                            "meta_json": str(region_debug_dir / "meta.json"),
                        }
                    )
                    log_fn(
                        f"[REFLECT_PLUGIN] round={ridx} region={region} reverted_from_template={reverted_cells}"
                    )
                    region_background_context = _collect_planner_background_context(
                        xlsx_path=output_path,
                        sheet_name=args.sheet,
                        planner_policy=planner_policy,
                        stage="refill",
                        focus_range=region,
                    )
                    if region_background_context:
                        (region_debug_dir / "planner_background_context.json").write_text(
                            json.dumps(region_background_context, ensure_ascii=False, indent=2),
                            encoding="utf-8",
                        )
                    try:
                        patch_result = agent_fill(
                            template_xlsx=output_path,
                            output_xlsx=output_path,
                            instruction=refill_instruction,
                            form={"pairs": patch_pairs},
                            model=args.agent_model,
                            base_url=args.agent_base_url,
                            api_key=args.agent_api_key or None,
                            tool_mode="code",
                            max_steps=max(1, int(args.reflect_patch_steps)),
                            skills_text=skills_text,
                            mapping_mode=patch_mapping_mode,
                            sheet_name=args.sheet,
                            allowed_write_range=region,
                            edit_log_jsonl=edit_log_jsonl,
                            round_idx=ridx,
                            reflect_hints=selected_region_hints,
                            context_summary={
                                "first_pass_write_log": first_pass_result.get("write_log", []),
                                "reflect_targets": patch_targets,
                                "screenshot_check": screenshot_check,
                                "planner_background_context": region_background_context,
                            },
                            compare_before_after_pngs={
                                "before": str(screenshot_check.get("before_sheet_png") or ""),
                                "after": str(screenshot_check.get("after_sheet_png") or ""),
                                "summary": json.dumps(
                                    {
                                        "need_fix": screenshot_check.get("need_fix"),
                                        "bad_range": screenshot_check.get("bad_range"),
                                        "reason": screenshot_check.get("reason"),
                                        "confidence": screenshot_check.get("confidence"),
                                    },
                                    ensure_ascii=False,
                                ),
                            },
                            history_dump_path=reflect_dir / f"patch_region_{job_idx}_history.json",
                            temperature=float(args.agent_temperature),
                            protected_write_cells=protected_write_cells,
                            planner_policy=planner_policy,
                            planner_stage="refill",
                            model_io_dir=((model_io_dir / f"reflect_round_{ridx}" / f"patch_region_{job_idx}") if model_io_dir is not None else None),
                        )
                    except Exception:
                        shutil.copyfile(region_backup, output_path)
                        region_after_restore_snapshot = region_debug_dir / "after_restore.xlsx"
                        shutil.copyfile(output_path, region_after_restore_snapshot)
                        debug_meta["status"] = "agent_fill_exception_restored"
                        debug_meta["output_files"] = {
                            "after_restore_xlsx": str(region_after_restore_snapshot),
                        }
                        _write_region_debug_meta(region_debug_dir, debug_meta)
                        raise
                    region_after_refill_snapshot = region_debug_dir / "after_refill.xlsx"
                    shutil.copyfile(output_path, region_after_refill_snapshot)
                    region_actual_edits = [
                        row for row in (patch_result.get("write_log", []) or [])
                        if isinstance(row, dict) and bool(row.get("is_actual_edit"))
                    ]
                    debug_meta["patch_result"] = {
                        "dirty": bool(patch_result.get("dirty")),
                        "usage": patch_result.get("usage", {}),
                        "write_log": patch_result.get("write_log", []),
                        "actual_edit_count": len(region_actual_edits),
                    }
                    debug_meta.setdefault("output_files", {})
                    debug_meta["output_files"]["after_refill_xlsx"] = str(region_after_refill_snapshot)
                    if not region_actual_edits:
                        shutil.copyfile(region_backup, output_path)
                        region_after_restore_snapshot = region_debug_dir / "after_restore.xlsx"
                        shutil.copyfile(output_path, region_after_restore_snapshot)
                        debug_meta["status"] = "no_actual_edits_restored"
                        debug_meta["output_files"]["after_restore_xlsx"] = str(region_after_restore_snapshot)
                        _write_region_debug_meta(region_debug_dir, debug_meta)
                        log_fn(
                            f"[REFLECT_PLUGIN] round={ridx} region={region} refill wrote no actual edits; restored backup."
                        )
                    else:
                        usage_summary["patch_agents"].append(
                            {
                                "round": ridx,
                                "region": region,
                                "usage": patch_result.get("usage", {}),
                            }
                        )
                        any_patch = True
                        debug_meta["status"] = "refill_kept"
                        _write_region_debug_meta(region_debug_dir, debug_meta)
                        log_fn(
                            f"[REFLECT_PLUGIN] round={ridx} region={region} refill done "
                            f"actual_edits={len(region_actual_edits)}"
                        )
                round_edits_after_rows = _actual_edit_records_for_round(edit_log_jsonl, ridx)
                (reflect_dir / "region_debug_index.json").write_text(
                    json.dumps(region_debug_index, ensure_ascii=False, indent=2, default=str),
                    encoding="utf-8",
                )
                new_actual_edit_rows = round_edits_after_rows[round_edits_before:]
                usage_summary["reflect_plugin_stats"]["patch_written_actual_edits_total"] += len(new_actual_edit_rows)
                log_fn(
                    f"[REFLECT_PLUGIN] round={ridx} patch_attempted_regions={patch_attempted_regions} "
                    f"patch_written_actual_edits={len(new_actual_edit_rows)} "
                    f"new_edit_cells={_format_edit_cell_list(new_actual_edit_rows)}"
                )
                if not any_patch:
                    log_fn(f"[REFLECT_PLUGIN] round={ridx} no patch applied; stop.")
                    break
        else:
            for ridx in range(1, max(1, int(args.reflect_max_rounds)) + 1):
                reflect_dir = job_dir / "reflect" / f"round_{ridx}"
                reflect_dir.mkdir(parents=True, exist_ok=True)
                try:
                    check = reflect_check_with_screenshot(
                        before_xlsx_path=before_fill_xlsx,
                        after_xlsx_path=output_path,
                        sheet_name=args.sheet,
                        instruction=args.instruction,
                        model=checker_model,
                        base_url=args.agent_base_url or None,
                        api_key=args.agent_api_key or None,
                        out_dir=reflect_dir / "check",
                        dpi=args.dpi,
                        checker_context_summary={
                            "struct_summary": first_pass_result.get("struct_summary", {}),
                            "first_pass_write_log": (first_pass_result.get("write_log", []) or [])[:40],
                        },
                        planner_policy=planner_policy,
                        model_io_dir=((model_io_dir / f"reflect_legacy_round_{ridx}" / "stage1_checker") if model_io_dir is not None else None),
                    )
                except Exception as exc:
                    log_fn(f"[REFLECT] checker failed at round={ridx}: {exc}")
                    break

                log_fn(f"[REFLECT] round={ridx} check={json.dumps(check, ensure_ascii=False)}")
                if bool(check.get("bad_range_parse_failed")):
                    log_fn(
                        "[REFLECT][WARN] "
                        f"round={ridx} screenshot checker returned non-A1 bad_ranges_raw="
                        f"{json.dumps(check.get('bad_ranges_raw', check.get('bad_range_raw', '')), ensure_ascii=False)} "
                        "so Stage 1 cannot apply a local patch."
                    )
                bad_ranges = [str(x).strip() for x in (check.get("bad_ranges") or []) if str(x).strip()]
                if not bad_ranges:
                    legacy_bad_range = str(check.get("bad_range") or "").strip()
                    if legacy_bad_range:
                        bad_ranges = [legacy_bad_range]
                need_fix = bool(check.get("need_fix")) and bool(bad_ranges)
                if not need_fix:
                    log_fn(f"[REFLECT] round={ridx} no local fix needed; stop.")
                    break
                for bad_range in bad_ranges:
                    try:
                        bad_range = _normalize_a1_range(bad_range)
                    except Exception:
                        log_fn(f"[REFLECT] round={ridx} invalid bad_range={bad_range!r}; skip.")
                        continue

                    log_fn(
                        f"[REFLECT] round={ridx} local_render focus={bad_range} "
                        f"pad(L,T,R,B)=({args.reflect_focus_pad_left_cols},{args.reflect_focus_pad_top_rows},"
                        f"{args.reflect_focus_pad_right_cols},{args.reflect_focus_pad_bottom_rows})"
                    )
                    prep = preprocess_excel_html_css(
                        xlsx_path=output_path,
                        out_dir=reflect_dir / "preprocess",
                        sheet_name=args.sheet,
                        dpi=args.dpi,
                        canvas_bg=args.render_canvas_bg,
                        png_scale=args.render_png_scale,
                        max_render_rows=args.render_max_rows,
                        max_render_cols=args.render_max_cols,
                        focus_range=bad_range,
                        focus_pad_left_cols=args.reflect_focus_pad_left_cols,
                        focus_pad_top_rows=args.reflect_focus_pad_top_rows,
                        focus_pad_right_cols=args.reflect_focus_pad_right_cols,
                        focus_pad_bottom_rows=args.reflect_focus_pad_bottom_rows,
                    )
                    qwen_out_dir = reflect_dir / "qwen_out"
                    qwen_outputs = run_qwen_small(
                        image_path=prep["sheet_png"],
                        xlsx_path=output_path,
                        out_dir=qwen_out_dir,
                        sheet_name=args.sheet,
                        ckpt_path=args.ckpt or None,
                        edges_json=prep["edges_json"],
                        bounds_json=prep["bounds_json"],
                        qwen_base_url=args.qwen_base_url or None,
                        qwen_model=args.qwen_model or None,
                        qwen_api_key=args.qwen_api_key or None,
                        skip_qwen=(args.qwen_mode == "lora"),
                        log_fn=log_fn,
                    )
                    reflect_pairs_path: Path | None = None
                    if args.qwen_mode == "lora":
                        if not args.qwen_lora_model or not args.qwen_lora_path:
                            raise PipelineError("reflect requires qwen_lora_model and qwen_lora_path when qwen_mode=lora")
                        reflect_slots_json = qwen_out_dir / "slots_in_bad_range.json"
                        kept_slots = _filter_slots_json_by_focus_range(
                            slots_json_path=qwen_outputs["slots_json"],
                            focus_range=bad_range,
                            out_json_path=reflect_slots_json,
                        )
                        if kept_slots <= 0:
                            log_fn(
                                f"[REFLECT] round={ridx} no slots in bad_range={bad_range}; "
                                "falling back to local free patch."
                            )
                        else:
                            reflect_pairs_path = run_qwen_lora(
                                image_path=prep["sheet_png"],
                                slots_json=reflect_slots_json,
                                out_dir=qwen_out_dir,
                                model_name=args.qwen_lora_model,
                                lora_path=args.qwen_lora_path,
                                device_map=args.qwen_lora_device_map,
                                max_new_tokens=args.qwen_lora_max_new_tokens,
                                temperature=args.qwen_lora_temperature,
                                slots_chunk_size=args.qwen_lora_slots_chunk_size,
                                crop_by_slot_chunk=args.qwen_lora_crop_by_slot_chunk,
                                crop_left_pad=args.qwen_lora_crop_left_pad,
                                crop_top_pad=args.qwen_lora_crop_top_pad,
                                crop_right_pad=args.qwen_lora_crop_right_pad,
                                crop_bottom_pad=args.qwen_lora_crop_bottom_pad,
                                max_image_size=args.qwen_lora_max_image_size,
                                resize_after_crop=args.qwen_lora_resize_after_crop,
                                annotate_slot_id=args.qwen_lora_annotate_slot_id,
                                log_fn=log_fn,
                                debug_dir=qwen_out_dir / "lora_debug",
                            )
                    else:
                        reflect_pairs_path = qwen_outputs["qwen_pairs"]

                    reflect_form_all = (
                        build_form_from_qwen(reflect_pairs_path, xlsx_path=output_path, sheet_name=args.sheet)
                        if reflect_pairs_path is not None
                        else {"pairs": []}
                    )
                    reflect_form = _filter_form_pairs_by_focus_range(reflect_form_all, bad_range)
                    reflect_invalid_slots = reflect_form.get("invalid_slots", []) or []
                    log_fn(
                        f"[REFLECT] round={ridx} bad_range={bad_range} pairs_total={len(reflect_form_all.get('pairs', []))} "
                        f"pairs_in_range={len(reflect_form.get('pairs', []))} "
                        f"invalid_total={len(reflect_form_all.get('invalid_slots', []))} "
                        f"invalid_in_range={len(reflect_invalid_slots)}"
                    )
                    patch_pairs = reflect_form.get("pairs", []) or []
                    patch_mapping_mode = "hint" if patch_pairs else "off"
                    if patch_mapping_mode == "hint" and not workflow_policy.get("prefer_hints_when_available", True):
                        patch_mapping_mode = "off"
                        log_fn(
                            f"[REFLECT] round={ridx} planner controller disabled hint mode in bad_range={bad_range}; "
                            "using local background/context only."
                        )
                    if not patch_pairs:
                        log_fn(
                            f"[REFLECT] round={ridx} no qwen hints in bad_range={bad_range}; "
                            "falling back to local free patch."
                        )
                    legacy_reflect_hints: list[dict[str, Any]] = []
                    for pair in patch_pairs:
                        if not isinstance(pair, dict):
                            continue
                        key = str(pair.get("key") or "").strip()
                        value_cell = str(pair.get("value_cell") or "").strip()
                        if not key or not value_cell:
                            continue
                        legacy_reflect_hints.append(
                            {
                                "key": key,
                                "target_range": bad_range,
                                "candidate_value_cell": value_cell,
                                "candidate_value": _read_anchor_value_from_workbook(output_path, args.sheet, value_cell),
                                "hint_source": "reflect_slot_kv",
                            }
                        )
                    for invalid in reflect_invalid_slots:
                        if not isinstance(invalid, dict):
                            continue
                        value_cell = str(invalid.get("value_cell") or invalid.get("candidate_value_cell") or "").strip()
                        if not value_cell:
                            continue
                        legacy_reflect_hints.append(
                            {
                                "key": str(invalid.get("key") or "").strip(),
                                "target_range": bad_range,
                                "candidate_value_cell": value_cell,
                                "candidate_value": _read_anchor_value_from_workbook(output_path, args.sheet, value_cell),
                                "hint_source": "reflect_invalid_slot",
                                "invalid_slot": True,
                                "avoid_cell": True,
                                "invalid_reason": str(invalid.get("reason") or "qwen_marked_invalid_slot"),
                            }
                        )

                    patch_instruction = (
                        f"{args.instruction}\n\n"
                        f"[Reflect patch round {ridx}]\n"
                        f"Detected likely error region: {bad_range}.\n"
                        "Only revise cells in this range; keep all other regions unchanged.\n"
                        "If no structured reflect hints are available, still inspect this local region and try one local correction pass."
                    )
                    bad_range_background_context = _collect_planner_background_context(
                        xlsx_path=output_path,
                        sheet_name=args.sheet,
                        planner_policy=planner_policy,
                        stage="refill",
                        focus_range=bad_range,
                    )
                    patch_result = agent_fill(
                        template_xlsx=output_path,
                        output_xlsx=output_path,
                        instruction=patch_instruction,
                        form={"pairs": patch_pairs},
                        model=args.agent_model,
                        base_url=args.agent_base_url,
                        api_key=args.agent_api_key or None,
                        tool_mode="code",
                        max_steps=max(1, int(args.reflect_patch_steps)),
                        skills_text=skills_text,
                        mapping_mode=patch_mapping_mode,
                        sheet_name=args.sheet,
                        allowed_write_range=bad_range,
                        edit_log_jsonl=edit_log_jsonl,
                        round_idx=ridx,
                        reflect_hints=legacy_reflect_hints,
                        context_summary={"planner_background_context": bad_range_background_context} if bad_range_background_context else None,
                        temperature=float(args.agent_temperature),
                        planner_policy=planner_policy,
                        planner_stage="refill",
                        model_io_dir=((model_io_dir / f"reflect_legacy_round_{ridx}" / f"patch_{_safe_console_text(bad_range).replace(':', '_')}") if model_io_dir is not None else None),
                    )
                    usage_summary["patch_agents"].append(
                        {"round": ridx, "region": bad_range, "usage": patch_result.get("usage", {})}
                    )
                    log_fn(f"[REFLECT] round={ridx} local patch done.")

    patch_prompt = sum(int((x.get("usage") or {}).get("prompt_tokens", 0)) for x in usage_summary["patch_agents"])
    patch_completion = sum(int((x.get("usage") or {}).get("completion_tokens", 0)) for x in usage_summary["patch_agents"])
    patch_total = sum(int((x.get("usage") or {}).get("total_tokens", 0)) for x in usage_summary["patch_agents"])
    usage_summary["patch_agents_totals"] = {
        "prompt_tokens": patch_prompt,
        "completion_tokens": patch_completion,
        "total_tokens": patch_total,
    }
    usage_summary["grand_total"] = {
        "prompt_tokens": int(usage_summary["first_pass_agent"].get("prompt_tokens", 0))
        + int(usage_summary["reflect_checker"].get("prompt_tokens", 0))
        + int(usage_summary["reflect_model_assess"].get("prompt_tokens", 0))
        + patch_prompt,
        "completion_tokens": int(usage_summary["first_pass_agent"].get("completion_tokens", 0))
        + int(usage_summary["reflect_checker"].get("completion_tokens", 0))
        + int(usage_summary["reflect_model_assess"].get("completion_tokens", 0))
        + patch_completion,
        "total_tokens": int(usage_summary["first_pass_agent"].get("total_tokens", 0))
        + int(usage_summary["reflect_checker"].get("total_tokens", 0))
        + int(usage_summary["reflect_model_assess"].get("total_tokens", 0))
        + patch_total,
    }
    if args.reflect_enable and args.reflect_plugin_enable:
        log_fn(
            f"[REFLECT_PLUGIN][FINAL] stats={json.dumps(usage_summary['reflect_plugin_stats'], ensure_ascii=False)}"
        )
    usage_summary_path = job_dir / "run_usage_summary.json"
    usage_summary_path.write_text(json.dumps(usage_summary, ensure_ascii=False, indent=2), encoding="utf-8")
    log_fn(f"[USAGE] summary={json.dumps(usage_summary['grand_total'], ensure_ascii=False)} path={usage_summary_path}")
    log_fn(f"Done. Filled file: {args.output}")


if __name__ == "__main__":
    main()
